import io
import zipfile
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy.exc import OperationalError
from sqlalchemy.orm import Session

from test_product_library import library, preview, confirm
from app.models import LibraryProduct
from app.product_library_quotations import RISK_WORDS

BASE = "/api/product-library"


def product(client, **values):
    from uuid import uuid4
    supplier = values.pop('manufacturer_name', '测试供应商')
    amount = values.pop('wholesale_price', None)
    notes = values.pop('price_note', None) or '采购条件已核对'
    payload = {'model': 'M1', 'name': '风扇', **values}
    response = client.post(BASE + '/products', json=payload)
    assert response.status_code == 201, response.text
    item = response.json()
    if amount is not None:
        base = f"{BASE}/products/{item['id']}"
        response = client.post(base + '/manual-quotes', json={'product_version': item['version'], 'supplier': supplier, 'variant': None, 'price_text': amount, 'conditions': notes, 'idempotency_key': str(uuid4())})
        assert response.status_code == 201, response.text
        source = response.json()
        response = client.post(base + f"/quotes/{source['quote_version_id']}/review", json={'product_version': item['version'] + 1, 'status': 'eligible', 'reason': '已核对出厂价', 'source_id': source['source_id']})
        assert response.status_code == 200, response.text
        item = client.get(base).json()
    return item


def test_price_filters_and_batch_partial_failure(library):
    client, _ = library
    rows = [product(client, model=f"M{i}", name=f"灯{i}", wholesale_price=p) for i, p in enumerate([None, "0.00", "10.00", "20.00"])]
    assert client.get(BASE + "/products", params={"price_mode": "between", "price_min": "0", "price_max": "10"}).json()["total"] == 2
    assert client.get(BASE + "/products", params={"price_mode": "above", "price_min": "10"}).json()["total"] == 1
    assert client.get(BASE + "/products", params={"price_mode": "below", "price_max": "10"}).json()["total"] == 1
    assert client.get(BASE + "/products", params={"price_status": "unquoted"}).json()["total"] == 1
    for params in [{"price_mode": "between", "price_min": "10", "price_max": "1"}, {"price_mode": "below"}, {"price_min": "1"}]:
        assert client.get(BASE + "/products", params=params).status_code == 422
    response = client.post(BASE + "/products/batch-delete", json={"items": [{"id": rows[0]["id"], "version": 1}, {"id": rows[1]["id"], "version": 2}, {"id": 999, "version": 1}]})
    assert response.status_code == 200
    assert {k: response.json()[k] for k in ["deleted", "missing", "stale"]} == {"deleted": 1, "missing": 1, "stale": 1}


def test_quotation_decimal_and_customer_workbook(library):
    client, _ = library
    p = product(client, name="=客户文本", parameters="普通参数", wholesale_price="1.00", price_note="PRIVATE_NOTE_913")
    data = quote(client, rules={"markup_mode": "percent", "markup_value": "0.50"})
    assert data["rows"][0]["final_price"] == "1.01"
    response = export(client, data)
    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.content))
    assert len(wb.worksheets) == 1
    assert list(wb.active.values)[0] == ("产品型号", "产品名称", "完整参数", "卖点", "最终报价", "MOQ", "报价说明")
    assert wb.active["B2"].value == "=客户文本" and wb.active["B2"].data_type == "s"
    assert Decimal(str(wb.active["E2"].value)) == Decimal("1.01")
    assert not any(c.data_type == "f" or c.comment for row in wb.active for c in row)
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        content = b"".join(archive.read(name) for name in archive.namelist())
        assert b"PRIVATE_NOTE_913" not in content and b"wholesale_price" not in content
    assert client.get(BASE + f"/products/{p['id']}").json() == p


def quote(client, **payload):
    scope = payload.get('scope', 'all')
    params = {'page_size': 100} | (payload.get('filters', {}) if scope == 'filtered' else {})
    products = client.get(BASE + '/products', params=params).json()['items']
    if scope == 'selected':
        ids = {item['id'] for item in payload['items']}
        products = [item for item in products if item['id'] in ids]
    selections = {}
    for item in products:
        detail = client.get(f"{BASE}/products/{item['id']}").json()
        current = next((q for q in detail['quotes'] if q['is_current'] and q['eligibility'] == 'eligible'), None)
        if current:
            selections[str(item['id'])] = {'quote_version_id': current['id'], 'source_id': current['source_ids'][0], 'supplier': current['supplier_name'], 'variant': current['variant'], 'final_conditions': '采购未税，运费另计', 'rules_confirmed': True}
    response = client.post(BASE + "/customer-quotations/preview", json={"scope": "all", "selections": selections, "conditions_confirmed_ids": [item['id'] for item in products], **payload})
    assert response.status_code == 200, response.text
    return response.json()


def export(client, data, **extra):
    return client.post(BASE + "/customer-quotations/export", json={"items": data["items"], "rules": data["rules"], "digest": data["digest"], "overrides": {}, "confirmed_risk_ids": [], "confirmed_all_text": True, "selections": data['selections'], "conditions_confirmed_ids": data.get("conditions_confirmed_ids", []), **extra})


@pytest.mark.parametrize("cost,rules,expected", [
    ("1.00", {"markup_mode": "percent", "markup_value": "0.40"}, "1.00"),
    ("1.00", {"markup_mode": "percent", "markup_value": "0.50"}, "1.01"),
    ("1.00", {"markup_mode": "percent", "markup_value": "1.50"}, "1.02"),
    ("10.00", {"rounding": "integer"}, "10.00"),
    ("10.01", {"rounding": "integer"}, "11.00"),
    ("10.00", {"markup_mode": "percent", "markup_value": "10", "tax_mode": "adjust", "tax_value": "5", "freight_mode": "fixed", "freight_value": "2"}, "13.55"),
    ("10.00", {"markup_mode": "fixed", "markup_value": "2", "tax_mode": "adjust", "tax_value": "5", "freight_mode": "fixed", "freight_value": "2"}, "14.60"),
    ("1.00", {"markup_mode": "percent", "markup_value": "0.49", "tax_mode": "adjust", "tax_value": "0.49"}, "1.01"),
    ("999999999.99", {}, "999999999.99"),
    ("999999999.99", {"rounding": "integer"}, None),
    ("999999999.99", {"markup_mode": "fixed", "markup_value": "0.01"}, None),
])
def test_decimal_order_rounding_and_upper_bound(library, cost, rules, expected):
    client, _ = library
    product(client, wholesale_price=cost)
    data = quote(client, rules=rules)
    assert data["rows"][0]["final_price"] == expected
    assert export(client, data).status_code == (422 if expected is None else 200)


@pytest.mark.parametrize("rules", [
    {"markup_mode": "bad"}, {"rounding": "bad"}, {"extra": "x"},
    {"markup_mode": "percent", "markup_value": "1000.01"}, {"tax_value": "1001"},
    {"freight_value": "-1"}, {"markup_value": "NaN"}, {"freight_value": "1.001"},
    {"markup_value": 1}, {"markup_value": None}, {"tax_value": "1e2"},
])
def test_invalid_rules_rejected(library, rules):
    client, _ = library
    assert client.post(BASE + "/customer-quotations/preview", json={"scope": "all", "rules": rules}).status_code == 422


def test_unquoted_manual_cost_and_text_confirmation(library):
    client, _ = library
    p = product(client, parameters="\n".join(RISK_WORDS))
    assert quote(client)["items"] == []
    data = quote(client, include_unquoted=True)
    row = data["rows"][0]
    assert row["final_price"] is None and "无法核对成本" in row["risks"]
    assert set(r["word"] for r in row["text_risks"]) == set(RISK_WORDS)
    assert export(client, data).status_code == 422
    response = export(client, data, confirmed_risk_ids=[p["id"]])
    assert response.status_code == 422
    overrides = {str(p["id"]): "12.34"}
    manual = quote(client, include_unquoted=True, overrides=overrides, rules={"markup_mode": "percent", "markup_value": "1000"})
    assert manual["rows"][0]["final_price"] is None
    assert export(client, manual, overrides=overrides, confirmed_risk_ids=[p["id"]]).status_code == 422
    for value in ["-1", "NaN", "1.001", "1000000000", 1, ""]:
        assert export(client, manual, overrides={str(p["id"]): value}, confirmed_risk_ids=[p["id"]]).status_code == 422


def test_low_cost_warning_does_not_block_and_rules_preserve_override(library):
    client, _ = library
    p = product(client, wholesale_price="10.00")
    overrides = {str(p["id"]): "9.00"}
    data = quote(client, overrides=overrides, rules={"markup_mode": "percent", "markup_value": "100"})
    assert data["rows"][0]["final_price"] == "9.00"
    assert "低于采购成本" in data["rows"][0]["risks"]
    assert export(client, data, overrides=overrides).status_code == 422
    assert export(client, data, overrides=overrides, confirmed_below_cost_ids=[p['id']]).status_code == 200


def test_snapshot_mismatch_and_stale(library):
    client, engine = library
    first = product(client, wholesale_price="10.00")
    second = product(client, model="M2", name="第二款", wholesale_price="20.00")
    data = quote(client)
    for items in [data["items"][::-1], data["items"][:1], data["items"] + [{"id": 999, "version": 1}]]:
        assert export(client, data, items=items).status_code == 409
    assert export(client, data, rules={**data["rules"], "markup_mode": "fixed"}).status_code == 409
    assert export(client, data, automatic_prices=["0"]).status_code == 422
    with Session(engine) as db:
        db.get(LibraryProduct, first["id"]).version += 1
        db.commit()
    assert client.request('DELETE', BASE + f"/products/{second['id']}", json={'version': second['version']}).status_code == 204
    response = export(client, data)
    assert response.status_code == 409
    assert {item["reason"] for item in response.json()["items"]} == {"missing", "stale"}


def test_batch_delete_commit_failure_rolls_back_and_validates_items(library, monkeypatch):
    client, engine = library
    rows = [product(client, model=f"M{i}", name=f"产品{i}") for i in range(3)]
    items = [{"id": p["id"], "version": p["version"]} for p in rows]
    for bad in [[], items + [items[0]], [{"id": True, "version": 1}], [{"id": 1, "version": 0}], [{"id": 1, "version": 1, "extra": 1}]]:
        assert client.post(BASE + "/products/batch-delete", json={"items": bad}).status_code == 422
    def fail_commit(*args):
        raise OperationalError("commit", {}, Exception("disk failure"))
    with monkeypatch.context() as patch:
        patch.setattr(Session, "commit", fail_commit)
        response = client.post(BASE + "/products/batch-delete", json={"items": items})
    assert response.status_code == 500
    assert client.get(BASE + "/products").json()["total"] == 3


def test_193_plus_8_import_and_filtered_all_scopes(library):
    client, _ = library
    from test_product_library_simple_import import workbook, HEADERS
    rows = [[None, f"M{i:03}", f"灯{i}", "3.7V", str(i), None, "便携", None] for i in range(193)]
    rows += [[None, None, None, "无名称", None, None, None, None] for _ in range(8)]
    receipt = client.post(BASE + '/imports/preview', params={'manufacturer_name': '厂家'}, headers={'X-Filename': 'test.xlsx'}, content=workbook('产品清单', HEADERS, rows)).json()
    assert sum(bool(row['errors']) for row in receipt['rows']) == 8
    decisions = [{'row_number': row['row_number'], 'action': 'skip' if row['errors'] else 'apply'} for row in receipt['rows']]
    result = client.post(BASE + f"/imports/{receipt['import_id']}/confirm", json={'decisions': decisions, 'confirmed_attachments_not_imported': True})
    assert result.status_code == 200 and sum(row['action'] == 'imported' for row in result.json()['rows']) == 193
    filters = {"q": "3.7V", "manufacturer": "厂家", "price_status": "quoted", "price_mode": "between", "price_min": "10", "price_max": "80", "sort": "price_asc"}
    ids = []
    for page in range(1, 5):
        data = client.get(BASE + "/products", params={**filters, "page": page}).json()
        assert data["total"] == 71
        ids += [p["id"] for p in data["items"]]
    assert len(set(ids)) == len(ids) == 71
    data = quote(client, scope="filtered", filters=filters)
    assert [item["id"] for item in data["items"]] == ids
    assert len(quote(client)["items"]) == 193
    selected = data["items"][::20]
    assert quote(client, scope="selected", items=selected)["items"] == selected


@pytest.mark.parametrize("tax", ["excluded", "included", "adjust"])
@pytest.mark.parametrize("freight", ["separate", "included", "fixed"])
def test_notes_and_complete_workbook_isolation(library, tax, freight):
    client, _ = library
    product(client, manufacturer_name="PRIVATE_SUPPLIER_417", wholesale_price="937.41", price_note="PRIVATE_NOTE_417", model="@literal", parameters="+text", selling_points="-text")
    data = quote(client, rules={"markup_mode": "percent", "markup_value": "17.13", "tax_mode": tax, "tax_value": "6.37", "freight_mode": freight, "freight_value": "23.49"})
    response = export(client, data)
    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.content))
    ws = wb.active
    assert ws.sheet_state == "visible" and len(wb.worksheets) == 1
    assert not any(dim.hidden for dim in ws.row_dimensions.values())
    assert not any(dim.hidden for dim in ws.column_dimensions.values())
    assert not any(cell.comment or cell.data_type == "f" for row in ws for cell in row)
    assert all(ws.cell(2, col).data_type == "s" for col in range(1, 5))
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        xml = b"".join(archive.read(name) for name in archive.namelist()).decode("utf-8", errors="replace")
    for sentinel in ["PRIVATE_SUPPLIER_417", "PRIVATE_NOTE_417", "937.41", "17.13", "6.37", "23.49", "wholesale_price", ">version<", "markup_value", "tax_value", "freight_value"]:
        assert sentinel not in xml and sentinel not in response.headers["content-disposition"]
    assert ws["F2"].value == data["moq"]
    assert (ws["G2"].value or "") == data["note"]
    assert "法定税率" not in data["note"] and "整单分摊" not in data["note"]


@pytest.mark.parametrize("field", ["model", "name", "parameters", "selling_points"])
def test_excel_illegal_text_reports_product_and_field_without_rewriting(library, field):
    client, _ = library
    p = product(client, wholesale_price="1.00", **{field: "客户文本\x01原文"})
    data = quote(client)
    assert data["rows"][0]["invalid_text_fields"] == [field]
    assert field in data["rows"][0]["error"]
    response = export(client, data)
    assert response.status_code == 422
    assert response.json()["code"] == "customer_text_invalid"
    assert response.json()["errors"] == [{"id": p["id"], "fields": [field]}]
    assert client.get(BASE + f"/products/{p['id']}").json() == p

def test_vnext_receipt_retains_source_identity_and_error_after_preview_purged(library):
    client, _ = library
    data = preview(client, [["A", "新增名称", None, None, "12.34"], ["B", None], ["C", "跳过名称"]]).json()
    payload = {'confirmed_attachments_not_imported': True, 'decisions': [{'row_number': 2, 'action': 'apply'}, {'row_number': 3, 'action': 'skip'}, {'row_number': 4, 'action': 'skip'}]}
    result = client.post(BASE + f"/imports/{data['import_id']}/confirm", json=payload).json()
    receipt = client.get(BASE + f"/imports/{data['import_id']}").json()
    assert receipt == result
    assert [(row['row_number'], row['values']['model'], row['values']['name'], row['action']) for row in receipt['rows']] == [(2, 'A', '新增名称', 'imported'), (3, 'B', None, 'skipped'), (4, 'C', '跳过名称', 'skipped')]
    assert receipt["rows"][1]["errors"]
    assert client.get(BASE + "/products").json()["total"] == 1
    assert client.post(BASE + f"/imports/{data['import_id']}/confirm", json=payload).json() == receipt


def test_vnext_auto_manual_and_excel_share_authoritative_decimal_calculation(library):
    client, _ = library
    p = product(client, wholesale_price="10.01", name="=客户名称", parameters="客户参数", selling_points="客户卖点", price_note="PRIVATE_COST")
    overrides = {str(p["id"]): "8.21"}
    rules = {"markup_mode": "percent", "markup_value": "10", "tax_mode": "adjust", "tax_value": "5", "freight_mode": "fixed", "freight_value": "2", "rounding": "integer"}
    data = quote(client, overrides=overrides, rules=rules)
    assert data["rows"][0]["automatic_price"] == "14.00"
    assert data["rows"][0]["final_price"] == "9.00"
    response = export(client, data, overrides=overrides, confirmed_below_cost_ids=[p['id']])
    wb = load_workbook(io.BytesIO(response.content))
    assert Decimal(str(wb.active['E2'].value)) == Decimal(data['rows'][0]['final_price'])
    assert wb.active['B2'].value == '=客户名称' and wb.active['B2'].data_type == 's'
    automatic = quote(client, rules=rules)
    assert automatic["rows"][0]["final_price"] == data["rows"][0]["automatic_price"]
    assert client.get(BASE + f"/products/{p['id']}").json() == p
