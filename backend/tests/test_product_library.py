import io
from datetime import timedelta

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

from app.database import Base, get_db
from app.main import app
from app.models import LibraryImport, LibraryProduct
from app.product_library_service import HEADERS, now


@pytest.fixture
def library(tmp_path):
    engine = create_engine(
        f"sqlite:///{(tmp_path / 'library.sqlite').as_posix()}",
        connect_args={"check_same_thread": False},
    )
    Base.metadata.create_all(engine)

    def session():
        with Session(engine) as db:
            yield db

    app.dependency_overrides[get_db] = session
    client = TestClient(
        app, headers={"X-Filename": "products.xlsx"}
    )  # No context manager: the real scheduler lifespan never runs.
    yield client, engine
    app.dependency_overrides.pop(get_db, None)
    client.close()
    engine.dispose()


def workbook(rows, modify=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "产品数据"
    ws.append(HEADERS)
    for row in rows:
        ws.append(row)
    if modify:
        modify(ws)
    output = io.BytesIO()
    wb.save(output)
    wb.close()
    return output.getvalue()


def preview(client, rows, manufacturer="厂家", modify=None):
    return client.post(
        "/api/product-library/imports/preview",
        params={"manufacturer_name": manufacturer},
        content=workbook(rows, modify),
        headers={
            "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        },
    )


def confirm(client, receipt, overrides=None):
    choices = [
        {
            "row_number": row["row_number"],
            "action": (overrides or {}).get(row["row_number"], row["default_action"]),
            **({"target_id": row["existing_id"], "target_version": row["version"]} if (overrides or {}).get(row["row_number"], row["default_action"]) in ("update", "append_source") else {}),
        }
        for row in receipt["rows"]
    ]
    return client.post(
        f"/api/product-library/imports/{receipt['import_id']}/confirm",
        json={"decisions": choices},
    )


def product(client, **extra):
    payload = {
        "manufacturer_name": "厂家",
        "model": "M1",
        "name": "风扇",
        "parameters": None,
        "selling_points": None,
        "wholesale_price": None,
        "price_note": None,
        **extra,
    }
    response = client.post("/api/product-library/products", json=payload)
    assert response.status_code == 201, response.text
    return response.json()


def test_hundred_and_partial_invalid_rows(library):
    client, _ = library
    rows = [
        [f"M{i}", f"产品{i}", "3.7V\r\n10W", "便携", "12.30", None] for i in range(100)
    ]
    data = preview(client, rows).json()
    assert len(data['rows']) == 100 and not any(row['errors'] for row in data['rows'])
    assert client.get("/api/product-library/products").json()["total"] == 0
    payload = {'confirmed_attachments_not_imported': True, 'decisions': [{'row_number': row['row_number'], 'action': 'apply'} for row in data['rows']]}
    endpoint = f"/api/product-library/imports/{data['import_id']}/confirm"
    result = client.post(endpoint, json=payload).json()
    assert sum(row['action'] == 'imported' for row in result['rows']) == 100
    assert client.post(endpoint, json=payload).json() == result
    rows = [[f"B{i}", f"名称{i}", None, None, None, None] for i in range(100)]
    for i in range(5):
        rows[i][0] = i
    data = preview(client, rows).json()
    assert sum(bool(row['errors']) for row in data['rows']) == 5
    assert [r["row_number"] for r in data["rows"] if r["errors"]] == [2, 3, 4, 5, 6]
    payload = {'confirmed_attachments_not_imported': True, 'decisions': [{'row_number': row['row_number'], 'action': 'skip' if row['errors'] else 'apply'} for row in data['rows']]}
    result = client.post(f"/api/product-library/imports/{data['import_id']}/confirm", json=payload).json()
    assert (sum(row['action'] == 'imported' for row in result['rows']), sum(bool(row['errors']) for row in result['rows'])) == (95, 5)


def test_update_diff_blank_zero_and_manual_clear(library):
    client, _ = library
    original = client.post('/api/product-library/products', json={'model': 'M1', 'name': '风扇', 'parameters': '旧参数', 'selling_points': '旧卖点'}).json()
    base = f"/api/product-library/products/{original['id']}"
    added = client.post(base + '/manual-quotes', json={'product_version': original['version'], 'supplier': '厂家', 'variant': None, 'price_text': '30.00', 'conditions': '未税', 'idempotency_key': 'original-price'}).json()
    reviewed = client.post(base + f"/quotes/{added['quote_version_id']}/review", json={'product_version': original['version'] + 1, 'source_id': added['source_id'], 'status': 'eligible', 'reason': '已核对出厂价'}).json()
    original = client.get(base).json()
    assert original['version'] == reviewed['product_version']
    data = preview(client, [["M1", "风扇", None, None, "0", None]]).json()
    row = data["rows"][0]
    assert row['product_id'] == original['id'] and 'parameters' not in row['differences']
    assert row['price_conflicts'][0]['incoming_cents'] == 0
    payload = {'confirmed_attachments_not_imported': True, 'decisions': [{'row_number': 2, 'action': 'apply', 'price_decisions': {'': 'keep'}}]}
    assert client.post(f"/api/product-library/imports/{data['import_id']}/confirm", json=payload).status_code == 200
    saved = client.get(f"/api/product-library/products/{original['id']}").json()
    assert saved['parameters'] == '旧参数' and saved['version'] == original['version'] + 1
    assert saved['quotes'][0]['amount_cents'] == 3000 and saved['quotes'][0]['eligibility'] == 'eligible'
    assert saved['sources'][-1]['price_text'] == '0.00'
    payload = {
        field: saved[field]
        for field in (
            "model",
            "name",
            "category",
            "parameters",
            "control_price",
            "selling_points",
            "version",
        )
    }
    payload["parameters"] = None
    response = client.put(f"/api/product-library/products/{saved['id']}", json=payload)
    assert response.json()["parameters"] is None and response.json()["version"] == saved['version'] + 1
    assert (
        client.put(
            f"/api/product-library/products/{saved['id']}", json=payload
        ).status_code
        == 409
    )


def test_existing_identity_appends_source_across_companies(library):
    client, _ = library
    first = client.post('/api/product-library/products', json={'model': 'M1', 'name': '风扇'}).json()
    response=client.post('/api/product-library/products',json={'model':'M1','name':'风扇'})
    assert response.status_code==409
    data=preview(client,[["M1","风扇"]],manufacturer='另一厂家').json()
    assert data['rows'][0]['product_id'] == first['id']
    payload = {'confirmed_attachments_not_imported': True, 'decisions': [{'row_number': 2, 'action': 'apply'}]}
    assert client.post(f"/api/product-library/imports/{data['import_id']}/confirm", json=payload).status_code == 200
    saved=client.get(f"/api/product-library/products/{first['id']}").json()
    assert saved['manufacturer_name'] is None and saved['sources'][0]['supplier_name'] == '另一厂家'
    assert saved['version']==first['version']+1
    repeated=preview(client,[["M1","风扇"]],manufacturer='另一厂家').json()
    assert repeated['rows'][0]['product_id'] == first['id']


@pytest.mark.parametrize(
    "rows,invalid",
    [
        ([["A", "同名"], ["A", "异名"], ["B", "合法"]], [False, False, False]),
        ([["A", "同名"], ["B", "同名"], [None, "同名"]], [False, False, False]),
        (
            [[None, "同名"], [None, "同名"], ["A", "同名"]],
            [True, True, False],
        ),
        ([[123, "同名"], ["123", "同名"]], [True, False]),
    ],
)
def test_file_duplicate_rules(library, rows, invalid):
    client, _ = library
    data = preview(client, rows).json()
    assert [bool(row['errors']) for row in data['rows']] == invalid
    payload = {'confirmed_attachments_not_imported': True, 'decisions': [{'row_number': row['row_number'], 'action': 'skip' if bad else 'apply'} for row, bad in zip(data['rows'], invalid)]}
    response = client.post(f"/api/product-library/imports/{data['import_id']}/confirm", json=payload)
    assert response.status_code == 200
    assert sum(row['action'] == 'imported' for row in response.json()['rows']) == invalid.count(False)


def test_partial_identity_matches_are_allowed_but_exact_null_model_is_duplicate(library):
    client, _ = library
    assert client.post('/api/product-library/products', json={'model': 'M1', 'name': '风扇'}).status_code == 201
    response = client.post(
        "/api/product-library/products",
        json={"name": "风扇", "allow_suspected_duplicate": True},
    )
    assert response.status_code == 201
    assert client.post(
        "/api/product-library/products",
        json={"name": "风扇", "allow_suspected_duplicate": True},
    ).json()["code"] == "product_exists"
    data = preview(client, [[None, "风扇"]]).json()
    assert data['rows'][0]['product_id'] == response.json()['id']


@pytest.mark.parametrize(
    "mutator",
    [
        lambda ws: setattr(ws, "title", "错误表"),
        lambda ws: setattr(ws["A1"], "value", " 产品型号"),
        lambda ws: setattr(ws["G2"], "value", "额外"),
        lambda ws: ws.merge_cells("A2:B2"),
        lambda ws: setattr(ws.row_dimensions[2], "hidden", True),
        lambda ws: setattr(ws.column_dimensions["C"], "hidden", True),
    ],
)
def test_structure_rejected(library, mutator):
    client, _ = library
    response = preview(client, [["M", "名称"]], modify=mutator)
    assert response.status_code == 400
    assert client.get("/api/product-library/products").json()["total"] == 0


@pytest.mark.parametrize(
    "field,value",
    [
        (0, 123),
        (0, True),
        (0, "=1+1"),
        (1, "X" * 201),
        (4, "-1"),
        (4, "1.001"),
        (4, "1e2"),
        (4, "NaN"),
        (4, True),
        (4, "12～20"),
    ],
)
def test_excel_bad_types_are_row_errors(library, field, value):
    client, _ = library
    row = ["001", "名称", None, None, None, None]
    row[field] = value
    response = preview(client, [row])
    assert response.status_code == 201, response.text
    assert response.json()["rows"][0]["errors"]


def test_search_sort_normalization_and_delete(library):
    from test_product_library_v11 import product as simple_product
    client, _ = library
    first = simple_product(
        client,
        manufacturer_name="  厂家\u3000 \nA ",
        model="001",
        name="风扇%",
        parameters="3.7V\r\n电池",
        selling_points="独特_卖点",
        wholesale_price="20",
    )
    second = simple_product(client, model="B", name="产品", wholesale_price="10")
    simple_product(client, model="C", name="未知")
    assert (
        first['source_companies'] == ['厂家 A'] and first["parameters"] == "3.7V\n电池"
    )
    for q in ("001", "风扇%", "电池", "独特_", "厂家 A", "%", "_"):
        assert (
            client.get("/api/product-library/products", params={"q": q}).json()["total"]
            == 1
        )
    assert (
        client.get(
            "/api/product-library/products", params={"q": "001", "manufacturer": "厂家"}
        ).json()["total"]
        == 0
    )
    assert [
        p["id"]
        for p in client.get(
            "/api/product-library/products", params={"sort": "price_asc"}
        ).json()["items"]
    ] == [second['id'], first['id'], 3]
    assert (
        client.request(
            "DELETE",
            f"/api/product-library/products/{second['id']}",
            json={"version": second['version'] + 1},
        ).status_code
        == 409
    )
    assert (
        client.request(
            "DELETE",
            f"/api/product-library/products/{second['id']}",
            json={"version": second['version']},
        ).status_code
        == 204
    )
    assert simple_product(client, model="D", name="后续")["id"] > 3


@pytest.mark.parametrize('change',['create','edit','delete','unrelated'])
def test_preview_baseline(library,change):
    client,_=library
    original=client.post('/api/product-library/products', json={'model': 'M1', 'name': '风扇'}).json()
    data=preview(client,[["NEW","新品"]] if change=='create' else [["M1","风扇"]]).json()
    if change=='create':
        assert client.post('/api/product-library/products', json={'model': 'NEW', 'name': '新品'}).status_code == 201
    elif change=='unrelated':
        assert client.post('/api/product-library/products', json={'model': 'OTHER', 'name': '无关产品'}).status_code == 201
    elif change=='delete':
        client.request('DELETE',f"/api/product-library/products/{original['id']}",json={'version':1})
    else:
        payload={field:original[field] for field in ('model','name','category','parameters','control_price','selling_points','version')}
        payload['parameters']='变化'
        assert client.put(f"/api/product-library/products/{original['id']}",json=payload).status_code==200
    response=client.post(f"/api/product-library/imports/{data['import_id']}/confirm", json={'confirmed_attachments_not_imported': True, 'decisions': [{'row_number': 2, 'action': 'apply'}]})
    assert response.status_code==(200 if change=='unrelated' else 409),response.text
    if change!='unrelated':
        assert response.json()['code']=='import_stale'


def test_expiry_decisions_and_idempotency(library):
    client, engine = library
    data = preview(client, [["M", "名称"], ["N", "名称2"]]).json()
    url = f"/api/product-library/imports/{data['import_id']}/confirm"
    for choices in (
        [],
        [{"row_number": 2, "action": "apply"}],
        [{"row_number": 2, "action": "apply"}] * 2,
    ):
        assert client.post(url, json={'confirmed_attachments_not_imported': True, 'decisions': choices}).status_code == 422
    skipped = preview(client, [['S', '跳过']]).json()
    skip_result = client.post(f"/api/product-library/imports/{skipped['import_id']}/confirm", json={'confirmed_attachments_not_imported': True, 'decisions': [{'row_number': 2, 'action': 'skip'}]})
    assert skip_result.status_code == 200 and skip_result.json()['rows'][0]['action'] == 'skipped'
    choices = [{'row_number': r['row_number'], 'action': 'apply'} for r in data['rows']]
    result = client.post(url, json={'confirmed_attachments_not_imported': True, 'decisions': choices}).json()
    reversed_choices = [
        {"row_number": r["row_number"], "action": "apply"}
        for r in reversed(data["rows"])
    ]
    assert client.post(url, json={'confirmed_attachments_not_imported': True, 'decisions': reversed_choices}).json() == result
    assert client.post(url, json={'confirmed_attachments_not_imported': True, 'decisions': [{'row_number': 2, 'action': 'skip'}, {'row_number': 3, 'action': 'apply'}]}).status_code == 409
    with Session(engine) as db:
        item = db.get(LibraryImport, data["import_id"])
        item.expires_at = now() - timedelta(hours=1)
        db.commit()
    assert client.get(url.removesuffix("/confirm")).json() == result
    data = preview(client, [["Z", "新"]]).json()
    with Session(engine) as db:
        item = db.get(LibraryImport, data["import_id"])
        item.expires_at = now() - timedelta(hours=1)
        db.commit()
    assert client.post(f"/api/product-library/imports/{data['import_id']}/confirm", json={'confirmed_attachments_not_imported': True, 'decisions': [{'row_number': 2, 'action': 'apply'}]}).status_code == 409


def test_atomic_rollback_on_second_insert(library, monkeypatch):
    client, engine = library
    data = preview(client, [["A", "A"], ["B", "B"]]).json()
    original_flush = Session.flush
    calls = 0

    def fail(db, *args, **kwargs):
        nonlocal calls
        if any(isinstance(p, LibraryProduct) for p in db.new):
            calls += 1
            if calls == 2:
                from sqlalchemy.exc import OperationalError

                raise OperationalError("injected", {}, Exception("failure"))
        return original_flush(db, *args, **kwargs)

    monkeypatch.setattr(Session, "flush", fail)
    payload = {'confirmed_attachments_not_imported': True, 'decisions': [{'row_number': row['row_number'], 'action': 'apply'} for row in data['rows']]}
    endpoint = f"/api/product-library/imports/{data['import_id']}/confirm"
    response = client.post(endpoint, json=payload)
    assert response.status_code == 500
    with Session(engine) as db:
        assert db.scalars(select(LibraryProduct)).all() == []
        assert db.get(LibraryImport, data["import_id"]).status == "ready"
    monkeypatch.setattr(Session, "flush", original_flush)
    assert sum(row['action'] == 'imported' for row in client.post(endpoint, json=payload).json()['rows']) == 2


def test_template_and_row_boundary(library):
    client, _ = library
    response = client.get("/api/product-library/template")
    wb = load_workbook(io.BytesIO(response.content))
    assert wb.sheetnames == ["产品清单"]
    assert tuple(c.value for c in wb["产品清单"][1]) == ("分类", "产品型号", "产品名称", "产品参数", "批发价", "控价", "卖点", "备注")
    assert wb["产品清单"]["B2"].number_format == "@"
    wb.close()
    assert preview(client, []).json()["code"] == "no_product_rows"
    assert (
        preview(
            client,
            [["001", "名称"]],
            modify=lambda ws: setattr(ws["A1002"], "value", "超限"),
        ).status_code
        == 413
    )
    assert (
        preview(
            client,
            [["001", "名称"]],
            modify=lambda ws: setattr(ws["A2000"], "number_format", "@"),
        ).status_code
        == 201
    )


def test_migration_isolated_roundtrip_and_constraints(tmp_path):
    from test_collection_batches_migration import run_command
    from sqlalchemy import inspect, text
    from sqlalchemy.exc import IntegrityError

    url = f"sqlite:///{(tmp_path/'migration.sqlite').as_posix()}"
    run_command(url, "20261015_15")
    engine = create_engine(url)
    with engine.begin() as connection:
        before = set(inspect(connection).get_table_names())
        connection.execute(
            text(
                "INSERT INTO competitor_groups (name,created_at) VALUES ('retained','2026-10-18')"
            )
        )
    engine.dispose()
    run_command(url, "20261018_16")
    engine = create_engine(url)
    with engine.begin() as connection:
        assert set(inspect(connection).get_table_names()) - before == {
            "library_products",
            "library_imports",
        }
        assert (
            connection.execute(text("SELECT name FROM competitor_groups")).scalar()
            == "retained"
        )
        for cents in (-1, 100000000000):
            with pytest.raises(IntegrityError):
                connection.execute(
                    text(
                        "INSERT INTO library_products(manufacturer_name,model,name,wholesale_price_cents,created_at,updated_at,version) VALUES ('A','M','P',:c,'2026-10-18','2026-10-18',1)"
                    ),
                    {"c": cents},
                )
        connection.execute(
            text(
                "INSERT INTO library_products(manufacturer_name,model,name,created_at,updated_at,version) VALUES ('A','M','P','2026-10-18','2026-10-18',1)"
            )
        )
        with pytest.raises(IntegrityError):
            connection.execute(
                text(
                    "INSERT INTO library_products(manufacturer_name,model,name,created_at,updated_at,version) VALUES ('A','M','P2','2026-10-18','2026-10-18',1)"
                )
            )
        assert connection.execute(text("PRAGMA integrity_check")).scalar() == "ok"
        assert connection.execute(text("PRAGMA foreign_key_check")).all() == []
    engine.dispose()
    run_command(url, "20261015_15", downgrade=True)
    engine = create_engine(url)
    with engine.connect() as connection:
        assert "library_products" not in inspect(connection).get_table_names()
        assert (
            connection.execute(text("SELECT name FROM competitor_groups")).scalar()
            == "retained"
        )
    engine.dispose()


def test_concurrent_confirm_and_deleted_product_replay(library):
    from concurrent.futures import ThreadPoolExecutor

    client, _ = library
    data = preview(client, [["M", "名称"]]).json()
    payload = {'confirmed_attachments_not_imported': True, 'decisions': [{'row_number': 2, 'action': 'apply'}]}
    endpoint = f"/api/product-library/imports/{data['import_id']}/confirm"
    with ThreadPoolExecutor(max_workers=2) as pool:
        responses = list(pool.map(lambda _: client.post(endpoint, json=payload), range(2)))
    assert all(r.status_code == 200 for r in responses)
    assert responses[0].json() == responses[1].json()
    product_id = responses[0].json()["rows"][0]["product_id"]
    current = client.get(f'/api/product-library/products/{product_id}').json()
    assert client.request('DELETE', f'/api/product-library/products/{product_id}', json={'version': current['version']}).status_code == 204
    assert client.post(endpoint, json=payload).json() == responses[0].json()
    assert client.get("/api/product-library/products").json()["total"] == 0


@pytest.mark.parametrize("failure", ["commit", "receipt"])
def test_receipt_and_commit_failure_rolls_back_updates(library, monkeypatch, failure):
    from sqlalchemy.exc import OperationalError

    client, engine = library
    original = client.post('/api/product-library/products', json={'model': 'M1', 'name': '风扇', 'parameters': 'before'}).json()
    data = preview(client, [["M1", "风扇", "after"], ["N", "新增"]]).json()
    endpoint = f"/api/product-library/imports/{data['import_id']}/confirm"
    payload = {'confirmed_attachments_not_imported': True, 'decisions': [{'row_number': 2, 'action': 'apply', 'update_fields': ['parameters']}, {'row_number': 3, 'action': 'apply'}]}
    original_method = getattr(Session, "commit" if failure == "commit" else "flush")

    def fail(db, *args, **kwargs):
        if failure == "commit" or any(
            isinstance(p, LibraryImport) and p.status == "committed" for p in db.dirty
        ):
            raise OperationalError("injected", {}, Exception("failure"))
        return original_method(db, *args, **kwargs)

    monkeypatch.setattr(Session, "commit" if failure == "commit" else "flush", fail)
    response = client.post(endpoint, json=payload)
    assert response.status_code == 500
    monkeypatch.setattr(
        Session, "commit" if failure == "commit" else "flush", original_method
    )
    assert (
        client.get(f"/api/product-library/products/{original['id']}").json() == original
    )
    assert client.get("/api/product-library/products").json()["total"] == 1
    with Session(engine) as db:
        assert db.get(LibraryImport, data['import_id']).status == 'ready'
    result = client.post(endpoint, json=payload)
    assert result.status_code == 200
    assert client.get(f"/api/product-library/products/{original['id']}").json()['parameters'] == 'after'


def test_xml_security_and_zip_limits(library):
    import zipfile

    client, _ = library
    original = workbook([["M", "名称"]])

    def package(extra_name, extra_content):
        out = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(original)) as source, zipfile.ZipFile(
            out, "w", zipfile.ZIP_DEFLATED
        ) as target:
            for item in source.infolist():
                target.writestr(item, source.read(item))
            target.writestr(extra_name, extra_content)
        return out.getvalue()

    for name, body, status in [
        ("xl/vbaProject.bin", b"not executable", 400),
        ("xl/embeddings/object.bin", b"object", 400),
        ("xl/externalLinks/link.xml", b"<x/>", 400),
        ("evil.xml", b'<!DOCTYPE x [<!ENTITY y "bad">]><x/>', 400),
        (
            "evil.rels",
            b'<Relationships><Relationship TargetMode="External" Target="https://example.com"/></Relationships>',
            400,
        ),
        ("huge.bin", b"0" * (50 * 1024 * 1024), 413),
    ]:
        response = client.post(
            "/api/product-library/imports/preview",
            params={"manufacturer_name": "厂家"},
            content=package(name, body),
        )
        assert response.status_code == status, response.text
    assert (
        client.post(
            "/api/product-library/imports/preview",
            params={"manufacturer_name": "厂家"},
            content=original,
            headers={"X-Filename": "wrong.xls"},
        ).status_code
        == 415
    )
    assert (
        client.post(
            "/api/product-library/imports/preview",
            params={"manufacturer_name": "厂家"},
            content=b"0" * (5 * 1024 * 1024 + 1),
        ).status_code
        == 413
    )


def test_false_dimension_and_exact_thousand_rows(library):
    import re
    import zipfile

    client, _ = library
    original = workbook([[f"M{i}", f"产品{i}"] for i in range(1000)])
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(original)) as source, zipfile.ZipFile(
        out, "w"
    ) as target:
        for item in source.infolist():
            body = source.read(item)
            if item.filename == "xl/worksheets/sheet1.xml":
                body = re.sub(
                    b'<dimension ref="[^"]+"', b'<dimension ref="A1:A1"', body
                )
            target.writestr(item, body)
    response = client.post(
        "/api/product-library/imports/preview",
        params={"manufacturer_name": "厂家"},
        content=out.getvalue(),
    )
    assert response.status_code == 201, response.text
    assert len(response.json()['rows']) == 1000
    assert not any(row['errors'] for row in response.json()['rows'])


def test_exact_package_limits(library):
    import zipfile
    from app.product_library_service import MIB

    client, _ = library
    original = workbook([["M", "名称"]])
    with zipfile.ZipFile(io.BytesIO(original)) as source:
        entries = [(item.filename, source.read(item)) for item in source.infolist()]

    def packed(extra, compression=zipfile.ZIP_DEFLATED):
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w", compression) as target:
            for name, body in [*entries, *extra]:
                target.writestr(name, body)
        return output.getvalue()

    def post(content):
        return client.post(
            "/api/product-library/imports/preview",
            params={"manufacturer_name": "厂家"},
            content=content,
        )

    for count, expected in [(1000, 201), (1001, 413)]:
        content = packed(
            [(f"padding/{i}.bin", b"") for i in range(count - len(entries))]
        )
        assert post(content).status_code == expected
    xml_total = sum(len(body) for _, body in entries)
    assert (
        post(packed([("padding.bin", b"0" * (50 * MIB - xml_total))])).status_code
        == 201
    )
    assert (
        post(packed([("padding.bin", b"0" * (50 * MIB - xml_total + 1))])).status_code
        == 413
    )
    empty = packed([("padding.bin", b"")], zipfile.ZIP_STORED)
    exact = packed([("padding.bin", b"0" * (5 * MIB - len(empty)))], zipfile.ZIP_STORED)
    assert len(exact) == 5 * MIB
    assert post(exact).status_code == 201
    assert post(exact + b"0").status_code == 413


def test_expired_preview_stays_gone_after_new_upload(library):
    client, engine = library
    data = preview(client, [["M", "名称"]]).json()
    with Session(engine) as db:
        item = db.get(LibraryImport, data["import_id"])
        item.expires_at = now() - timedelta(hours=1)
        db.commit()
    assert preview(client, [["N", "新名称"]]).status_code == 201
    response = client.get(f"/api/product-library/imports/{data['import_id']}")
    assert response.status_code == 410


def test_readonly_skips_extra_sheet_cells_and_styled_extremes(library, monkeypatch):
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet
    from openpyxl.styles import Font

    client, _ = library
    wb = Workbook()
    ws = wb.active
    ws.title = "产品数据"
    ws.append(HEADERS)
    ws.append(["001", "名称"])
    ws.cell(1048576, 16384).font = Font(bold=True)
    extra = wb.create_sheet("额外资料")
    for i in range(2000):
        extra.append([f"不导入{i}", "=1+1"])
    output = io.BytesIO()
    wb.save(output)
    wb.close()
    original = ReadOnlyWorksheet._cells_by_row
    read_sheets = []

    def observe(self, *args, **kwargs):
        read_sheets.append(self.title)
        return original(self, *args, **kwargs)

    monkeypatch.setattr(ReadOnlyWorksheet, "_cells_by_row", observe)
    response = client.post(
        "/api/product-library/imports/preview",
        params={"manufacturer_name": "厂家"},
        content=output.getvalue(),
    )
    assert response.status_code == 201, response.text
    assert len(response.json()['rows']) == 1
    assert read_sheets and set(read_sheets) == {"产品数据"}


@pytest.mark.parametrize("relation_type", ["package", "oleObject"])
def test_embedded_relationship_rejected_outside_embeddings(library, relation_type):
    import zipfile

    client, _ = library
    original = workbook([["M", "名称"]])
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(original)) as source, zipfile.ZipFile(
        output, "w"
    ) as target:
        for item in source.infolist():
            target.writestr(item, source.read(item))
        target.writestr(
            "xl/worksheets/_rels/sheet1.xml.rels",
            f'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="object" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/{relation_type}" Target="../media/book.xlsx"/></Relationships>',
        )
        target.writestr("xl/media/book.xlsx", b"embedded")
    response = client.post(
        "/api/product-library/imports/preview",
        params={"manufacturer_name": "厂家"},
        content=output.getvalue(),
    )
    assert response.status_code == 400
    assert response.json()["code"] == "invalid_workbook"


def test_preview_requires_filename(library):
    client, _ = library
    client.headers.pop("X-Filename")
    response = preview(client, [["M", "名称"]])
    assert response.status_code == 415
    assert response.json()["code"] == "unsupported_excel_format"


def test_preview_work_runs_off_event_loop(monkeypatch):
    import asyncio
    import threading
    from starlette.requests import Request
    import app.product_library as routes

    caller_thread = threading.get_ident()
    worker_threads = []

    def service(content, manufacturer, workbook_name, db):
        worker_threads.append(threading.get_ident())
        assert content == b"binary" and manufacturer == "厂家"
        return {"status": "ready"}

    monkeypatch.setattr("app.product_library_simple_import.preview", service)

    async def receive():
        return {"type": "http.request", "body": b"binary", "more_body": False}

    request = Request(
        {"type": "http", "headers": [(b"x-filename", b"data.xlsx")]}, receive
    )
    result = asyncio.run(
        routes.preview_import(request, manufacturer_name="厂家", db=None)
    )
    assert result == {"status": "ready"}
    assert worker_threads and worker_threads[0] != caller_thread


@pytest.mark.parametrize(
    "coordinate,status", [("A2", 201), ("G2", 400), ("A1002", 413)]
)
def test_error_cell_without_cached_value_is_not_ignored(library, coordinate, status):
    import zipfile

    client, _ = library
    original = workbook([])
    output = io.BytesIO()
    row = "".join(char for char in coordinate if char.isdigit())
    with zipfile.ZipFile(io.BytesIO(original)) as source, zipfile.ZipFile(
        output, "w"
    ) as target:
        for item in source.infolist():
            body = source.read(item)
            if item.filename == "xl/worksheets/sheet1.xml":
                body = body.replace(
                    b"</sheetData>",
                    f'<row r="{row}"><c r="{coordinate}" t="e"/></row></sheetData>'.encode(),
                )
            target.writestr(item, body)
    response = client.post(
        "/api/product-library/imports/preview",
        params={"manufacturer_name": "厂家"},
        content=output.getvalue(),
    )
    assert response.status_code == status, response.text
    if status == 201:
        assert response.json()["rows"][0]["errors"]
        assert response.json()["rows"][0]["row_number"] == 2


def test_dtd_rejected_in_xml_part_with_nonxml_suffix(library):
    import zipfile

    client, _ = library
    original = workbook([["M", "名称"]])
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(original)) as source, zipfile.ZipFile(
        output, "w"
    ) as target:
        for item in source.infolist():
            name, body = item.filename, source.read(item)
            if name == "xl/worksheets/sheet1.xml":
                name = "xl/worksheets/sheet1.data"
                body = b"<!DOCTYPE worksheet>" + body
            elif name in ("xl/_rels/workbook.xml.rels", "[Content_Types].xml"):
                body = body.replace(b"sheet1.xml", b"sheet1.data")
            target.writestr(name, body)
    response = client.post(
        "/api/product-library/imports/preview",
        params={"manufacturer_name": "厂家"},
        content=output.getvalue(),
    )
    assert response.status_code == 400
    assert response.json()["code"] == "invalid_workbook"


@pytest.mark.parametrize("params", [{"page": 0}, {"page_size": 101}, {"q": "x" * 201}])
def test_request_validation_uses_product_library_error_contract(library, params):
    client, _ = library
    response = client.get("/api/product-library/products", params=params)
    assert response.status_code == 422
    assert response.json()["code"] == "invalid_product"
    assert "message" in response.json() and "detail" not in response.json()


@pytest.mark.parametrize("body", [b"{invalid", b"[]"])
def test_invalid_json_body_uses_library_error_contract(library, body):
    client, _ = library
    response = client.post(
        "/api/product-library/products",
        content=body,
        headers={"Content-Type": "application/json"},
    )
    assert response.status_code == 422
    assert response.json()["code"] == "invalid_product"
    response = client.post(
        "/api/product-library/imports/00000000-0000-0000-0000-000000000000/confirm",
        content=body,
        headers={"Content-Type": "application/json"},
    )
    assert response.status_code == 422
    assert response.json()["code"] == "invalid_decisions"


def test_restart_process_reads_and_replays_durable_receipt(library):
    import json
    import os
    from pathlib import Path
    import subprocess
    import sys

    client, engine = library
    data = preview(client, [["M", "名称"]]).json()
    payload = {'confirmed_attachments_not_imported': True, 'decisions': [{'row_number': 2, 'action': 'apply'}]}
    result = client.post(f"/api/product-library/imports/{data['import_id']}/confirm", json=payload).json()
    script = r"""
import json, sys
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from app.database import get_db
from app.main import app
engine = create_engine(sys.argv[1], connect_args={"check_same_thread": False})
def isolated():
    with Session(engine) as db:
        yield db
app.dependency_overrides[get_db] = isolated
client = TestClient(app)
url = "/api/product-library/imports/" + sys.argv[2]
read = client.get(url)
replay = client.post(url + "/confirm", json={"confirmed_attachments_not_imported": True, "decisions": [{"row_number": 2, "action": "apply"}]})
assert read.status_code == 200 and replay.status_code == 200
assert read.json() == replay.json()
print(json.dumps(read.json(), ensure_ascii=True))
client.close()
engine.dispose()
"""
    environment = dict(os.environ)
    environment["PYTHONPATH"] = str(Path(__file__).parents[1])
    completed = subprocess.run(
        [sys.executable, "-c", script, str(engine.url), data["import_id"]],
        env=environment,
        capture_output=True,
        text=True,
        timeout=30,
    )
    assert completed.returncode == 0, completed.stderr
    assert json.loads(completed.stdout) == result


def test_cached_formula_is_invalid_even_when_cache_is_numeric(library):
    import zipfile

    client, _ = library
    original = workbook([["M", "名称", None, None, "=1+1"]])
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(original)) as source, zipfile.ZipFile(
        output, "w"
    ) as target:
        for item in source.infolist():
            body = source.read(item)
            if item.filename == "xl/worksheets/sheet1.xml":
                from xml.etree import ElementTree as ET

                namespace = (
                    "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
                )
                root = ET.fromstring(body)
                cell = root.find(f".//{namespace}c[@r='E2']")
                formula = cell.find(f"{namespace}f")
                cached = cell.find(f"{namespace}v")
                if cached is None:
                    cached = ET.SubElement(cell, f"{namespace}v")
                cached.text = "2"
                assert formula.text == "1+1" and cached.text == "2"
                body = ET.tostring(root, encoding="utf-8")
            target.writestr(item, body)
    response = client.post(
        "/api/product-library/imports/preview",
        params={"manufacturer_name": "厂家"},
        content=output.getvalue(),
    )
    assert response.status_code == 201
    assert response.json()["rows"][0]["errors"]
    assert any(
        error["code"] == "formula_or_error"
        for error in response.json()["rows"][0]["errors"]
    )


@pytest.mark.parametrize("value", ["¥12", "1,000", "Infinity", "1000000000", "#DIV/0!"])
def test_additional_price_validation_cases(library, value):
    client, _ = library
    data = preview(client, [["001", "名称", None, None, value]]).json()
    assert data["rows"][0]["errors"]
    assert data["rows"][0]["values"]["model"] == "001"


def test_date_cells_rejected_and_leading_zero_text_preserved(library):
    from datetime import datetime

    client, _ = library
    data = preview(
        client,
        [
            ["001", "名称"],
            ["002", "日期参数", datetime(2026, 10, 18)],
            ["003", "日期价格", None, None, datetime(2026, 10, 18)],
        ],
    ).json()
    assert [bool(row['errors']) for row in data['rows']] == [False, True, True]
    payload = {'confirmed_attachments_not_imported': True, 'decisions': [{'row_number': row['row_number'], 'action': 'skip' if row['errors'] else 'apply'} for row in data['rows']]}
    result = client.post(f"/api/product-library/imports/{data['import_id']}/confirm", json=payload).json()
    saved = client.get(
        f"/api/product-library/products/{result['rows'][0]['product_id']}"
    ).json()
    assert saved["model"] == "001"


@pytest.mark.parametrize("kind", ["order", "duplicate", "missing", "hidden"])
def test_additional_header_and_sheet_structure_cases(library, kind):
    client, _ = library

    def mutate(ws):
        if kind == "order":
            ws["A1"], ws["B1"] = HEADERS[1], HEADERS[0]
        elif kind == "duplicate":
            ws["B1"] = HEADERS[0]
        elif kind == "missing":
            ws["F1"] = None
        else:
            ws.parent.create_sheet("填写说明")
            ws.sheet_state = "hidden"

    response = preview(client, [["M", "名称"]], modify=mutate)
    assert response.status_code == 400
    assert response.json()["code"] == "invalid_structure"
    assert client.get("/api/product-library/products").json()["total"] == 0


def test_manufacturer_case_width_and_newline_fidelity(library):
    client, _ = library
    manufacturers = ["ABC", "abc", "ＡＢＣ"]
    for index, manufacturer in enumerate(manufacturers):
        item = client.post('/api/product-library/products', json={'name': '风扇', 'model': f'M{index}', 'parameters': '3.7V\r3000mAh\r\n10W\n±5%', 'selling_points': ' 保留\n第二行 '}).json()
        response = client.post(f"/api/product-library/products/{item['id']}/manual-quotes", json={'product_version': item['version'], 'supplier': manufacturer, 'variant': None, 'price_text': '0', 'conditions': '已核对', 'idempotency_key': f'supplier-{index}'})
        assert response.status_code == 201, response.text
    assert client.get("/api/product-library/products").json()["total"] == 3
    assert set(client.get("/api/product-library/manufacturers").json()["items"]) == set(
        manufacturers
    )
    for manufacturer in manufacturers:
        saved = client.get(
            "/api/product-library/products", params={"manufacturer": manufacturer}
        ).json()["items"]
        assert len(saved) == 1
        assert saved[0]["parameters"] == "3.7V\n3000mAh\n10W\n±5%"
        assert saved[0]["selling_points"] == "保留\n第二行"


def test_real_sqlite_busy_confirm_retries_original_receipt(library):
    client, engine = library
    data = preview(client, [["M", "名称"]]).json()
    original_override = app.dependency_overrides[get_db]

    def short_timeout_session():
        from sqlalchemy import text

        with Session(engine) as db:
            db.execute(text("PRAGMA busy_timeout=1"))
            db.commit()
            yield db

    app.dependency_overrides[get_db] = short_timeout_session
    connection = engine.raw_connection()
    payload = {'confirmed_attachments_not_imported': True, 'decisions': [{'row_number': 2, 'action': 'apply'}]}
    endpoint = f"/api/product-library/imports/{data['import_id']}/confirm"
    try:
        connection.execute("BEGIN IMMEDIATE")
        response = client.post(endpoint, json=payload)
        assert response.status_code == 503
        assert response.json()["code"] == "database_busy"
    finally:
        connection.rollback()
        connection.close()
        app.dependency_overrides[get_db] = original_override
    with Session(engine) as db:
        assert db.get(LibraryImport, data['import_id']).status == 'ready'
    assert client.get("/api/product-library/products").json()["total"] == 0
    assert client.post(endpoint, json=payload).json()['rows'][0]['action'] == 'imported'
