import io

from openpyxl import Workbook, load_workbook
from fastapi import HTTPException
import pytest
from sqlalchemy import event, text
from sqlalchemy.engine import Engine

from app.product_library_simple_import import HEADERS, parse, template_bytes
from test_product_library import library


def workbook(sheet, headers, rows, extra=None):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(headers)
    for row in rows:
        ws.append(row)
    if extra:
        wb.create_sheet(extra)
    result = io.BytesIO()
    wb.save(result)
    return result.getvalue()


def test_template_and_multi_spec_parse():
    data = template_bytes()
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ['产品清单']
    assert tuple(cell.value for cell in wb.active[1]) == HEADERS
    assert wb.active['B2'].number_format == '@'
    result = parse(workbook('产品清单', HEADERS, [['风扇', '001', '风扇', '参数', '红：12；蓝：13', '控价', '卖点', '条件']]))
    assert result['rows'][0]['prices'] == [('红', 1200), ('蓝', 1300)]
    assert result['rows'][0]['values']['model'] == '001'
    assert len(result['sha256']) == 64


def test_legacy_with_instructions_and_numeric_model_error():
    data = workbook('产品数据', ['产品型号', '产品名称', '产品参数', '产品卖点', '批发价', '价格备注'], [[123, '名称', None, None, 10, '原备注']], '填写说明')
    result = parse(data)
    assert result['template_version'] == 1
    assert result['rows'][0]['errors']
    assert result['rows'][0]['raw_values']['price_note'] == '原备注'


def test_invalid_model_does_not_claim_null_model_identity(library):
    client, _ = library
    data = workbook('产品数据', ['产品型号', '产品名称', '产品参数', '产品卖点', '批发价', '价格备注'], [[123, '名称'], [None, '名称']])
    response = client.post('/api/product-library/imports/preview?manufacturer_name=A', headers={'X-Filename': 'test.xlsx'}, content=data)
    assert response.status_code == 201
    rows = response.json()['rows']
    assert rows[0]['errors'] and not rows[1]['errors']


def test_supplier_attribution_conflict_blocks_row(library):
    client, _ = library
    data = workbook('产品清单', HEADERS, [[None, 'M', '名称', None, '12', None, None, '供应商：另一公司']])
    response = client.post('/api/product-library/imports/preview?manufacturer_name=A', headers={'X-Filename': 'test.xlsx'}, content=data)
    assert response.status_code == 201
    preview = response.json()
    assert preview['rows'][0]['errors'][0]['field'] == 'notes'
    result = client.post(f"/api/product-library/imports/{preview['import_id']}/confirm", json={'confirmed_attachments_not_imported': True, 'decisions': [{'row_number': 2, 'action': 'apply'}]})
    assert result.status_code == 422


@pytest.mark.parametrize('label', ['供应商名称', '厂家名称', '供货商名称'])
def test_supplier_name_label_conflict_blocks_row(library, label):
    client, _ = library
    data = workbook('产品清单', HEADERS, [[None, 'M', '名称', None, '12', None, None, f'{label}：另一公司']])
    result = client.post('/api/product-library/imports/preview?manufacturer_name=A', headers={'X-Filename': 'test.xlsx'}, content=data).json()
    assert result['rows'][0]['errors']


def test_ambiguous_price_preview_and_source_are_preserved(library):
    client, _ = library
    data = workbook('产品清单', HEADERS, [[None, 'M', '名称', None, '定制价：26', None, None, '条件']])
    preview = client.post('/api/product-library/imports/preview?manufacturer_name=A', headers={'X-Filename': 'test.xlsx'}, content=data).json()
    assert preview['rows'][0]['price_status'] == 'pending' and preview['rows'][0]['price_reason']
    result = client.post(f"/api/product-library/imports/{preview['import_id']}/confirm", json={'confirmed_attachments_not_imported': True, 'decisions': [{'row_number': 2, 'action': 'apply'}]})
    assert result.status_code == 200
    detail = client.get('/api/product-library/products/1').json()
    assert detail['quotes'] == [] and detail['sources'][0]['price_text'] == '定制价：26'
    assert detail['sources'][0]['raw_values']['values']['批发价'] == '定制价：26'


@pytest.mark.parametrize('note', ['定制价', '供货价', '一件代发'])
def test_non_wholesale_note_cannot_make_price_eligible(library, note):
    client, _ = library
    data = workbook('产品清单', HEADERS, [[None, 'M', '名称', None, '26', None, None, note]])
    preview = client.post('/api/product-library/imports/preview?manufacturer_name=A', headers={'X-Filename': 'test.xlsx'}, content=data).json()
    assert preview['rows'][0]['price_status'] == 'pending'
    response = client.post(f"/api/product-library/imports/{preview['import_id']}/confirm", json={'confirmed_attachments_not_imported': True, 'decisions': [{'row_number': 2, 'action': 'apply'}]})
    assert response.status_code == 200
    assert client.get('/api/product-library/products/1').json()['quotes'][0]['eligibility'] == 'pending'


def test_reject_extra_sheet_and_preserve_formula_as_error():
    with pytest.raises(HTTPException):
        parse(workbook('产品清单', HEADERS, [['类', 'M', '名称']], '其他'))
    row = parse(workbook('产品清单', HEADERS, [['类', 'M', '名称', '=1+1']]))['rows'][0]
    assert row['errors'][0]['field'] == 'parameters'


def test_import_multi_supplier_replay_and_price_conflict(library):
    client, _ = library
    endpoint = '/api/product-library/imports/preview?manufacturer_name=A'
    rows = [['风扇', '001', '风扇', '原参数', '红：12；蓝：13', '控价', '卖点', '含税'], ['灯', '002', '灯', None, None, None, None, None]]
    content = workbook('产品清单', HEADERS, rows)
    def upload(data, supplier='A'):
        response = client.post('/api/product-library/imports/preview', params={'manufacturer_name': supplier}, headers={'X-Filename': 'price.xlsx'}, content=data)
        assert response.status_code == 201, response.text
        return response.json()
    def commit(preview, decisions):
        return client.post(f"/api/product-library/imports/{preview['import_id']}/confirm", json={'decisions': decisions, 'confirmed_attachments_not_imported': True})
    first = upload(content)
    decisions = [{'row_number': row['row_number'], 'action': 'apply'} for row in first['rows']]
    response = commit(first, decisions)
    assert response.status_code == 200, response.text
    assert commit(first, decisions).json() == response.json()
    repeated = upload(content)
    assert commit(repeated, decisions).status_code == 200
    item = client.get('/api/product-library/products/1').json()
    assert len(item['sources']) == 1 and len(item['quotes']) == 2
    other = upload(content, 'B')
    assert commit(other, decisions).status_code == 200
    assert len(client.get('/api/product-library/products/1').json()['sources']) == 2
    rows[0][4] = '红：15；蓝：13'
    rows[0][3] = '新参数'
    changed = upload(workbook('产品清单', HEADERS, rows))
    assert changed['rows'][0]['price_conflicts'][0]['variant'] == '红'
    changed_decisions = [{'row_number': 2, 'action': 'apply', 'update_fields': ['parameters'], 'price_decisions': {'红': 'update'}}, {'row_number': 3, 'action': 'skip'}]
    result = commit(changed, changed_decisions)
    assert result.status_code == 200, result.text
    item = client.get('/api/product-library/products/1').json()
    assert item['parameters'] == '新参数'
    assert len([q for q in item['quotes'] if q['supplier_name'] == 'A' and q['variant'] == '红']) == 2


def test_import_write_failure_rolls_back_and_can_retry(library):
    client, engine = library
    content = workbook('产品清单', HEADERS, [['类', 'M', '产品', None, '12', None, None, '含税']])
    preview = client.post('/api/product-library/imports/preview?manufacturer_name=A', headers={'X-Filename': 'test.xlsx'}, content=content).json()
    endpoint = f"/api/product-library/imports/{preview['import_id']}/confirm"
    payload = {'confirmed_attachments_not_imported': True, 'decisions': [{'row_number': 2, 'action': 'apply'}]}
    def fail(connection, cursor, statement, parameters, context, executemany):
        if statement.startswith('INSERT INTO library_simple_quote_versions'):
            raise RuntimeError('injected import failure')
    event.listen(Engine, 'before_cursor_execute', fail)
    try:
        with pytest.raises(RuntimeError, match='injected'):
            client.post(endpoint, json=payload)
    finally:
        event.remove(Engine, 'before_cursor_execute', fail)
    with engine.connect() as connection:
        assert connection.execute(text('SELECT count(*) FROM library_products')).scalar_one() == 0
        assert connection.execute(text('SELECT count(*) FROM library_simple_sources')).scalar_one() == 0
        assert connection.execute(text("SELECT status FROM library_imports WHERE import_id=:id"), {'id': preview['import_id']}).scalar_one() == 'ready'
    assert client.post(endpoint, json=payload).status_code == 200


def test_import_rejects_stale_preview(library):
    client, _ = library
    content = workbook('产品清单', HEADERS, [['类', 'M', '产品', None, '12']])
    preview = client.post('/api/product-library/imports/preview?manufacturer_name=A', headers={'X-Filename': 'test.xlsx'}, content=content).json()
    assert client.post('/api/product-library/products', json={'name': '产品', 'model': 'M'}).status_code == 201
    response = client.post(f"/api/product-library/imports/{preview['import_id']}/confirm", json={'confirmed_attachments_not_imported': True, 'decisions': [{'row_number': 2, 'action': 'apply'}]})
    assert response.status_code == 409


def test_legacy_same_amount_does_not_become_eligible_evidence(library):
    client, _ = library
    def do_import(data):
        preview = client.post('/api/product-library/imports/preview?manufacturer_name=A', headers={'X-Filename': 'test.xlsx'}, content=data).json()
        result = client.post(f"/api/product-library/imports/{preview['import_id']}/confirm", json={'confirmed_attachments_not_imported': True, 'decisions': [{'row_number': 2, 'action': 'apply'}]})
        assert result.status_code == 200, result.text
    do_import(workbook('产品清单', HEADERS, [[None, 'M', '产品', None, '12']]))
    do_import(workbook('产品数据', ['产品型号', '产品名称', '产品参数', '产品卖点', '批发价', '价格备注'], [['M', '产品', None, None, 12, '旧备注']], '填写说明'))
    item = client.get('/api/product-library/products/1').json()
    assert len(item['sources']) == 2
    assert len(item['quotes']) == 1 and item['quotes'][0]['eligibility'] == 'eligible'
    assert len(item['quotes'][0]['source_ids']) == 1


def test_explicit_non_wholesale_source_is_never_auto_eligible(library):
    client, _ = library
    data = workbook('产品清单', HEADERS, [[None, 'M', '产品', None, '12', '平台控价 30', None, '一件代发报价']])
    preview = client.post('/api/product-library/imports/preview?manufacturer_name=A', headers={'X-Filename': 'test.xlsx'}, content=data).json()
    response = client.post(f"/api/product-library/imports/{preview['import_id']}/confirm", json={'confirmed_attachments_not_imported': True, 'decisions': [{'row_number': 2, 'action': 'apply'}]})
    assert response.status_code == 200
    product = client.get('/api/product-library/products/1').json()
    assert product['control_price'] == '平台控价 30'
    assert product['quotes'][0]['eligibility'] == 'pending'
    quote = product['quotes'][0]
    review = client.post('/api/product-library/products/1/quotes/%s/review' % quote['id'], json={'product_version': product['version'], 'source_id': quote['source_ids'][0], 'status': 'eligible', 'reason': '确认'})
    assert review.status_code == 422
