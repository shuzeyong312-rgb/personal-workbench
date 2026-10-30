import io
import zipfile

import pytest
from openpyxl import load_workbook

from test_product_library import library

BASE = '/api/product-library'


def product(client, wholesale_price, price_note='采购条件已核对'):
    response = client.post(BASE + '/products', json={'model': 'M1', 'name': '风扇'})
    assert response.status_code == 201, response.text
    item = response.json()
    base = f"{BASE}/products/{item['id']}"
    source = client.post(base + '/manual-quotes', json={'product_version': item['version'], 'supplier': '测试供应商', 'variant': None, 'price_text': wholesale_price, 'conditions': price_note, 'idempotency_key': 'customer-test'}).json()
    reviewed = client.post(base + f"/quotes/{source['quote_version_id']}/review", json={'product_version': item['version'] + 1, 'status': 'eligible', 'reason': '已核对出厂价', 'source_id': source['source_id']})
    assert reviewed.status_code == 200, reviewed.text
    return client.get(base).json()


def quote(client, **payload):
    products = client.get(BASE + '/products').json()['items']
    selections = {}
    for item in products:
        detail = client.get(f"{BASE}/products/{item['id']}").json()
        current = next(q for q in detail['quotes'] if q['is_current'] and q['eligibility'] == 'eligible')
        selections[str(item['id'])] = {'quote_version_id': current['id'], 'source_id': current['source_ids'][0], 'supplier': current['supplier_name'], 'variant': current['variant'], 'final_conditions': '采购未税，运费另计', 'rules_confirmed': True}
    response = client.post(BASE + '/customer-quotations/preview', json={'scope': 'all', 'selections': selections, 'conditions_confirmed_ids': [p['id'] for p in products], **payload})
    assert response.status_code == 200, response.text
    return response.json()


def export(client, data, **extra):
    return client.post(BASE + '/customer-quotations/export', json={'items': data['items'], 'rules': data['rules'], 'digest': data['digest'], 'overrides': {}, 'confirmed_risk_ids': [], 'confirmed_all_text': True, 'selections': data['selections'], 'conditions_confirmed_ids': data['conditions_confirmed_ids'], **extra})


def test_customer_conditions_roundtrip_and_preview_integrity(library):
    client, _ = library
    original = product(client, wholesale_price="937.41", price_note="PRIVATE_INTERNAL_NOTE")
    overrides = {str(original['id']): '1000.21'}
    note = '=客户说明\r\n含税，运费另计\r交期双方确认'
    data = quote(client, moq=120, note=note, overrides=overrides, rules={'rounding': 'integer'})
    assert data['note'] == '=客户说明\n含税，运费另计\n交期双方确认'
    assert data['rows'][0]['original_price'] == '937.41'
    assert data['rows'][0]['final_price'] == '1001.00'
    options = dict(moq=data['moq'], note=data['note'], overrides=overrides)
    for changed in [dict(moq=121), dict(note='其他说明'), dict(overrides={str(original['id']): '1002'})]:
        assert export(client, data, **(options | changed)).status_code == 409
    response = export(client, data, **options)
    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.content))
    ws = wb.active
    assert list(ws.values)[1][4:] == (1001, 120, data['note'])
    assert ws.cell(ws.max_row, 7).value == data['terms']
    assert ws['G2'].data_type == 's' and ws['G2'].alignment.wrap_text
    assert len(wb.worksheets) == 1 and ws.sheet_state == 'visible'
    assert not any(c.data_type == 'f' or c.comment for row in ws for c in row)
    assert not any(d.hidden for d in ws.column_dimensions.values())
    assert not any(d.hidden for d in ws.row_dimensions.values())
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        xml = b''.join(archive.read(n) for n in archive.namelist())
        for value in [b'937.41', b'PRIVATE_INTERNAL_NOTE', b'wholesale_price', b'original_price']:
            assert value not in xml
    assert client.get(BASE + f"/products/{original['id']}").json() == original


@pytest.mark.parametrize('moq', [0, -1, True, 1.2, '10', None, 1000000000])
def test_invalid_moq_on_both_endpoints(library, moq):
    client, _ = library
    product(client, wholesale_price='10')
    data = quote(client)
    assert client.post(BASE + '/customer-quotations/preview', json={'scope': 'all', 'moq': moq}).status_code == 422
    assert export(client, data, moq=moq).status_code == 422


@pytest.mark.parametrize('note', [None, 123, 'x'*2001, 'bad\x00text', '\ud800', '\uffff'])
def test_invalid_note_on_both_endpoints(library, note):
    client, _ = library
    product(client, wholesale_price='10')
    data = quote(client)
    # ASCII JSON also exercises unpaired surrogate validation at the API boundary.
    import json
    assert client.post(BASE + '/customer-quotations/preview', content=json.dumps({'scope': 'all', 'note': note}), headers={'content-type': 'application/json'}).status_code == 422
    if isinstance(note, str) and '\ud800' in note:
        return
    assert export(client, data, note=note).status_code == 422


def test_note_safety_confirmation_and_empty_and_length_boundary(library):
    client, _ = library
    product(client, wholesale_price='10')
    data = quote(client, note='内部价需核查', moq=999999999)
    assert data['note_risks'][0]['field'] == 'note'
    assert export(client, data, note=data['note'], moq=data['moq']).status_code == 422
    assert export(client, data, note=data['note'], moq=data['moq'], confirmed_note=True).status_code == 200
    assert export(client, data, note=data['note'], moq=data['moq'], confirmed_note=True, confirmed_all_text=False).status_code == 422
    for note in ['', '字'*2000, '+literal\n@literal\n-literal']:
        data = quote(client, note=note)
        response = export(client, data, note=note)
        assert response.status_code == 200
        assert (load_workbook(io.BytesIO(response.content)).active['G2'].value or '') == note
