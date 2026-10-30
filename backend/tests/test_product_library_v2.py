import copy
import io
import zipfile

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import event, text
from sqlalchemy.orm import Session

from app.models import LibraryImport, LibraryProduct
from app.product_library_v2_import import SHEETS
from test_product_library import library, product, preview as preview_v1, confirm as confirm_v1

BASE = '/api/product-library'


def workbook(tables=None):
    wb=Workbook()
    wb.remove(wb.active)
    for sheet,headers in SHEETS.items():
        ws=wb.create_sheet(sheet)
        ws.append(headers)
        for row in (tables or {}).get(sheet,[]):
            ws.append(row)
    output=io.BytesIO()
    wb.save(output)
    return output.getvalue()


def upload(client,tables,company='资料公司'):
    return client.post(BASE+'/imports/preview',params={'manufacturer_name':company},content=workbook(tables))


def plan(preview, overrides=None):
    choices=[]
    for row in preview['rows']:
        choice=dict(row_key=row['row_key'],action=row['default_action'])
        choice.update((overrides or {}).get((row['row_key']['sheet'],row['row_key']['temp_id']),{}))
        if choice['action'] in ('update_main','use_existing','append_source','update','reuse'):
            choice.update(target_id=row['target_id'],target_version=row['target_version'])
        choices.append(choice)
    return dict(decisions=choices,confirmed_attachments_not_imported=True)


def confirm(client,preview,payload=None):
    return client.post(BASE+f"/imports/{preview['import_id']}/confirm",json=payload or plan(preview))


def tables():
    return {
        '产品数据':[['P','H8','暖手袋',None,'原文参数']],
        '规格款式':[['V1','P','5000mAh'],['V2','P','10000mAh']],
        '资料索引':[['S','P','报价公司',None,'原件.pdf','2','2026-10-01','quote','摘录']],
        '价格明细':[
            ['A','P','V1','批发价',None,'12.30','CNY','件','20','件','excluded','excluded',None,None,'2026-10-01','S','false','20件起'],
            ['B','P','V1','一件代发价',None,'15','CNY','件',None,None,None,None,None,None,None,'S'],
            ['C','P','V2','批发价',None,'20','CNY','件',None,None,None,None,None,None,None,'S'],
            ['D','P','V2','一件代发价',None,'25','CNY','件',None,None,None,None,None,None,None,'S'],
        ],
        '参数明细':[['P','V1','S','供电/电池','容量','5000','mAh',None,'原文5000']],
    }


def detail(client,product_id=1):
    response=client.get(BASE+f'/products/{product_id}')
    assert response.status_code==200,response.text
    return response.json()


def test_five_sheet_atomic_import_replay_and_unknowns(library):
    client,engine=library
    data=upload(client,tables()).json()
    assert data['template_version']==2 and data['identity_rule_version']==4
    with Session(engine) as db:
        assert db.query(LibraryProduct).count()==0
    assert confirm(client,data,{'decisions':plan(data)['decisions']}).json()['code']=='attachments_not_confirmed'
    payload=plan(data)
    response=confirm(client,data,payload)
    assert response.status_code==200,response.text
    assert confirm(client,data,payload).json()==response.json()
    p=detail(client)
    assert len(p['prices'])==4 and len(p['variants'])==2 and len(p['sources'])==2
    assert p['wholesale_price'] is None and p['price_state']=='multiple'
    assert p['variants'][0]['structured_attributes'][0]['source_id']==p['sources'][0]['id']
    assert p['sources'][0]['claimed_manufacturer_name'] is None
    repeat=upload(client,tables()).json()
    assert all(r['default_action'] in ('use_existing','reuse') for r in repeat['rows'])
    assert confirm(client,repeat).json()['code']=='no_changes_selected'


def test_cross_source_does_not_overwrite_main_or_default(library):
    client,_=library
    old=product(client,model='X',name='产品',parameters='旧参数',wholesale_price='9.00')
    data=upload(client,{'产品数据':[['P','X','产品',None,'新参数',None,'2.00']]}).json()
    assert confirm(client,data).status_code==200
    p=detail(client,old['id'])
    assert p['manufacturer_name']=='厂家' and p['parameters']=='旧参数' and p['wholesale_price']=='9.00'
    assert len(p['prices'])==2 and any('新参数' in (s['source_excerpt'] or '') for s in p['sources'])
    assert p['version']==old['version']+1
    assert client.get(BASE+'/products',params={'manufacturer':'资料公司'}).json()['total']==1


@pytest.mark.parametrize('kind,mutation',[
    ('orphan',lambda t:t['价格明细'][0].__setitem__(2,'missing')),
    ('duplicate_id',lambda t:t['规格款式'][1].__setitem__(0,'V1')),
    ('cross_product',lambda t:t['产品数据'].append(['Q','Q','另一个'])),
])
def test_structural_references(library,kind,mutation):
    client,_=library
    t=tables();mutation(t)
    if kind=='cross_product':
        t['规格款式'][0][1]='Q'
    response=upload(client,t)
    assert response.status_code==400,response.text


def test_skipped_source_blocks_dependent_insert(library):
    client,_=library
    data=upload(client,tables()).json()
    payload=plan(data,{('资料索引','S'):{'action':'skip'}})
    assert confirm(client,data,payload).json()['code']=='invalid_decisions'
    assert client.get(BASE+'/products').json()['total']==0


def test_import_failure_rolls_back_all_relations_and_receipt(library):
    client,engine=library
    data=upload(client,tables()).json()
    def fail(connection,cursor,statement,parameters,context,executemany):
        if statement.startswith('INSERT INTO library_product_prices'):
            from sqlalchemy.exc import OperationalError
            raise OperationalError(statement,parameters,Exception('injected'))
    event.listen(engine,'before_cursor_execute',fail)
    try:
        assert confirm(client,data).json()['code']=='import_save_failed'
    finally:
        event.remove(engine,'before_cursor_execute',fail)
    with Session(engine) as db:
        assert db.query(LibraryProduct).count()==0
        assert db.get(LibraryImport,data['import_id']).status=='ready'
    assert confirm(client,data).status_code==200


def test_conflicting_attributes_preserve_facts(library):
    client,_=library
    t={'产品数据':[['P','GS','风扇']], '参数明细':[['P',None,None,'包装','尺寸','10cm',None,None,'10cm'],['P',None,None,'包装','尺寸','12cm',None,None,'12cm']]}
    data=upload(client,t).json()
    assert confirm(client,data).status_code==200
    values=detail(client)['structured_attributes']
    assert {a['value'] for a in values}=={'10cm','12cm'} and all(a['review_status']=='needs_review' for a in values)


def test_detail_versions_default_projection_and_delete_dependencies(library):
    client,_=library
    p=product(client,wholesale_price='12.30')
    p=detail(client,p['id']);price=p['prices'][0];source=p['sources'][0]
    response=client.request('DELETE',BASE+f"/products/{p['id']}/sources/{source['id']}",json={'product_version':p['version'],'version':source['version']})
    assert response.json()['code']=='related_record_in_use'
    response=client.put(BASE+f"/products/{p['id']}/prices/{price['id']}",json={'product_version':p['version'],'version':price['version'],'amount_text':'15.00'})
    assert response.status_code==200,response.text
    p=detail(client,p['id'])
    assert p['wholesale_price']=='15.00' and p['version']==2
    payload={field:p[field] for field in ('manufacturer_name','model','name','parameters','selling_points','wholesale_price','price_note','version')}
    payload['wholesale_price']=None
    assert client.put(BASE+f"/products/{p['id']}",json=payload).status_code==200
    p=detail(client,p['id'])
    assert p['wholesale_price'] is None and len(p['prices'])==1 and not p['prices'][0]['is_default_wholesale']


def test_quote_explicit_selection_conditions_and_no_internal_xlsx(library):
    client,_=library
    t=tables();t['资料索引'][0][2]='INTERNAL_SENTINEL_COMPANY'
    t['价格明细'].append(['CONTROL','P',None,'平台控价',None,'59','CNY','件',None,None,None,None,None,None,None,'S'])
    data=upload(client,t).json()
    assert confirm(client,data).status_code==200
    p=detail(client);v=p['variants'][0];cost=p['prices'][0];control=p['prices'][-1]
    base=dict(scope='selected',items=[{'id':p['id'],'version':p['version']}],include_unquoted=True)
    invalid=client.post(BASE+'/customer-quotations/preview',json={**base,'selections':{'1':{'variant_id':v['id'],'price_id':control['id'],'price_version':control['version']}}})
    assert invalid.status_code==422
    selections={'1':{'variant_id':v['id'],'price_id':cost['id'],'price_version':cost['version']}}
    q=client.post(BASE+'/customer-quotations/preview',json={**base,'selections':selections}).json()
    assert q['rows'][0]['error'] and q['rows'][0]['final_price'] is None
    q=client.post(BASE+'/customer-quotations/preview',json={**base,'selections':selections,'conditions_confirmed_ids':[1]}).json()
    assert q['rows'][0]['final_price']=='12.30' and '5000mAh' in q['rows'][0]['name']
    payload={k:q[k] for k in ('items','rules','digest','moq','note','selections','conditions_confirmed_ids')}
    payload.update(overrides={},confirmed_risk_ids=[],confirmed_all_text=True)
    response=client.post(BASE+'/customer-quotations/export',json=payload)
    assert response.status_code==200,response.text
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        assert all(b'INTERNAL_SENTINEL' not in archive.read(name) for name in archive.namelist())
    ws=load_workbook(io.BytesIO(response.content)).active
    assert ws.max_column==7 and '5000mAh' in ws.cell(2,2).value


@pytest.mark.parametrize('amount,currency,unit,default,valid',[
    ('0','CNY','件','true',True), (None,None,None,'false',True),
    ('1.234','USD','箱','false',True), ('1.234','CNY','件','false',False),
    ('NaN','CNY','件','false',False), ('1e3','USD','件','false',False),
    ('12',None,'件','true',False), ('12','CNY','箱','true',False),
    ('12','BAD','件','false',False), ('-1','CNY','件','false',False),
])
def test_price_value_boundaries(library,amount,currency,unit,default,valid):
    client,_=library
    t={'产品数据':[['P','M','产品']], '价格明细':[['A','P',None,'批发价',None,amount,currency,unit,None,None,None,None,None,None,None,None,default]]}
    data=upload(client,t).json()
    row=next(r for r in data['rows'] if r['kind']=='prices')
    assert bool(row['errors']) is not valid
    assert row['default_action']==('insert' if valid else 'skip')
    response=confirm(client,data)
    assert response.status_code==200,response.text
    p=detail(client)
    assert len(p['prices'])==(1 if valid else 0)
    if valid and amount is None:
        assert p['prices'][0]['amount_text'] is None


def test_parameter_dedup_and_variant_version_once(library):
    client,_=library
    data=upload(client,tables()).json()
    assert confirm(client,data).status_code==200
    p=detail(client)
    t=tables()
    t['参数明细']=[['P','V1','S','基本','长','10','cm',None,'10'],['P','V1','S','基本','宽','5','cm',None,'5'],['P','V1','S','基本','宽','5','cm',None,'5']]
    data=upload(client,t).json()
    duplicate=[row for row in data['rows'] if row['kind']=='attribute'][-1]
    assert duplicate['allowed_actions']==['skip']
    assert confirm(client,data).status_code==200
    changed=detail(client)
    assert changed['variants'][0]['version']==p['variants'][0]['version']+1
    assert changed['version']==p['version']+1
    assert len(changed['variants'][0]['structured_attributes'])==3


def test_changed_related_record_invalidates_import_and_quote(library):
    client,_=library
    data=upload(client,tables()).json();assert confirm(client,data).status_code==200
    p=detail(client)
    ready=upload(client,tables()).json()
    q=client.post(BASE+'/customer-quotations/preview',json={'scope':'selected','items':[{'id':p['id'],'version':p['version']}],'include_unquoted':True}).json()
    source=p['sources'][0]
    assert client.put(BASE+f"/products/{p['id']}/sources/{source['id']}",json={'product_version':p['version'],'version':source['version'],'notes':'new'}).status_code==200
    assert confirm(client,ready).json()['code']=='preview_stale'
    payload={key:q[key] for key in ('items','rules','digest','moq','note')}
    payload.update(overrides={},confirmed_risk_ids=[],confirmed_all_text=True)
    assert client.post(BASE+'/customer-quotations/export',json=payload).json()['code']=='quotation_stale'


def test_same_receipt_concurrency_and_reordered_idempotent_replay(library):
    from concurrent.futures import ThreadPoolExecutor
    client,_=library
    data=upload(client,tables()).json();payload=plan(data)
    with ThreadPoolExecutor(max_workers=2) as pool:
        responses=list(pool.map(lambda _:confirm(client,data,payload),range(2)))
    assert [response.status_code for response in responses]==[200,200]
    assert responses[0].json()==responses[1].json()
    payload['decisions'].reverse()
    assert confirm(client,data,payload).json()==responses[0].json()
    assert len(detail(client)['prices'])==4


def test_cross_product_foreign_keys_and_product_delete_cascade(library):
    from sqlalchemy.exc import IntegrityError
    from app.models import LibraryProductPrice
    from app.product_library_service import now
    client,engine=library
    a=product(client,model='A',wholesale_price='12.00');b=product(client,model='B')
    with Session(engine) as db:
        db.add(LibraryProductPrice(product_id=b['id'],source_id=a['sources'][0]['id'],price_type='批发价',currency_code='CNY',price_unit='件',amount_text='12',tax_mode='unknown',freight_mode='unknown',is_default_wholesale=False,created_at=now(),updated_at=now(),version=1))
        with pytest.raises(IntegrityError):
            db.commit()
    assert client.request('DELETE',BASE+f"/products/{a['id']}",json={'version':a['version']}).status_code==204
    with engine.connect() as connection:
        assert connection.exec_driver_sql('PRAGMA foreign_key_check').all()==[]
        assert connection.exec_driver_sql('SELECT COUNT(*) FROM library_product_prices').scalar()==0


def test_generic_price_semantics_modes_and_no_quote(library):
    client,_=library
    t={'产品数据':[['A','GS008','示例1'],['B','DM82','示例2'],['C','S17','无报价'],['D','W01','多模式']],
       '价格明细':[['P1','A',None,'其他',None,'26',None,None,None,None,None,None,None,None,None,None,'false',None,'出货价：含义未确认'],
                 ['P2','A',None,'其他',None,'59',None,None,None,None,None,None,None,None,None,None,'false',None,'零售控价'],
                 ['P3','B',None,'平台控价','阿里','22.9','CNY'],['P4','B',None,'平台控价','C端','39.9','CNY']],
       '参数明细':[['D',None,None,'性能','功率','5','W','自然风','5W'],['D',None,None,'性能','功率','10','W','强风','10W']]}
    t['资料索引']=[['S1','A','资料公司',None,'示例.pdf','1',None,'quote','出货价：含义未确认'],['S2','A','资料公司',None,'示例.pdf','1',None,'quote','零售控价']]
    t['价格明细'][0][15]='S1'
    t['价格明细'][1][15]='S2'
    data=upload(client,t).json();response=confirm(client,data)
    assert response.status_code==200,response.text
    assert not any(r['errors'] for r in data['rows']), [(r['row_key'],r['errors']) for r in data['rows'] if r['errors']]
    assert detail(client,1)['prices'][0]['price_type']=='其他'
    assert detail(client,1)['prices'][0]['currency_code'] is None
    assert {price['price_channel'] for price in detail(client,2)['prices']}=={'阿里','C端'}
    assert detail(client,3)['prices']==[] and detail(client,3)['wholesale_price'] is None
    assert {item['condition'] for item in detail(client,4)['structured_attributes']}=={'自然风','强风'}


def test_legacy_default_unknown_conditions_require_confirmation(library):
    client,_=library
    p=product(client,wholesale_price='12.30')
    payload={'scope':'all'}
    q=client.post(BASE+'/customer-quotations/preview',json=payload).json()
    assert q['rows'][0]['final_price'] is None and q['rows'][0]['error']
    q=client.post(BASE+'/customer-quotations/preview',json={**payload,'conditions_confirmed_ids':[p['id']]}).json()
    assert q['rows'][0]['final_price']=='12.30'
    assert detail(client,p['id'])['wholesale_price']=='12.30'


def test_legacy_multi_price_edit_and_first_company_change_are_explicit(library):
    client,_=library
    p=product(client,model='X',name='产品',wholesale_price='9.00')
    ready=upload(client,{'产品数据':[['P','X','产品',None,None,None,'2.00']]}).json()
    assert confirm(client,ready).status_code==200
    p=detail(client)
    payload={field:p[field] for field in ('manufacturer_name','model','name','parameters','selling_points','wholesale_price','price_note','version')}
    assert client.put(BASE+'/products/1',json={**payload,'wholesale_price':'7.00'}).json()['code']=='default_price_selection_required'
    assert client.put(BASE+'/products/1',json={**payload,'manufacturer_name':'确认的新公司'}).json()['code']=='first_company_confirmation_required'
    assert client.put(BASE+'/products/1',json={**payload,'manufacturer_name':'确认的新公司','confirm_manufacturer_change':True}).status_code==200


def test_v1_media_warning_requires_confirmation_and_receipt_records_it(library):
    from test_product_library import workbook as v1_workbook
    client,_=library
    output=io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(v1_workbook([['M','产品']])) ) as original, zipfile.ZipFile(output,'w') as result:
        for entry in original.infolist():
            result.writestr(entry,original.read(entry.filename))
        result.writestr('xl/media/unused.png',b'not-imported-image')
    ready=client.post(BASE+'/imports/preview',params={'manufacturer_name':'公司'},content=output.getvalue()).json()
    assert ready['attachment_warning']
    assert confirm_v1(client,ready).json()['code']=='attachments_not_confirmed'
    payload={'decisions':[{'row_number':2,'action':'insert'}],'confirmed_attachments_not_imported':True}
    response=client.post(BASE+f"/imports/{ready['import_id']}/confirm",json=payload)
    assert response.status_code==200,response.text
    assert response.json()['attachments_not_imported'] is True


def test_unknown_company_stays_pending_and_workbook_name_is_not_original_file(library):
    client,_=library
    assert client.post(BASE+'/imports/preview',params={'manufacturer_name':'公司'},content=workbook(),headers={'X-Filename':'a'*501+'.xlsx'}).status_code==422
    data=upload(client,{'产品数据':[['P','X','资料']]},company='').json()
    parent=next(row for row in data['rows'] if row['kind']=='product')
    assert parent['allowed_actions']==['skip_product'] and parent['errors']
    assert confirm(client,data).json()['code']=='no_changes_selected'
    data=upload(client,{'产品数据':[['P','X','资料']]}).json()
    assert confirm(client,data).status_code==200
    source=detail(client)['sources'][0]
    assert source['source_file_name'] is None and 'products.xlsx' in source['notes']


def test_variant_parameters_must_be_removed_explicitly_before_deletion(library):
    client,_=library
    data=upload(client,{'产品数据':[['P','X','资料']], '规格款式':[['V','P','蓝']], '参数明细':[['P','V',None,'基本','尺寸','10','cm',None,'10cm']]}).json()
    assert confirm(client,data).status_code==200
    p=detail(client);variant=p['variants'][0]
    response=client.request('DELETE',BASE+f"/products/{p['id']}/variants/{variant['id']}",json={'product_version':p['version'],'version':variant['version']})
    assert response.json()['code']=='related_record_in_use'


def test_long_synthesized_source_can_be_edited_without_losing_excerpt(library):
    client,_=library
    ready=preview_v1(client,[['M','资料','参'*10000,'卖'*10000]]).json()
    assert confirm_v1(client,ready).status_code==200
    p=detail(client);source=p['sources'][0]
    assert len(source['source_excerpt'])>20000
    response=client.put(BASE+f"/products/{p['id']}/sources/{source['id']}",json={'product_version':p['version'],'version':source['version'],'notes':'补充核对记录'})
    assert response.status_code==200,response.text
    assert response.json()['source_excerpt']==source['source_excerpt']


def test_new_product_uses_single_explicit_source_when_batch_company_unknown(library):
    client,_=library
    data=upload(client,{'产品数据':[['P','X','资料']], '资料索引':[['S','P','明确来源公司',None,'原件.pdf','3',None,'quote','页码摘录']]},company='').json()
    parent=next(row for row in data['rows'] if row['kind']=='product')
    assert not parent['errors'],parent['errors']
    response=confirm(client,data)
    assert response.status_code==200,response.text
    p=detail(client)
    assert p['manufacturer_name']=='明确来源公司'
    assert {source['source_company'] for source in p['sources']}=={'明确来源公司'}


def test_mixed_workbook_resolves_first_company_per_product(library):
    client,_=library
    data=upload(client,{'产品数据':[['P','X','产品甲'],['Q','Y','产品乙']], '资料索引':[['S','P','公司甲'],['T','Q','公司乙']]},company='').json()
    response=confirm(client,data)
    assert response.status_code==200,response.text
    assert detail(client,1)['manufacturer_name']=='公司甲'
    assert detail(client,2)['manufacturer_name']=='公司乙'


def test_existing_product_retains_full_main_row_with_explicit_source(library):
    client,engine=library
    old=product(client,model='X',name='产品',parameters='旧参数',selling_points='旧卖点',price_note='旧备注',wholesale_price='9.00')
    incoming={'产品数据':[['P','X','产品',None,'完整新参数\n第二段','完整新卖点\n第二项',None,'新价格条件备注']],
              '资料索引':[['S','P','资料提供公司',None,'原始资料.pdf','第3页',None,'quote','简短页码摘录']]}
    data=upload(client,incoming,company='').json()
    assert next(row for row in data['rows'] if row['kind']=='product')['default_action']=='use_existing'
    response=confirm(client,data)
    assert response.status_code==200,response.text
    p=detail(client,old['id'])
    for field in ('manufacturer_name','parameters','selling_points','price_note','wholesale_price'):
        assert p[field]==old[field]
    full=next((source for source in p['sources'] if '完整新参数\n第二段' in (source['source_excerpt'] or '')),None)
    assert full is not None,'主行全文在预览清除后必须仍可查询'
    assert '完整新卖点\n第二项' in full['source_excerpt'] and '新价格条件备注' in full['source_excerpt']
    assert full['source_company']=='资料提供公司'
    assert full['source_file_name'] is None and full['source_page'] is None
    explicit=next(source for source in p['sources'] if source['source_file_name']=='原始资料.pdf')
    assert explicit['source_excerpt']=='简短页码摘录' and explicit['source_page']=='第3页'
    assert p['version']==old['version']+1
    with Session(engine) as db:
        assert db.get(LibraryImport,data['import_id']).preview is None
    repeated=upload(client,incoming,company='').json()
    assert confirm(client,repeated).json()['code']=='no_changes_selected'


@pytest.mark.parametrize('existing',[False,True])
def test_multiple_source_companies_require_explicit_main_row_attribution(library,existing):
    client,_=library
    if existing:
        product(client,model='X',name='产品',parameters='旧参数')
    incoming={'产品数据':[['P','X','产品',None,'主行原文']], '资料索引':[['S','P','公司甲'],['T','P','公司乙']]}
    for company in ('','不在资料中的公司'):
        data=upload(client,incoming,company=company).json()
        parent=next(row for row in data['rows'] if row['kind']=='product')
        assert parent['source_company_options']==['公司乙','公司甲']
        assert parent['allowed_actions']==['skip_product'] and parent['main_source_company'] is None
        assert confirm(client,data).json()['code']=='no_changes_selected'
    data=upload(client,incoming,company='公司乙').json()
    parent=next(row for row in data['rows'] if row['kind']=='product')
    assert parent['main_source_company']=='公司乙'
    assert confirm(client,data).status_code==200
    p=detail(client)
    assert p['manufacturer_name']==('厂家' if existing else '公司乙')
    assert {source['source_company'] for source in p['sources']} >= {'公司甲','公司乙'}
    raw=next(source for source in p['sources'] if '主行原文' in (source['source_excerpt'] or ''))
    assert raw['source_company']=='公司乙'


def test_main_row_provenance_cannot_be_silently_skipped(library):
    client,_=library
    incoming={'产品数据':[['P','X','产品',None,'必须保留的原文']], '资料索引':[['S','P','公司','明确厂家','原件.pdf','1',None,'quote','短摘录']]}
    data=upload(client,incoming,company='').json()
    raw=next(row for row in data['rows'] if row.get('main_row_provenance'))
    payload=plan(data,{(raw['row_key']['sheet'],raw['row_key']['temp_id']):{'action':'skip'}})
    assert confirm(client,data,payload).json()['code']=='invalid_decisions'
    assert client.get(BASE+'/products').json()['total']==0
    payload=plan(data,{('产品数据','P'):{'action':'skip_product'},(raw['row_key']['sheet'],raw['row_key']['temp_id']):{'action':'skip'}})
    assert confirm(client,data,payload).json()['code']=='no_changes_selected'
    assert confirm(client,data).status_code==200


def test_old_v2_ready_requires_repreview_but_committed_receipt_still_replays(library):
    client,engine=library
    ready=upload(client,{'产品数据':[['P','X','待导入']]}).json()
    saved=upload(client,{'产品数据':[['P','Y','已导入']]}).json()
    payload=plan(saved)
    result=confirm(client,saved,payload).json()
    with Session(engine) as db:
        item=db.get(LibraryImport,ready['import_id'])
        preview=copy.deepcopy(item.preview);preview.pop('import_plan_version');item.preview=preview
        item=db.get(LibraryImport,saved['import_id'])
        old=copy.deepcopy(item.result);old.pop('import_plan_version');item.result=old
        db.commit()
    assert client.get(BASE+f"/imports/{ready['import_id']}").json()['code']=='preview_rule_changed'
    assert confirm(client,ready).json()['code']=='preview_rule_changed'
    result.pop('import_plan_version')
    assert client.get(BASE+f"/imports/{saved['import_id']}").json()==result
    assert confirm(client,saved,payload).json()==result


def test_main_row_source_failure_rolls_back_explicit_source_and_receipt(library):
    from sqlalchemy.exc import OperationalError
    client,engine=library
    old=product(client,model='X',name='产品',parameters='旧参数')
    incoming={'产品数据':[['P','X','产品',None,'主行事务哨兵']], '资料索引':[['S','P','明确公司',None,None,None,None,'quote','短摘录']]}
    data=upload(client,incoming,company='').json()
    def fail(connection,cursor,statement,parameters,context,executemany):
        if statement.startswith('INSERT INTO library_product_sources') and '主行事务哨兵' in str(parameters):
            raise OperationalError(statement,parameters,Exception('injected main-row source failure'))
    event.listen(engine,'before_cursor_execute',fail)
    try:
        assert confirm(client,data).json()['code']=='import_save_failed'
    finally:
        event.remove(engine,'before_cursor_execute',fail)
    assert detail(client,old['id'])==old
    with Session(engine) as db:
        assert db.get(LibraryImport,data['import_id']).status=='ready'
    assert confirm(client,data).status_code==200
    assert any('主行事务哨兵' in (source['source_excerpt'] or '') for source in detail(client)['sources'])


def test_explicit_full_excerpt_reused_without_duplicate_provenance(library):
    client,_=library
    excerpt='产品临时ID: P\n产品型号: X\n产品名称: 产品\n产品参数原文: 新原文'
    data=upload(client,{'产品数据':[['P','X','产品',None,'新原文']], '资料索引':[['S','P','明确公司',None,'原件.pdf','3',None,'quote',excerpt]]},company='').json()
    assert sum(row['kind']=='sources' for row in data['rows'])==1
    assert confirm(client,data).status_code==200
    assert detail(client)['sources'][0]['source_excerpt']==excerpt
