"""Five-sheet import plans, validated again under the confirmation write lock."""
import io
import json
import zipfile
from collections import defaultdict
from datetime import timedelta
from decimal import Decimal
from uuid import uuid4
from xml.etree import ElementTree as ET

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
from sqlalchemy import select, text
from sqlalchemy.exc import SQLAlchemyError

from app.models import LibraryImport, LibraryProduct
from app.product_library_service import (FIELDS as MAIN_FIELDS, HEADERS as V1_HEADERS, baseline, begin_write, database_error,
    digest, error, normalized_text, now, price, product_view, read_excel, receipt, require_identity_schema)
from app.product_library_details import (ATTACHMENT_WARNING, FIELDS, MODELS, PRICE_KEY, SOURCE_KEY, attributes,
    decimal_text, detail_view, product_details, project_price, records, text_value, touch, validate_detail)

SHEETS = {
    '产品数据': ('产品临时ID', '产品型号', '产品名称', '产品分类', '产品参数原文', '产品卖点', '默认批发价', '价格备注'),
    '参数明细': ('产品临时ID', '规格临时ID', '资料临时ID', '参数分组', '参数名称', '参数值', '单位', '适用条件', '原文摘录'),
    '规格款式': ('规格临时ID', '产品临时ID', '规格名称', '规格编码', '颜色/款式', '规格备注'),
    '价格明细': ('价格临时ID', '产品临时ID', '规格临时ID', '价格类型', '价格渠道', '金额', '币种', '计价单位', '起订量', '起订单位', '税费口径', '运费口径', '运费金额', '定制条件', '报价日期', '资料临时ID', '是否默认批发价', '价格备注', '原文摘录'),
    '资料索引': ('资料临时ID', '产品临时ID', '报价来源公司', '原件中标注的生产厂家', '原始文件名', '页码/工作表/区域', '原件日期', '资料类型', '原文摘录', '备注'),
}
SHEET_KIND = dict(zip(SHEETS, ('product', 'attribute', 'variants', 'prices', 'sources')))
PRODUCT_ACTIONS = ('insert', 'update_main', 'use_existing', 'append_source', 'skip_product')


def template_bytes():
    wb = Workbook()
    wb.remove(wb.active)
    for name, headers in SHEETS.items():
        ws = wb.create_sheet(name)
        ws.append(headers)
        ws.freeze_panes = 'A2'
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = 24
            from openpyxl.styles import Font
            column[0].font = Font(bold=True)
    notes = wb.create_sheet('填写说明')
    for line in [
        '五张业务表及表头必须完整保留；所有ID和金额以文本填写。主产品最多1000行，关联表合计5000行，5 MiB。',
        '产品身份为型号+名称（不含来源公司），同一产品只写一个主行。临时ID只在本工作簿使用。',
        '资料类型：quote/catalog/specification/other；税费：unknown/included/excluded；运费另可 extra/negotiable。',
        '价格类型：出厂价/批发价/一件代发价/定制价/平台控价/建议零售价/零售价/其他；其他必须保留原文。',
        '币种为大写ISO 4217代码；未知值留空，不猜测为0或CNY。计价单位和起订单位分开，日期YYYY-MM-DD。',
        '是否默认批发价：true/false；仅人民币/件、金额明确的批发价可设默认。已有产品须在预览显式确认。',
        '更新时空值保留旧值；清空或删除请在详情操作。参数冲突保留所有事实并标为待核对。',
        ATTACHMENT_WARNING,
    ]:
        notes.append([line])
    notes.column_dimensions['A'].width = 110
    output = io.BytesIO()
    wb.save(output)
    wb.close()
    return output.getvalue()


def physical_rows(wb, content, sheet):
    ws = wb[sheet]
    if ws.sheet_state != 'visible':
        error(400, 'invalid_structure', f'{sheet}不能隐藏')
    last = 1
    with zipfile.ZipFile(io.BytesIO(content)) as archive, archive.open(ws._worksheet_path) as source:
        for _, element in ET.iterparse(source, events=('end',)):
            tag = element.tag.rsplit('}', 1)[-1]
            if tag == 'mergeCell' or tag in ('row', 'col') and element.attrib.get('hidden') in ('1', 'true', 'True'):
                error(400, 'invalid_structure', f'{sheet}不能合并或隐藏行列')
            if tag == 'c':
                nonempty = element.attrib.get('t') == 'e' or any(child.tag.rsplit('}', 1)[-1] == 'f' or child.tag.rsplit('}', 1)[-1] in ('v', 't') and child.text is not None for child in element.iter())
                if nonempty:
                    column, number = coordinate_from_string(element.attrib['r'])
                    if column_index_from_string(column) > len(SHEETS[sheet]):
                        error(400, 'invalid_structure', f'{sheet}有未知列')
                    last = max(last, number)
                element.clear()
            elif tag == 'row':
                element.clear()
    limit = 1000 if sheet == '产品数据' else 5000
    if last - 1 > limit:
        error(413, 'workbook_limit_exceeded', f'{sheet}物理行超限')
    ws.reset_dimensions()
    rows = ws.iter_rows(max_row=last, max_col=len(SHEETS[sheet]))
    if tuple(cell.value for cell in next(rows)) != SHEETS[sheet]:
        error(400, 'invalid_structure', f'{sheet}表头必须严格匹配V2模板')
    result = []
    for number, cells in enumerate(rows, 2):
        if all(cell.value is None for cell in cells):
            continue
        if any(cell.data_type in ('f', 'e') for cell in cells):
            error(400, 'invalid_structure', f'{sheet}第{number}行含公式或错误单元格')
        result.append((number, dict(zip(SHEETS[sheet], (cell.value for cell in cells)))))
    return result, last - 1


def row_id(sheet, number, temp_id):
    return {'sheet': sheet, 'row_number': number, 'record_type': SHEET_KIND[sheet], 'temp_id': temp_id}


def key(row):
    return json.dumps(row['row_key'], ensure_ascii=False, sort_keys=True)


def parse(wb, content):
    if '产品数据' not in wb.sheetnames:
        error(400, 'invalid_structure', '缺少产品数据')
    first = tuple(next(wb['产品数据'].iter_rows(max_row=1, max_col=8, values_only=True)))
    if first[:6] == V1_HEADERS and not any(first[6:]):
        if set(wb.sheetnames) & (set(SHEETS) - {'产品数据'}):
            error(400, 'invalid_structure', '混合版本，请使用完整V2模板')
        return None
    if set(wb.sheetnames) - {'填写说明'} != set(SHEETS):
        error(400, 'invalid_structure', 'V2必须具有且仅具有五张业务表')
    rows, tables, total = [], {}, 0
    for sheet in SHEETS:
        table, count = physical_rows(wb, content, sheet)
        tables[sheet] = table
        if sheet != '产品数据':
            total += count
    if total > 5000:
        error(413, 'workbook_limit_exceeded', '关联表合计最多5000个物理行')
    products, variants, sources = {}, {}, {}
    ids = defaultdict(set)
    for sheet, table in tables.items():
        kind = SHEET_KIND[sheet]
        for number, raw in table:
            id_field = {'product': '产品临时ID', 'variants': '规格临时ID', 'sources': '资料临时ID', 'prices': '价格临时ID'}.get(kind)
            try:
                temp_id = text_value(raw[id_field], id_field, 100, True) if id_field else None
                product_ref = text_value(raw['产品临时ID'], '产品临时ID', 100, True)
                variant_ref = text_value(raw.get('规格临时ID'), '规格临时ID', 100) if kind in ('attribute', 'prices') else None
                source_ref = text_value(raw.get('资料临时ID'), '资料临时ID', 100) if kind in ('attribute', 'prices') else None
            except ValueError as exc:
                error(400, 'invalid_structure', f'{sheet}第{number}行：{exc}')
            if temp_id is not None:
                if temp_id in ids[kind]:
                    error(400, 'invalid_structure', f'{sheet}第{number}行临时ID重复')
                ids[kind].add(temp_id)
            row = dict(row_key=row_id(sheet, number, temp_id), kind=kind, product_ref=product_ref, variant_ref=variant_ref, source_ref=source_ref, raw=raw, errors=[], dependencies=[])
            rows.append(row)
            if kind == 'product':
                products[temp_id] = row
            elif kind == 'variants':
                variants[temp_id] = row
            elif kind == 'sources':
                sources[temp_id] = row
    if not products:
        error(400, 'no_product_rows', '没有产品行')
    for row in rows:
        for ref, table in ((row['product_ref'], products), (row['variant_ref'], variants), (row['source_ref'], sources)):
            if ref is not None and (ref not in table or table[ref]['product_ref'] != row['product_ref']):
                error(400, 'invalid_structure', '引用不存在或跨产品', **row['row_key'])
    return rows


def main_excerpt(raw):
    return '\n'.join(f'{name}: {value}' for name, value in raw.items() if value is not None)


def normalize_rows(rows, company, db, workbook_name=None):
    for row in rows:
        raw, kind = row['raw'], row['kind']
        try:
            if kind == 'product':
                data = {field: normalized_text(raw.get(header), field) for header, field in (('产品型号','model'),('产品名称','name'),('产品参数原文','parameters'),('产品卖点','selling_points'),('价格备注','price_note'))}
                data['category'] = text_value(raw['产品分类'], '产品分类')
                data['wholesale_price'] = price(raw['默认批发价'])
            elif kind == 'sources':
                data = dict(zip(FIELDS[kind], (raw.get(header) for header in ('报价来源公司','原件中标注的生产厂家','原始文件名','页码/工作表/区域','原件日期','资料类型','原文摘录','备注'))))
                data['source_company'] = text_value(data['source_company'], 'source_company') or company
                data = validate_detail(kind, data, db, None, synthetic=row.get('synthetic', False))
            elif kind == 'variants':
                data = dict(variant_name=raw['规格名称'], variant_code=raw['规格编码'], color=raw['颜色/款式'], notes=raw['规格备注'])
                data = validate_detail(kind, data, db, None)
            elif kind == 'prices':
                data = dict(zip(('price_type','price_channel','amount_text','currency_code','price_unit','min_order_quantity','min_order_unit','tax_mode','freight_mode','customization_terms','quote_date','terms_text','source_excerpt'), (raw.get(header) for header in ('价格类型','价格渠道','金额','币种','计价单位','起订量','起订单位','税费口径','运费口径','定制条件','报价日期','价格备注','原文摘录'))))
                flag = text_value(raw.get('是否默认批发价'), '是否默认批发价')
                if flag not in (None, 'true', 'false'):
                    raise ValueError('默认标记仅允许文本true/false')
                data['is_default_wholesale'] = flag == 'true'
                freight = price(raw.get('运费金额'))
                data['freight_amount_cents'] = None if freight is None else int(Decimal(freight)*100)
                data = validate_detail(kind, data, db, None, defer_references=True)
            else:
                data = {field: text_value(raw.get(header), field, 10000 if field in ('value','source_excerpt') else 200, field == 'name') for header, field in (('参数分组','group'),('参数名称','name'),('参数值','value'),('单位','unit'),('适用条件','condition'),('原文摘录','source_excerpt'))}
                if not data['value'] and not data['source_excerpt']:
                    raise ValueError('参数缺少值和摘录')
                data.update(source_id=None, review_status='confirmed' if data['value'] else 'needs_review')
            row['incoming'] = data
        except (ValueError, KeyError) as exc:
            row['incoming'] = {}
            row['errors'].append({**row['row_key'], 'field': 'record', 'message': str(exc)})
    sources = {row['row_key']['temp_id']: row for row in rows if row['kind'] == 'sources'}
    for product in [row for row in rows if row['kind'] == 'product']:
        ref = product['product_ref']
        group_sources = [row for row in sources.values() if row['product_ref'] == ref]
        if not group_sources:
            temp_id = f'@source:{ref}'
            while temp_id in sources:
                temp_id += ':'
            synthetic = dict(row_key=row_id('资料索引', product['row_key']['row_number'], temp_id), kind='sources', product_ref=ref, variant_ref=None, source_ref=None, raw={}, synthetic=True, errors=[], dependencies=[], incoming=dict(source_company=company, source_type='other', source_file_name=None, source_page=None, source_date=None, source_excerpt=main_excerpt(product['raw']), notes=f'整理工作簿：{workbook_name or "未提供"}；原件文件名未知', claimed_manufacturer_name=None))
            rows.append(synthetic)
            sources[temp_id] = synthetic
            group_sources = [synthetic]
        for row in [r for r in rows if r['product_ref'] == ref and r['kind'] in ('attribute','prices')]:
            if row['source_ref'] is None:
                if len(group_sources) != 1:
                    row['errors'].append({**row['row_key'], 'field':'资料临时ID', 'message':'多个来源时必须明确指定来源'})
                else:
                    row['source_ref'] = group_sources[0]['row_key']['temp_id']
            source = sources.get(row['source_ref'])
            if row['kind'] == 'prices' and source and not row['errors'] and not source['errors']:
                date = source['incoming'].get('source_date')
                if date and row['incoming'].get('quote_date') not in (None, date):
                    row['errors'].append({**row['row_key'], 'field':'报价日期', 'message':'报价日期与来源日期冲突'})
                else:
                    row['incoming']['quote_date'] = row['incoming'].get('quote_date') or date
        amount = product['incoming'].get('wholesale_price')
        if amount is not None:
            defaults = [r for r in rows if r['product_ref'] == ref and r['kind'] == 'prices' and r['incoming'].get('is_default_wholesale')]
            if defaults:
                if len(defaults) != 1 or defaults[0]['variant_ref'] is not None or defaults[0]['incoming'].get('amount_text') != decimal_text(amount, currency='CNY') or len(group_sources) != 1 or defaults[0]['source_ref'] != group_sources[0]['row_key']['temp_id']:
                    product['errors'].append({**product['row_key'], 'field':'默认批发价', 'message':'快捷默认价与价格明细冲突'})
            elif len(group_sources) == 1:
                data = dict(price_type='批发价', amount_text=decimal_text(amount, currency='CNY'), currency_code='CNY', price_unit='件', is_default_wholesale=True)
                data = validate_detail('prices', data, db, None, defer_references=True)
                rows.append(dict(row_key=row_id('价格明细', product['row_key']['row_number'], derived_id(rows, 'prices', f'@price:{ref}')), kind='prices', product_ref=ref, variant_ref=None, source_ref=group_sources[0]['row_key']['temp_id'], raw={}, synthetic=True, errors=[], dependencies=[], incoming=data))
            else:
                product['errors'].append({**product['row_key'], 'field':'默认批发价', 'message':'快捷默认价需要唯一来源'})
    for row in rows:
        row['raw'] = {field:value if value is None or isinstance(value,(str,int,float,bool)) else str(value) for field,value in row['raw'].items()}
    return rows


def snapshot(db, rows):
    identities = {(r['incoming'].get('model'), r['incoming'].get('name')) for r in rows if r['kind'] == 'product' and not r['errors']}
    products = db.scalars(select(LibraryProduct).order_by(LibraryProduct.id)).all()
    return digest([{**product_view(p), **product_details(db,p)} for p in products if (p.model,p.name) in identities])


def match_key(kind, data):
    fields = SOURCE_KEY if kind == 'sources' else PRICE_KEY if kind == 'prices' else ('variant_code',) if data.get('variant_code') else ('variant_name','color')
    return tuple(data.get(field) for field in fields)


def plan(db, rows, company=None):
    products = { (p.model,p.name): p for p in db.scalars(select(LibraryProduct)) }
    product_rows = {r['product_ref']: r for r in rows if r['kind'] == 'product'}
    refs = {(r['kind'],r['row_key']['temp_id']): r for r in rows if r['row_key']['temp_id'] is not None}
    identities = defaultdict(list)
    for row in product_rows.values():
        if not row['errors']:
            identities[(row['incoming']['model'],row['incoming']['name'])].append(row)
    for group in identities.values():
        if len(group) > 1:
            for row in group:
                row['errors'].append({**row['row_key'], 'field':'产品名称', 'message':'同一身份有多条主产品行，请合并后重新预览'})
    for row in sorted(rows, key=lambda r: ('product','sources','variants','attribute','prices').index(r['kind'])):
        kind, data = row['kind'], row['incoming']
        parent = product_rows[row['product_ref']]
        product = products.get((parent['incoming'].get('model'),parent['incoming'].get('name')))
        row.update(old=None, effective=data, candidate=[], target_id=None, target_version=None)
        if kind == 'product':
            if product is None and not company:
                row['errors'].append({**row['row_key'],'field':'source_company','message':'新产品的首录资料公司未知，请保留待处理并核实；不得编造公司名'})
            if product:
                row.update(target_id=product.id,target_version=product.version,old=product_view(product), effective=product_view(product), allowed_actions=['update_main','use_existing','append_source','skip_product'],default_action='use_existing')
            else:
                row.update(allowed_actions=['insert','skip_product'],default_action='insert')
        else:
            row['dependencies'] = [parent['row_key']]
            for ref_kind, ref in (('sources',row['source_ref']),('variants',row['variant_ref'])):
                if ref:
                    dep = refs[(ref_kind,ref)]
                    row['dependencies'].append(dep['row_key'])
            incoming = dict(data)
            if kind in ('prices','attribute'):
                for field, ref_kind, ref in (('source_id','sources',row['source_ref']),('variant_id','variants',row['variant_ref'])):
                    if ref:
                        incoming[field] = refs[(ref_kind,ref)].get('target_id') or f'new:{ref}'
            candidates = records(db,product.id,kind) if product and kind in MODELS else []
            matches = [record for record in candidates if match_key(kind,detail_view(record,kind)) == match_key(kind,incoming)]
            if kind == 'attribute':
                container = product
                if row['variant_ref']:
                    dep = refs[('variants',row['variant_ref'])]
                    container = db.get(MODELS['variants'],dep['target_id']) if dep.get('target_id') else None
                existing = container.structured_attributes or [] if container else []
                exact = any(all(item.get(f) == incoming.get(f) for f in ('group','name','value','unit','condition','source_id','source_excerpt')) for item in existing)
                row.update(allowed_actions=['reuse','skip'] if exact else ['insert','skip'],default_action='reuse' if exact else 'insert',target_id=container.id if container else None,target_version=container.version if container else None)
            elif matches:
                record = matches[0]
                old = detail_view(record,kind)
                effective = {f: data.get(f) if data.get(f) is not None else old[f] for f in FIELDS[kind]}
                if kind == 'prices':
                    effective['is_default_wholesale'] = old['is_default_wholesale']
                changed = any(effective.get(f) != old[f] for f in FIELDS[kind] if f not in ('source_id','variant_id'))
                row.update(target_id=record.id,target_version=record.version,old=old,effective=effective,allowed_actions=['update','skip'] if changed else ['reuse','skip'],default_action='skip' if changed else 'reuse')
            else:
                row.update(allowed_actions=['insert','skip'],default_action='insert')
        if row['errors'] or parent['errors']:
            row.update(allowed_actions=['skip_product'] if kind == 'product' else ['skip'], default_action='skip_product' if kind == 'product' else 'skip')
        row['category'] = 'invalid' if row['errors'] else 'duplicate' if row['target_id'] else 'new'
    attribute_facts=set()
    for row in rows:
        if row['kind']=='attribute' and not row['errors']:
            fact=(row['product_ref'],row['variant_ref'],row['source_ref'],digest(row['incoming']))
            if fact in attribute_facts:
                row.update(category='duplicate',allowed_actions=['skip'],default_action='skip')
                row['errors'].append({**row['row_key'],'field':'record','message':'同一参数事实重复行'})
            attribute_facts.add(fact)
    # Detect duplicate association keys before allowing any sequential writes.
    groups = defaultdict(list)
    for row in rows:
        if row['kind'] in MODELS and not row['errors']:
            data = {**row['incoming'], 'variant_id':row['variant_ref'], 'source_id':row['source_ref']}
            groups[(row['product_ref'], row['kind'], match_key(row['kind'],data))].append(row)
    for group in groups.values():
        if len(group) > 1:
            same = all(row['incoming'] == group[0]['incoming'] for row in group)
            for row in (group[1:] if same else group):
                row.update(category='duplicate' if same else 'invalid',allowed_actions=['skip'],default_action='skip')
                row['errors'].append({**row['row_key'],'field':'record','message':'同键重复行' if same else '同键内容冲突，请修正或全部跳过'})
    return rows


def preview(content, company, db, workbook_name=None):
    parsed = read_excel(content, parser=parse)
    if parsed is None:
        return None
    try:
        company = text_value(company,'source_company')
        workbook_name = text_value(workbook_name,'整理工作簿名',500)
        db.execute(text('BEGIN'))
        require_identity_schema(db)
        rows = plan(db,normalize_rows(parsed,company,db,workbook_name),company)
        stamp = snapshot(db,rows)
        db.rollback()
        created = now()
        result = dict(import_id=str(uuid4()),status='ready',template_version=2,identity_rule_version=4,expires_at=(created+timedelta(minutes=30)).isoformat(),manufacturer_name=company or '',rows=rows,attachment_warning=ATTACHMENT_WARNING,counts={kind:sum(r['kind']==kind for r in rows) for kind in ('product','variants','sources','prices','attribute')})
        begin_write(db)
        db.add(LibraryImport(import_id=result['import_id'],manufacturer_name=company or '',created_at=created,expires_at=created+timedelta(minutes=30),status='ready',preview=result,baseline_digest=stamp))
        db.commit()
        return result
    except ValueError as exc:
        db.rollback()
        error(422,'invalid_product',str(exc))
    except HTTPException:
        db.rollback()
        raise
    except SQLAlchemyError as exc:
        database_error(db,exc,'import_save_failed')


def decisions(payload, rows):
    if set(payload) - {'decisions','confirmed_attachments_not_imported'} or not isinstance(payload.get('decisions'),list):
        error(422,'invalid_decisions','必须提供完整逐行决策')
    if payload.get('confirmed_attachments_not_imported') is not True:
        error(422,'attachments_not_confirmed',ATTACHMENT_WARNING)
    allowed = {key(row):row for row in rows}
    result = {}
    for choice in payload['decisions']:
        if not isinstance(choice,dict) or set(choice)-{'row_key','action','target_id','target_version','set_default'} or not {'row_key','action'}.issubset(choice):
            error(422,'invalid_decisions','决策字段无效')
        token = json.dumps(choice['row_key'],ensure_ascii=False,sort_keys=True)
        row = allowed.get(token)
        if row is None or token in result or choice['action'] not in row['allowed_actions']:
            error(422,'invalid_decisions','行或动作无效/重复')
        target_action = choice['action'] in ('update_main','use_existing','append_source','update','reuse')
        if target_action and (choice.get('target_id') != row['target_id'] or choice.get('target_version') != row['target_version'] or type(choice.get('target_id')) is not int or type(choice.get('target_version')) is not int):
            error(422,'invalid_decisions','必须选择预览目标ID和版本')
        if not target_action and ('target_id' in choice or 'target_version' in choice):
            error(422,'invalid_decisions','该动作不接受目标')
        if 'set_default' in choice and (row['kind']!='prices' or type(choice['set_default']) is not bool or choice['action']=='skip'):
            error(422,'invalid_decisions','默认命令无效')
        result[token] = choice
    if set(result) != set(allowed):
        error(422,'invalid_decisions','每行必须恰有一个决策')
    return result


def confirm(import_id,payload,db):
    try:
        begin_write(db)
        item = receipt(db,import_id)
        if item.status == 'committed':
            decision_hash = decision_digest(payload)
            if item.decision_digest != decision_hash:
                error(409,'idempotency_conflict','已按不同决策提交')
            result = item.result
            db.rollback()
            return result
        require_identity_schema(db)
        rows = item.preview['rows']
        selected = decisions(payload,rows)
        decision_hash = decision_digest(payload)
        if snapshot(db,rows) != item.baseline_digest:
            error(409,'preview_stale','产品及关联资料已变化，请重新预览')
        parents = {row['product_ref']:row for row in rows if row['kind']=='product'}
        active, resolved, touched, outcomes = {}, {}, {}, []
        parameter_containers = {}
        ordered = sorted(rows,key=lambda row:(('product','sources','variants','attribute','prices').index(row['kind']), 0 if selected[key(row)].get('set_default') is False else 1))
        # Validate dependency decisions before changing any records.
        for row in ordered:
            action = selected[key(row)]['action']
            parent_action = selected[key(parents[row['product_ref']])]['action']
            if action in ('skip','skip_product') or parent_action == 'skip_product':
                continue
            if parent_action=='append_source' and row['kind'] in ('variants','attribute') and action not in ('reuse','skip'):
                error(422,'invalid_decisions','追加来源不能创建规格或修改参数')
            for dep in row['dependencies'][1:]:
                dep_key = json.dumps(dep,ensure_ascii=False,sort_keys=True)
                if selected[dep_key]['action'] == 'skip':
                    error(422,'invalid_decisions','被跳过的来源/规格不可被引用；请选择reuse或跳过依赖行')
        for row in ordered:
            kind, data, choice = row['kind'], dict(row['incoming']), selected[key(row)]
            action = choice['action']
            parent = parents[row['product_ref']]
            if action in ('skip','skip_product') or selected[key(parent)]['action']=='skip_product':
                outcomes.append(dict(row_key=row['row_key'],outcome='skipped',target_id=None,reason_code='user_skipped',errors=row['errors']))
                continue
            changed = False
            if kind=='product':
                if action=='insert':
                    data.pop('wholesale_price',None)
                    product = LibraryProduct(**data,manufacturer_name=item.manufacturer_name,created_at=now(),updated_at=now(),version=1)
                    db.add(product)
                    db.flush()
                    changed = True
                else:
                    product = db.get(LibraryProduct,choice['target_id'])
                    if product is None or product.version!=choice['target_version']:
                        error(409,'preview_stale','产品已变化')
                    if action=='update_main':
                        for field,value in data.items():
                            if field!='wholesale_price' and value is not None and getattr(product,field)!=value:
                                setattr(product,field,value)
                                changed = True
                active[row['product_ref']] = product
                record = product
            else:
                product = active[row['product_ref']]
                if kind in ('prices','attribute'):
                    for field,ref_kind,ref in (('source_id','sources',row['source_ref']),('variant_id','variants',row['variant_ref'])):
                        if ref:
                            data[field] = resolved[(ref_kind,ref)].id
                if kind=='attribute':
                    data.pop('variant_id',None)
                    container = resolved[('variants',row['variant_ref'])] if row['variant_ref'] else product
                    old = list(container.structured_attributes or [])
                    fact_fields = ('group','name','value','unit','condition','source_id','source_excerpt')
                    if action=='insert' and not any(all(fact.get(f)==data.get(f) for f in fact_fields) for fact in old):
                        semantic = ('group','name','unit','condition','source_id')
                        for fact in old:
                            if all(fact.get(f)==data.get(f) for f in semantic) and fact.get('value')!=data.get('value'):
                                fact['review_status']='needs_review'
                                data['review_status']='needs_review'
                        old.append(data)
                        container.structured_attributes=attributes(old,db,product.id)
                        if container is not product:
                            parameter_containers[container.id] = container
                        changed=True
                    record=container
                else:
                    record = db.get(MODELS[kind],choice['target_id']) if action in ('update','reuse') else None
                    if record is not None and (record.product_id!=product.id or record.version!=choice['target_version']):
                        error(409,'preview_stale','关联资料已变化')
                    if record:
                        data = {field:value if value is not None else getattr(record,field) for field,value in data.items()}
                    if kind=='prices':
                        data['is_default_wholesale'] = choice.get('set_default', data.get('is_default_wholesale',False) if selected[key(parent)]['action']=='insert' else record.is_default_wholesale if record else False)
                    data=validate_detail(kind,data,db,product.id,synthetic=row.get('synthetic',False))
                    if record is None:
                        record=MODELS[kind](product_id=product.id,**data,created_at=now(),updated_at=now(),version=1)
                        db.add(record)
                        changed=True
                    elif action=='update' or 'set_default' in choice:
                        if any(getattr(record,f)!=v for f,v in data.items()):
                            for field,value in data.items():
                                setattr(record,field,value)
                            touch(record)
                            changed=True
                    db.flush()
                    resolved[(kind,row['row_key']['temp_id'])]=record
            if changed:
                touched[product.id]=product
            outcomes.append(dict(row_key=row['row_key'],outcome='inserted' if action=='insert' else 'updated' if changed else 'reused',target_id=record.id,product_id=product.id,reason_code=action,errors=row['errors']))
        if not touched:
            error(422,'no_changes_selected','没有选择新增或更新')
        for product in touched.values():
            db.flush()
            project_price(db,product)
            if not any(row['kind']=='product' and selected[key(row)]['action']=='insert' and active.get(row['product_ref']) is product for row in rows):
                touch(product)
        for container in parameter_containers.values():
            touch(container)
        result=dict(import_id=import_id,status='committed',template_version=2,identity_rule_version=4,attachments_not_imported=True,rows=outcomes,counts={})
        for kind in ('product','sources','variants','attribute','prices'):
            result['counts'][kind]={outcome:sum(r['row_key']['record_type']==kind and r['outcome']==outcome for r in outcomes) for outcome in ('inserted','updated','reused','skipped')}
        item.status='committed'
        item.decision_digest=decision_hash
        item.result=result
        item.preview=None
        item.baseline_digest=None
        db.flush()
        db.commit()
        return result
    except ValueError as exc:
        db.rollback()
        error(422,'invalid_product',str(exc))
    except HTTPException:
        db.rollback()
        raise
    except SQLAlchemyError as exc:
        database_error(db,exc,'import_save_failed')


def decision_digest(payload):
    if not isinstance(payload, dict) or not isinstance(payload.get('decisions'),list):
        error(422,'invalid_decisions','决策格式无效')
    return digest({**payload,'decisions':sorted(payload['decisions'],key=lambda value:json.dumps(value,ensure_ascii=False,sort_keys=True))})


def derived_id(rows, kind, candidate):
    used={row['row_key']['temp_id'] for row in rows if row['kind']==kind}
    while candidate in used:
        candidate+=':'
    return candidate
