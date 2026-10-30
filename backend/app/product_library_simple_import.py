"""Single-sheet product import and legacy six-column compatibility."""
import hashlib
import io
import json
import re
import zipfile
from datetime import timedelta
from uuid import uuid4
from xml.etree import ElementTree as ET

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string
from sqlalchemy import func, select

from app.models import LibraryImport, LibraryProduct, LibrarySimpleQuoteEvidence, LibrarySimpleQuoteKey, LibrarySimpleQuoteVersion, LibrarySimpleSource
from app.product_library_service import begin_write, digest, error, now, read_excel, require_identity_schema, utc
from app.product_library_simple_prices import has_non_wholesale_semantics, parse_wholesale_prices

HEADERS = ('分类', '产品型号', '产品名称', '产品参数', '批发价', '控价', '卖点', '备注')
FIELDS = ('category', 'model', 'name', 'parameters', 'price_text', 'control_price', 'selling_points', 'notes')
LIMITS = (200, 200, 200, 10000, 10000, 2000, 10000, 10000)


def template_bytes():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '产品清单'
    sheet.append(HEADERS)
    sheet.freeze_panes = 'A2'
    for column in range(1, 9):
        sheet.column_dimensions[chr(64 + column)].width = 24
        sheet.cell(1, column).font = Font(bold=True)
    for row in range(2, 1002):
        sheet.cell(row, 2).number_format = '@'
    result = io.BytesIO()
    workbook.save(result)
    workbook.close()
    return result.getvalue()


def parse(content):
    def parse_new(workbook, original):
        if workbook.sheetnames != ['产品清单']:
            if '产品清单' in workbook.sheetnames:
                error(400, 'invalid_structure', '八列模板必须仅包含产品清单工作表')
            if '产品数据' in workbook.sheetnames:
                return None
            error(400, 'invalid_structure', '八列模板必须仅包含产品清单工作表')
        sheet = workbook['产品清单']
        if sheet.sheet_state != 'visible':
            error(400, 'invalid_structure', '产品清单不能隐藏')
        last = 1
        with zipfile.ZipFile(io.BytesIO(original)) as archive, archive.open(sheet._worksheet_path) as source:
            for _, element in ET.iterparse(source, events=('end',)):
                tag = element.tag.rsplit('}', 1)[-1]
                if tag == 'mergeCell' or tag in ('row', 'col') and element.attrib.get('hidden') in ('1', 'true', 'True'):
                    error(400, 'invalid_structure', '不能合并单元格或隐藏行列')
                if tag == 'c':
                    nonempty = element.attrib.get('t') == 'e' or any(child.tag.rsplit('}', 1)[-1] == 'f' or child.tag.rsplit('}', 1)[-1] in ('v', 't') and child.text is not None for child in element.iter())
                    if nonempty:
                        column, number = coordinate_from_string(element.attrib['r'])
                        if column_index_from_string(column) > 8:
                            error(400, 'invalid_structure', '八列以外存在非空内容')
                        if number > 1001:
                            error(413, 'workbook_limit_exceeded', '最多1000个物理数据行')
                        last = max(last, number)
                    element.clear()
                elif tag == 'row':
                    element.clear()
        sheet.reset_dimensions()
        cells = sheet.iter_rows(max_row=last, max_col=8)
        if tuple(cell.value for cell in next(cells)) != HEADERS:
            error(400, 'invalid_structure', '表头必须严格匹配八列模板')
        rows = []
        for number, row in enumerate(cells, 2):
            if all(cell.value is None for cell in row):
                continue
            values, errors, raw = {}, [], {}
            for field, label, limit, cell in zip(FIELDS, HEADERS, LIMITS, row):
                value = cell.value
                raw[label] = value if value is None or isinstance(value, (str, int, float, bool)) else str(value)
                if cell.data_type in ('f', 'e'):
                    errors.append({'field': field, 'message': '公式或错误单元格不能导入'})
                    value = None
                elif field == 'price_text' and isinstance(value, (int, float)) and not isinstance(value, bool):
                    value = str(value)
                elif value is not None and not isinstance(value, str):
                    errors.append({'field': field, 'message': '请将单元格设为文本并重新输入原文'})
                    value = None
                if isinstance(value, str):
                    value = value.replace('\r\n', '\n').replace('\r', '\n').strip() or None
                if value and len(value) > limit:
                    errors.append({'field': field, 'message': f'最多{limit}字符'})
                values[field] = value
            if not values['name']:
                errors.append({'field': 'name', 'message': '名称必填'})
            rows.append({'row_number': number, 'values': values, 'raw_values': raw, 'errors': errors,
                         'prices': parse_wholesale_prices(values['price_text'])})
        if not rows:
            error(400, 'no_product_rows', '没有可导入的产品行')
        return {'template_version': 3, 'sheet_name': '产品清单', 'rows': rows}
    result = read_excel(content, parser=parse_new)
    if result is None:
        legacy, _ = read_excel(content)
        rows = []
        for row in legacy:
            values = {'category': None, 'model': row['model'], 'name': row['name'], 'parameters': row['parameters'],
                      'price_text': row['wholesale_price'], 'control_price': None, 'selling_points': row['selling_points'], 'notes': row['price_note']}
            rows.append({'row_number': row['row_number'], 'values': values, 'raw_values': row['raw_values'],
                         'errors': row['errors'], 'prices': parse_wholesale_prices(values['price_text'])})
        result = {'template_version': 1, 'sheet_name': '产品数据', 'rows': rows}
    result['sha256'] = hashlib.sha256(content).hexdigest()
    return result


def baseline(db, rows, supplier, fingerprint):
    names = {row['values']['name'] for row in rows if row['values']['name']}
    products = db.scalars(select(LibraryProduct).where(LibraryProduct.name.in_(names)).order_by(LibraryProduct.id)).all()
    product_ids = [product.id for product in products]
    keys = db.scalars(select(LibrarySimpleQuoteKey).where(LibrarySimpleQuoteKey.product_id.in_(product_ids)).order_by(LibrarySimpleQuoteKey.id)).all()
    key_ids = [key.id for key in keys]
    quotes = db.scalars(select(LibrarySimpleQuoteVersion).where(LibrarySimpleQuoteVersion.quote_key_id.in_(key_ids)).order_by(LibrarySimpleQuoteVersion.id)).all()
    from sqlalchemy import or_
    sources = db.scalars(select(LibrarySimpleSource).where(or_(LibrarySimpleSource.product_id.in_(product_ids),
        (LibrarySimpleSource.file_sha256 == fingerprint) & (LibrarySimpleSource.supplier_name == supplier))).order_by(LibrarySimpleSource.id)).all()
    return digest({'products': [(p.id, p.version, p.model, p.name, bool(p.archived_at)) for p in products],
                   'keys': [(k.id, k.product_id, k.supplier_name, k.variant) for k in keys],
                   'quotes': [(q.id, q.quote_key_id, q.version_number, q.amount_cents, q.eligibility, q.is_current) for q in quotes],
                   'sources': [(s.id, s.product_id, s.version) for s in sources]})


def preview(content, supplier, filename, db):
    if not isinstance(supplier, str) or not (supplier := ' '.join(supplier.split())) or len(supplier) > 200:
        error(422, 'invalid_supplier', '请填写不超过200字符的供应商名称')
    require_identity_schema(db)
    parsed = parse(content)
    identities = {}
    for row in parsed['rows']:
        key = (row['values']['model'], row['values']['name'])
        if not row['errors']:
            identities.setdefault(key, []).append(row['row_number'])
    products = db.scalars(select(LibraryProduct)).all()
    for row in parsed['rows']:
        values = row['values']
        if not values['price_text']:
            row['price_status'], row['price_reason'] = 'empty', '未提供批发价'
        elif row['prices'] is None:
            row['price_status'], row['price_reason'] = 'pending', '价格类型、规格、金额或单位不明确，请核对原文'
        elif has_non_wholesale_semantics(values['notes']):
            row['price_status'], row['price_reason'] = 'pending', '来源备注含非出厂价格语义，请核对'
        elif parsed['template_version'] == 1:
            row['price_status'], row['price_reason'] = 'pending', '旧六列价格须确认出厂语义'
        else:
            row['price_status'], row['price_reason'] = 'eligible_on_confirm', '确认导入后按出厂批发价保存'
        source_names = re.findall(r'(?:供应商|厂家|供货商)(?:名称|姓名|名字)?\s*[:：]\s*([^\n;；，,]+)', values['notes'] or '')
        if any(' '.join(name.split()) != supplier for name in source_names):
            row['errors'].append({'field': 'notes', 'message': '来源标注与上传供应商不一致，请先拆分文件或核对归属'})
        identity = (values['model'], values['name'])
        if len(identities.get(identity, [])) > 1:
            row['errors'].append({'field': 'name', 'message': f"文件内身份重复：{identities[identity]}"})
        match = next((p for p in products if (p.model, p.name) == identity), None)
        row['product_id'] = match.id if match else None
        row['product_version'] = match.version if match else None
        row['archived'] = bool(match and match.archived_at)
        row['archived_sources'] = [{'supplier': source.supplier_name, 'notes': source.notes} for source in db.scalars(
            select(LibrarySimpleSource).where(LibrarySimpleSource.product_id == match.id).order_by(LibrarySimpleSource.id)
        ).all()] if row['archived'] else []
        row['differences'] = {field: {'current': getattr(match, field), 'incoming': values[field]}
                              for field in ('category', 'parameters', 'control_price', 'selling_points')
                              if match and values[field] is not None and getattr(match, field) != values[field]}
        row['potential_duplicates'] = [{'id': p.id, 'version': p.version, 'model': p.model, 'name': p.name}
                                       for p in products if not match and p.name == values['name'] and (p.model is None) != (values['model'] is None)]
        row['price_conflicts'] = []
        for variant, amount in row['prices'] or []:
            if match:
                current = db.execute(select(LibrarySimpleQuoteVersion, LibrarySimpleQuoteKey).join(LibrarySimpleQuoteKey).where(
                    LibrarySimpleQuoteKey.product_id == match.id, LibrarySimpleQuoteKey.supplier_name == supplier,
                    LibrarySimpleQuoteKey.variant == variant, LibrarySimpleQuoteVersion.is_current.is_(True))).first()
                if current and current[0].amount_cents != amount:
                    row['price_conflicts'].append({'variant': variant, 'current_cents': current[0].amount_cents, 'incoming_cents': amount})
    parsed.update(supplier=supplier, filename=filename, import_id=str(uuid4()), identity_rule_version=4, attachment_warning='原始文件和图片不会保存，请自行保留原件。')
    db.add(LibraryImport(import_id=parsed['import_id'], manufacturer_name=supplier, created_at=now(), expires_at=now()+timedelta(hours=1),
                         status='ready', preview=parsed, baseline_digest=baseline(db, parsed['rows'], supplier, parsed['sha256'])))
    db.commit()
    return parsed


def confirm(import_id, payload, db):
    if set(payload) != {'decisions', 'confirmed_attachments_not_imported'} or payload['confirmed_attachments_not_imported'] is not True or not isinstance(payload['decisions'], list):
        error(422, 'invalid_decisions', '请确认原件保存提示并逐行处理')
    decision_hash = digest({**payload, 'decisions': sorted(payload['decisions'], key=lambda value: json.dumps(value, ensure_ascii=False, sort_keys=True))})
    db.rollback()
    begin_write(db)
    require_identity_schema(db)
    item = db.get(LibraryImport, import_id)
    if item is None:
        error(404, 'import_not_found', '导入预览不存在')
    if item.status == 'committed':
        if item.decision_digest != decision_hash:
            error(409, 'import_already_committed', '该预览已按其他决策提交')
        return item.result
    data = item.preview
    if not data or data.get('template_version') not in (1, 3) or data.get('identity_rule_version') != 4 or not data.get('sha256'):
        error(409, 'import_stale', '旧预览已失效，请重新预览')
    if utc(item.expires_at) < now() or item.baseline_digest != baseline(db, data['rows'], data['supplier'], data['sha256']):
        error(409, 'import_stale', '产品或来源已变化，请重新预览')
    decisions = {}
    for choice in payload['decisions']:
        if not isinstance(choice, dict) or set(choice) - {'row_number', 'action', 'update_fields', 'price_decisions', 'allow_suspected_duplicate'} or type(choice.get('row_number')) is not int or choice.get('action') not in ('apply', 'skip') or choice['row_number'] in decisions:
            error(422, 'invalid_decisions', '行决策无效')
        decisions[choice['row_number']] = choice
    if set(decisions) != {row['row_number'] for row in data['rows']}:
        error(422, 'invalid_decisions', '必须处理每一行')
    results = []
    for row in data['rows']:
        choice = decisions[row['row_number']]
        if choice['action'] == 'skip':
            results.append({'row_number': row['row_number'], 'action': 'skipped', 'values': row['values'], 'errors': row['errors']})
            continue
        if row['errors'] or row['archived']:
            error(422, 'invalid_decisions', '错误或归档产品行必须跳过')
        if row['potential_duplicates'] and choice.get('allow_suspected_duplicate') is not True:
            error(422, 'invalid_decisions', '潜在重复产品需明确确认独立新增')
        updates = choice.get('update_fields', [])
        if not isinstance(updates, list) or any(not isinstance(field, str) or field not in row['differences'] for field in updates) or len(set(updates)) != len(updates):
            error(422, 'invalid_decisions', '主资料更新字段无效')
        price_choices = choice.get('price_decisions', {})
        if not isinstance(price_choices, dict) or set(price_choices) != {str(c['variant'] or '') for c in row['price_conflicts']} or any(v not in ('keep', 'update') for v in price_choices.values()):
            error(422, 'invalid_decisions', '请逐项处理报价冲突')
        values = row['values']
        product = db.get(LibraryProduct, row['product_id']) if row['product_id'] else None
        if product is None:
            product = LibraryProduct(model=values['model'], name=values['name'], category=values['category'], parameters=values['parameters'],
                                     control_price=values['control_price'], selling_points=values['selling_points'], created_at=now(), updated_at=now(), version=1)
            db.add(product)
            db.flush()
        else:
            for field in updates:
                setattr(product, field, values[field])
        source = db.scalar(select(LibrarySimpleSource).where(LibrarySimpleSource.file_sha256 == data['sha256'], LibrarySimpleSource.supplier_name == data['supplier'],
                                                           LibrarySimpleSource.sheet_name == data['sheet_name'], LibrarySimpleSource.row_number == row['row_number']))
        if source is None:
            source = LibrarySimpleSource(product_id=product.id, supplier_name=data['supplier'], source_type='excel', file_sha256=data['sha256'],
                                         file_name=data['filename'], sheet_name=data['sheet_name'], row_number=row['row_number'], raw_values={'template_version': data['template_version'], 'values': row['raw_values']},
                                         price_text=values['price_text'], notes=values['notes'], created_at=now(), version=1)
            db.add(source)
            db.flush()
            for variant, amount in row['prices'] or []:
                key = db.scalar(select(LibrarySimpleQuoteKey).where(LibrarySimpleQuoteKey.product_id == product.id, LibrarySimpleQuoteKey.supplier_name == data['supplier'], LibrarySimpleQuoteKey.variant == variant))
                if key is None:
                    key = LibrarySimpleQuoteKey(product_id=product.id, supplier_name=data['supplier'], variant=variant)
                    db.add(key)
                    db.flush()
                current = db.scalar(select(LibrarySimpleQuoteVersion).where(LibrarySimpleQuoteVersion.quote_key_id == key.id, LibrarySimpleQuoteVersion.is_current.is_(True)))
                status = 'eligible' if data['template_version'] == 3 and not has_non_wholesale_semantics(values['notes']) else 'pending'
                if current and current.amount_cents != amount and price_choices.get(str(variant or '')) == 'keep':
                    continue
                if current and current.amount_cents == amount and current.eligibility == 'eligible' and status == 'pending':
                    continue
                if current and current.amount_cents == amount and current.eligibility == status:
                    quote = current
                else:
                    version = (db.scalar(select(func.max(LibrarySimpleQuoteVersion.version_number)).where(LibrarySimpleQuoteVersion.quote_key_id == key.id)) or 0) + 1
                    if current:
                        current.is_current = False
                        db.flush()
                    quote = LibrarySimpleQuoteVersion(quote_key_id=key.id, version_number=version, amount_cents=amount, currency='CNY', unit='件', eligibility=status, is_current=True, created_at=now())
                    db.add(quote)
                    db.flush()
                db.add(LibrarySimpleQuoteEvidence(quote_version_id=quote.id, quote_key_id=key.id, product_id=product.id, source_id=source.id))
            product.version += 1
            product.updated_at = now()
        elif source.product_id != product.id:
            error(409, 'import_stale', '来源身份已变化，请重新预览')
        elif updates:
            product.version += 1
            product.updated_at = now()
        results.append({'row_number': row['row_number'], 'action': 'imported', 'product_id': product.id, 'source_id': source.id, 'values': row['values'], 'errors': row['errors']})
    item.status = 'committed'
    item.decision_digest = decision_hash
    item.result = {'import_id': import_id, 'template_version': data['template_version'], 'simple_import': True,
                   'attachments_not_imported': True,
                   'supplier': data['supplier'], 'filename': data['filename'], 'sheet_name': data['sheet_name'],
                   'sha256': data['sha256'], 'decisions': sorted(payload['decisions'], key=lambda choice: choice['row_number']), 'rows': results}
    db.commit()
    return item.result
