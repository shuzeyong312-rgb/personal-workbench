"""HTTP endpoints for the independent manufacturer product library."""

import io

from fastapi import APIRouter, Body, Depends, Query, Request
from fastapi.responses import Response
from starlette.concurrency import run_in_threadpool
from urllib.parse import unquote
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from sqlalchemy.orm import Session

from app.database import get_db
from app.product_library_service import (
    HEADERS,
    batch_delete_service,
    MIB,
    error,
    confirm_import_service,
    create_product_service,
    delete_product_service,
    get_import_service,
    get_product_service,
    list_products_service,
    manufacturers_service,
    preview_import_service,
    update_product_service,
)
from app.product_library_quotations import preview_quotation_service, export_quotation_service

router = APIRouter(prefix="/api/product-library", tags=["product-library"])


@router.post("/products/{product_id}/{kind}", status_code=201)
def create_detail(product_id: int, kind: str, payload: dict = Body(...), db: Session = Depends(get_db)):
    from app.product_library_details import mutate_detail
    return mutate_detail(product_id, kind, None, payload, "POST", db)


@router.put("/products/{product_id}/{kind}/{record_id}")
def update_detail(product_id: int, kind: str, record_id: int, payload: dict = Body(...), db: Session = Depends(get_db)):
    from app.product_library_details import mutate_detail
    return mutate_detail(product_id, kind, record_id, payload, "PUT", db)


@router.delete("/products/{product_id}/{kind}/{record_id}", status_code=204)
def delete_detail(product_id: int, kind: str, record_id: int, payload: dict = Body(...), db: Session = Depends(get_db)):
    from app.product_library_details import mutate_detail
    mutate_detail(product_id, kind, record_id, payload, "DELETE", db)
    return Response(status_code=204)


@router.get("/products")
def list_products(
    q: str = Query("", max_length=200),
    manufacturer: str | None = None,
    sort: str = "updated_desc",
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    price_status: str = "all",
    price_mode: str | None = None,
    price_min: str | None = None,
    price_max: str | None = None,
    db: Session = Depends(get_db),
):
    return list_products_service(
        q=q, manufacturer=manufacturer, sort=sort, page=page, page_size=page_size, db=db,
        price_status=price_status, price_mode=price_mode, price_min=price_min, price_max=price_max,
    )


@router.get("/manufacturers")
def manufacturers(db: Session = Depends(get_db)):
    return manufacturers_service(db=db)


@router.post("/products/batch-delete")
def batch_delete(payload: dict = Body(...), db: Session = Depends(get_db)):
    return batch_delete_service(payload, db)


@router.post("/customer-quotations/preview")
def preview_quotation(payload: dict = Body(...), db: Session = Depends(get_db)):
    return preview_quotation_service(payload, db)


@router.post("/customer-quotations/export")
def export_quotation(payload: dict = Body(...), db: Session = Depends(get_db)):
    return export_quotation_service(payload, db)


@router.get("/products/{product_id}")
def get_product(product_id: int, db: Session = Depends(get_db)):
    return get_product_service(product_id=product_id, db=db)


@router.post("/products", status_code=201)
def create_product(payload: dict = Body(...), db: Session = Depends(get_db)):
    return create_product_service(payload=payload, db=db)


@router.put("/products/{product_id}")
def update_product(
    product_id: int, payload: dict = Body(...), db: Session = Depends(get_db)
):
    return update_product_service(product_id=product_id, payload=payload, db=db)


@router.delete("/products/{product_id}", status_code=204)
def delete_product(
    product_id: int, payload: dict = Body(...), db: Session = Depends(get_db)
):
    return delete_product_service(product_id=product_id, payload=payload, db=db)


@router.get("/template")
def template(template_version: int = 1):
    if template_version not in (1, 2):
        error(422, 'invalid_template_version', '模板版本仅支持1或2')
    if template_version == 2:
        from app.product_library_v2_import import template_bytes
        return Response(template_bytes(), media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': 'attachment; filename="product-library-v2.xlsx"'})
    wb = Workbook()
    ws = wb.active
    ws.title = "产品数据"
    ws.append(HEADERS)
    ws.freeze_panes = "A2"
    for col in range(1, 7):
        ws.column_dimensions[chr(64 + col)].width = 26
        ws.cell(1, col).font = Font(bold=True)
        for row in range(2, 1002):
            cell = ws.cell(row, col)
            if col != 5:
                cell.number_format = "@"
            cell.alignment = Alignment(wrap_text=col in (3, 4, 6))
    notes = wb.create_sheet("填写说明")
    for line in [
        "示例仅供说明，不是待导入产品。厂家在上传页面统一填写。",
        "型号/名称最多200字符，名称必填；型号缺失留空，数字型号请存为文本以保留前导零。",
        "参数/卖点最多10000字符，保留换行、原单位及符号；价格备注最多2000字符。",
        "批发价人民币元：0～999999999.99，最多两位小数；未知留空，不填0。",
        "阶梯价、起售价、面议、价格范围只写价格备注，不推算确定价。",
        "示例有型号：001A / 便携风扇 / 3.7V 3000mAh / 轻便 / 12.30",
        "示例无型号：型号留空 / 台灯；价格未知：批发价留空；备注价格示例：100件起12元，面议。",
        "导入更新时空值保留旧值；明确清空请使用手动编辑。",
        "最多1000行，5 MiB，只支持.xlsx；不使用公式、合并或隐藏行列。",
    ]:
        notes.append([line])
    notes.column_dimensions["A"].width = 100
    output = io.BytesIO()
    wb.save(output)
    wb.close()
    from urllib.parse import quote

    return Response(
        output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": "attachment; filename*=UTF-8''"
            + quote("产品库导入模板.xlsx")
        },
    )


@router.post("/imports/preview", status_code=201)
async def preview_import(
    request: Request, manufacturer_name: str = Query(...), db: Session = Depends(get_db)
):
    filename = request.headers.get("X-Filename")
    if not filename or not unquote(filename).lower().endswith(".xlsx"):
        error(415, "unsupported_excel_format", "请提供 .xlsx 文件名（X-Filename）")
    if len(unquote(filename)) > 500:
        error(422, "invalid_product", "工作簿文件名最多500字符", field="X-Filename")
    content = bytearray()
    async for chunk in request.stream():
        content.extend(chunk)
        if len(content) > 5 * MIB:
            error(413, "file_too_large", "文件不能超过5 MiB")
    def preview():
        from app.product_library_v2_import import preview as preview_v2
        result = preview_v2(bytes(content), manufacturer_name, db, unquote(filename))
        return result if result is not None else preview_import_service(bytes(content), manufacturer_name, db, unquote(filename))
    return await run_in_threadpool(preview)


@router.get("/imports/{import_id}")
def get_import(import_id: str, db: Session = Depends(get_db)):
    return get_import_service(import_id=import_id, db=db)


@router.post("/imports/{import_id}/confirm")
def confirm_import(
    import_id: str, payload: dict = Body(...), db: Session = Depends(get_db)
):
    from app.models import LibraryImport
    item = db.get(LibraryImport, import_id)
    version = ((item.result if item.status == 'committed' else item.preview) or {}).get('template_version') if item else None
    db.rollback()
    if version == 2:
        from app.product_library_v2_import import confirm as confirm_v2
        return confirm_v2(import_id, payload, db)
    return confirm_import_service(import_id=import_id, payload=payload, db=db)
