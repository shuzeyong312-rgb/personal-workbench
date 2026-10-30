"""Product library: bounded Excel previews and atomic, durable confirmations."""

import hashlib
import io
import json
import re
import zipfile
from collections import defaultdict
from datetime import datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from uuid import UUID, uuid4
from xml.etree import ElementTree as ET

from fastapi import HTTPException
from fastapi.responses import Response
from openpyxl import load_workbook
from sqlalchemy import inspect, or_, select, text
from sqlalchemy.exc import IntegrityError, OperationalError, SQLAlchemyError

from app.models import LibraryImport, LibraryProduct

FIELDS = (
    "model",
    "name",
    "parameters",
    "selling_points",
    "wholesale_price",
    "price_note",
)
HEADERS = ("产品型号", "产品名称", "产品参数", "产品卖点", "批发价", "价格备注")
LIMITS = {
    "manufacturer_name": 200,
    "model": 200,
    "name": 200,
    "parameters": 10000,
    "selling_points": 10000,
    "price_note": 2000,
}
IDENTITY_RULE_VERSION = 4
V2_IMPORT_PLAN_VERSION = 2
MIB = 1024 * 1024


def error(status, code, message, **extra):
    raise HTTPException(status, {"code": code, "message": message, **extra})


def now():
    return datetime.now(timezone.utc)


def utc(value):
    return value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value


def normalized_text(value, field):
    if value is None:
        value = ""
    if not isinstance(value, str):
        raise ValueError("必须是文本，不能使用数字、布尔或日期")
    value = value.replace("\r\n", "\n").replace("\r", "\n").strip()
    if field == "manufacturer_name":
        value = " ".join(value.split())
    if not value and field in ("manufacturer_name", "name"):
        raise ValueError("必填字段不能为空")
    if len(value) > LIMITS[field]:
        raise ValueError(f"最多 {LIMITS[field]} 个字符")
    return value or None


def price(value, *, excel=False):
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if (
        isinstance(value, bool)
        or not isinstance(value, (str, int, float, Decimal))
        or (not excel and not isinstance(value, str))
    ):
        raise ValueError("批发价必须是十进制金额文本")
    raw = str(value).strip()
    if isinstance(value, str) and not re.fullmatch(r"[0-9]+(?:\.[0-9]+)?", raw):
        raise ValueError("批发价须为非负十进制金额，不含符号或范围")
    try:
        amount = Decimal(raw)
        if (
            not amount.is_finite()
            or amount < 0
            or amount > Decimal("999999999.99")
            or amount.as_tuple().exponent < -2
        ):
            raise ValueError("批发价范围为 0～999999999.99，最多两位小数")
        return f"{amount:.2f}"
    except InvalidOperation:
        raise ValueError("批发价格式错误")


def validate(data, *, excel=False):
    values, errors = {}, []
    for field in ("manufacturer_name", *FIELDS):
        try:
            values[field] = (
                price(data.get(field), excel=excel)
                if field == "wholesale_price"
                else normalized_text(data.get(field), field)
            )
        except ValueError as exc:
            values[field] = None
            errors.append(
                {"field": field, "code": "invalid_field", "message": str(exc)}
            )
    return values, errors


def product_view(product):
    from sqlalchemy.orm import object_session
    from app.models import LibrarySimpleQuoteKey, LibrarySimpleQuoteVersion, LibrarySimpleSource
    db = object_session(product)
    source_companies = sorted(set(db.scalars(select(LibrarySimpleSource.supplier_name).where(LibrarySimpleSource.product_id == product.id, LibrarySimpleSource.archived_at.is_(None))))) if db else []
    quoted = bool(db.scalar(select(LibrarySimpleQuoteVersion.id).join(LibrarySimpleQuoteKey).where(
        LibrarySimpleQuoteKey.product_id == product.id, LibrarySimpleQuoteVersion.is_current.is_(True),
        LibrarySimpleQuoteVersion.eligibility == 'eligible').limit(1))) if db else False
    return {
        "category": product.category,
        "control_price": product.control_price,
        "source_companies": source_companies,
        "price_state": "multiple" if quoted else "unquoted",
        "id": product.id,
        "manufacturer_name": product.manufacturer_name,
        **{
            field: getattr(product, field)
            for field in FIELDS
            if field != "wholesale_price"
        },
        "wholesale_price": None,
        "version": product.version,
        "created_at": utc(product.created_at).isoformat(),
        "updated_at": utc(product.updated_at).isoformat(),
    }


def assign(product, data):
    for field, value in data.items():
        if field == "wholesale_price":
            product.wholesale_price_cents = (
                None if value is None else int(Decimal(value) * 100)
            )
        else:
            setattr(product, field, value)


def begin_write(db):
    db.execute(text("BEGIN IMMEDIATE"))


def database_error(db, exc, code):
    db.rollback()
    if isinstance(exc, IntegrityError) and "UNIQUE constraint failed: library_products." in str(exc):
        error(409, "identity_conflict", "精确产品身份已存在，本次未保存，全部回滚；请重新预览核对")
    if isinstance(exc, IntegrityError):
        error(409, 'related_record_conflict', '关联记录重复或约束冲突，本次全部回滚；请重新核对')
    if isinstance(exc, OperationalError) and (
        "locked" in str(exc).lower() or "busy" in str(exc).lower()
    ):
        error(503, "database_busy", "数据库忙，请按原请求重试")
    error(500, code, "本次未保存，全部回滚")


def product_or_404(db, product_id):
    if product_id < 1:
        error(404, "product_not_found", "产品不存在")
    product = db.get(LibraryProduct, product_id)
    if product is None or product.archived_at is not None:
        error(404, "product_not_found", "产品不存在")
    return product


def require_identity_schema(db):
    schema = inspect(db.connection())
    if not {'library_simple_sources','library_simple_quote_keys','library_simple_quote_versions','library_simple_quote_evidence'}.issubset(schema.get_table_names()) or not {'category','control_price','archived_at'}.issubset({column['name'] for column in schema.get_columns('library_products')}):
        error(503, 'identity_schema_required', '产品库V2尚未完成授权迁移；本次未写入')
    indexes = {row[1]: row for row in db.execute(text('PRAGMA index_list("library_products")'))}
    for name, columns, predicate in (
        ("uq_library_products_model_name", ["model", "name"], "MODEL IS NOT NULL"),
        ("uq_library_products_null_model_name", ["name"], "MODEL IS NULL"),
    ):
        info = list(db.execute(text(f'PRAGMA index_xinfo("{name}")')))
        sql = db.scalar(text("SELECT sql FROM sqlite_master WHERE name=:name"), {"name": name}) or ""
        if (name not in indexes or not indexes[name][2]
            or [r[2] for r in info if r[5]] != columns
            or any(r[4] != "BINARY" for r in info if r[5])
            or sql.upper().split("WHERE")[-1].strip() != predicate):
            error(503, "identity_schema_required", "产品身份规则尚未启用：数据库需完成授权迁移；本次未写入")


def candidates_for(db, values, exclude=None):
    products = [p for p in all_products(db) if p.id != exclude]
    return [p for p in products if p.model == values["model"] and p.name == values["name"]]


def manual_payload(payload, editing=False):
    required = (
        {"manufacturer_name", *FIELDS} if editing else {"manufacturer_name", "name"}
    )
    allowed = {"manufacturer_name", *FIELDS, "allow_suspected_duplicate", "category", "structured_attributes", "confirm_manufacturer_change"} | (
        {"version"} if editing else set()
    )
    if set(payload) - allowed or not required.issubset(payload):
        error(422, "invalid_product", "字段缺失或含未知字段")
    if "allow_suspected_duplicate" in payload and not isinstance(
        payload["allow_suspected_duplicate"], bool
    ):
        error(422, "invalid_product", "重复确认必须是布尔值")
    if 'confirm_manufacturer_change' in payload and type(payload['confirm_manufacturer_change']) is not bool:
        error(422, 'invalid_product', '首录资料公司确认必须为布尔值')
    values, errors = validate(payload)
    if errors:
        error(422, "invalid_product", "产品字段校验失败", errors=errors)
    if "category" in payload:
        from app.product_library_details import text_value
        try:
            values['category'] = text_value(payload['category'], 'category')
        except ValueError as exc:
            error(422, 'invalid_product', str(exc))
    if "structured_attributes" in payload:
        values['structured_attributes'] = payload['structured_attributes']
    return values


def version_value(payload):
    version = payload.get("version")
    if type(version) is not int or version < 1:
        error(422, "invalid_product", "version 必须为正整数")
    return version


def check_duplicates(db, values, exclude=None):
    exact = candidates_for(db, values, exclude)
    if exact:
        if exact[0].archived_at is not None:
            error(409, 'product_archived', '该产品已归档，请恢复原产品或取消', candidates=[product_view(p) for p in exact])
        error(
            409,
            "product_exists",
            "同型号同名称产品已存在，请编辑已有产品",
            candidates=[product_view(p) for p in exact],
        )


def read_excel(content, parser=None):
    if len(content) > 5 * MIB:
        error(413, "file_too_large", "文件不能超过 5 MiB")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if (
                len(entries) > 1000
                or sum(entry.file_size for entry in entries) > 50 * MIB
            ):
                error(413, "workbook_limit_exceeded", "工作簿解压大小或条目超限")
            seen, total = set(), 0
            for entry in entries:
                name = entry.filename.lower()
                if name in seen:
                    error(400, "invalid_workbook", "工作簿含重复包条目")
                seen.add(name)
                if any(
                    token in name
                    for token in ("vbaproject", "embeddings/", "externallinks/")
                ):
                    error(400, "invalid_workbook", "不允许宏、嵌入对象或外链")
                with archive.open(entry) as member:
                    chunks = []
                    while True:
                        chunk = member.read(min(MIB, 50 * MIB - total + 1))
                        if not chunk:
                            break
                        total += len(chunk)
                        if total > 50 * MIB:
                            error(413, "workbook_limit_exceeded", "工作簿解压超限")
                        chunks.append(chunk)
                data = b"".join(chunks)
                # XML parts may use other suffixes and NUL-containing encodings.
                probe = data.replace(b"\x00", b"").lower()
                if b"<!doctype" in probe or b"<!entity" in probe:
                    error(400, "invalid_workbook", "不允许 DTD 或实体声明")
                if name.endswith((".xml", ".rels")):
                    for _, element in ET.iterparse(io.BytesIO(data), events=("end",)):
                        tag = element.tag.rsplit("}", 1)[-1]
                        relation = (
                            element.attrib.get("Type", "").rsplit("/", 1)[-1].lower()
                        )
                        if name.endswith(".rels") and (
                            element.attrib.get("TargetMode", "").lower() == "external"
                            or relation in ("package", "oleobject")
                        ):
                            error(400, "invalid_workbook", "不允许外部关系或嵌入对象")
                        if tag in ("oleObjects", "oleObject"):
                            error(400, "invalid_workbook", "不允许嵌入对象")
                        element.clear()
                    if name == "[content_types].xml" and (
                        b"macroenabled" in probe
                        or b"vba" in probe
                        or b"oleobject" in probe
                    ):
                        error(415, "unsupported_excel_format", "仅支持无宏 .xlsx")
        wb = load_workbook(
            io.BytesIO(content),
            read_only=True,
            data_only=False,
            keep_links=False,
            keep_vba=False,
        )
        try:
            if parser is not None:
                return parser(wb, content)
            if any(name in wb.sheetnames for name in ("参数明细", "规格款式", "价格明细", "资料索引")):
                error(400, "invalid_structure", "混合版本：请使用完整V2模板")
            if "产品数据" not in wb.sheetnames:
                error(400, "invalid_structure", "缺少产品数据工作表")
            ws = wb["产品数据"]
            if ws.sheet_state != "visible":
                error(400, "invalid_structure", "产品数据工作表不能隐藏")
            # Inspect actual XML before reading cells; dimensions and styled tails are not data.
            from openpyxl.utils.cell import (
                coordinate_from_string,
                column_index_from_string,
            )

            last = 1
            with zipfile.ZipFile(io.BytesIO(content)) as archive, archive.open(
                ws._worksheet_path
            ) as source:
                sheet_data = None
                for event, element in ET.iterparse(source, events=("start", "end")):
                    tag = element.tag.rsplit("}", 1)[-1]
                    if event == "start":
                        if tag == "sheetData":
                            sheet_data = element
                        continue
                    if tag == "mergeCell" or (
                        tag in ("row", "col")
                        and element.attrib.get("hidden") in ("1", "true", "True")
                    ):
                        error(
                            400,
                            "invalid_structure",
                            "产品数据不能含合并单元格或隐藏行列",
                        )
                    if tag == "c":
                        nonempty = element.attrib.get("t") == "e" or any(
                            child.tag.rsplit("}", 1)[-1] == "f"
                            or (
                                child.tag.rsplit("}", 1)[-1] in ("v", "t")
                                and child.text is not None
                            )
                            for child in element.iter()
                        )
                        if nonempty:
                            column, number = coordinate_from_string(element.attrib["r"])
                            if column_index_from_string(column) > 6:
                                error(400, "invalid_structure", "六列以外存在非空内容")
                            if number > 1001:
                                error(
                                    413,
                                    "workbook_limit_exceeded",
                                    "最多1000个物理数据行",
                                )
                            last = max(last, number)
                        element.clear()
                    elif tag == "row":
                        element.clear()
                        if sheet_data is not None:
                            sheet_data.clear()
            ws.reset_dimensions()
            data_rows = ws.iter_rows(min_row=1, max_row=last, max_col=6)
            if tuple(cell.value for cell in next(data_rows)) != HEADERS:
                error(400, "invalid_structure", "表头必须严格匹配模板六列")
            rows, blank = [], 0
            for number, row_cells in enumerate(data_rows, 2):
                raw = {field: cell.value for field, cell in zip(FIELDS, row_cells)}
                if not any(cell.data_type in ("f", "e") for cell in row_cells) and all(
                    value is None or (isinstance(value, str) and not value.strip())
                    for value in raw.values()
                ):
                    blank += 1
                    continue
                values, errors = validate(
                    {"manufacturer_name": "placeholder", **raw}, excel=True
                )
                values.pop("manufacturer_name")
                for field, cell in zip(FIELDS, row_cells):
                    if cell.data_type in ("f", "e"):
                        errors.append(
                            {
                                "field": field,
                                "code": "formula_or_error",
                                "message": "公式或错误单元格不可导入",
                            }
                        )
                rows.append({"row_number": number, **values, "errors": errors, "raw_values": {field: value if isinstance(value, (str, int, float, bool)) or value is None else str(value) for field, value in raw.items()}})
            if not rows:
                error(400, "no_product_rows", "没有可导入的非空产品行")
            return rows, blank
        finally:
            wb.close()
    except HTTPException:
        raise
    except Exception:
        error(400, "invalid_workbook", "工作簿损坏、加密或格式无效")


def digest(value):
    return hashlib.sha256(
        json.dumps(
            value, ensure_ascii=False, sort_keys=True, separators=(",", ":")
        ).encode()
    ).hexdigest()


def all_products(db):
    return db.scalars(
        select(LibraryProduct)
        .order_by(LibraryProduct.id)
    ).all()


def baseline(products):
    from sqlalchemy.orm import object_session
    from app.product_library_details import product_details
    return digest([{**product_view(product), **product_details(object_session(product),product)} for product in products])


def identity_products(db, rows):
    identities = {(row['model'], row['name']) for row in rows if not row['errors']}
    return [product for product in db.scalars(select(LibraryProduct).order_by(LibraryProduct.id)) if (product.model, product.name) in identities]


def field_diff(row, product):
    old = product_view(product)
    return {field: {"old": old[field], "incoming": row[field],
                    "effective": row[field] if field != "wholesale_price" and row[field] is not None else old[field],
                    "changed": field != "wholesale_price" and row[field] is not None and row[field] != old[field],
                    "preserved": row[field] is None} for field in FIELDS}


def preview_rows(rows, products):
    valid = [row for row in rows if not row["errors"]]
    groups = defaultdict(list)
    for row in valid:
        groups[(row["model"], row["name"])].append(row)
    file_duplicates = {}
    for group in groups.values():
        for row in group[1:]:
            file_duplicates[row["row_number"]] = group[0]
    for row in rows:
        row.update(candidate=[], file_candidates=[], existing_id=None, version=None, diff={}, allowed_actions=["skip"], default_action="skip", identity_reason="")
        if row["errors"]:
            row["category"] = "invalid"
            continue
        exact = [p for p in products if p.model == row["model"] and p.name == row["name"]]
        duplicate_of = file_duplicates.get(row["row_number"])
        if duplicate_of:
            row["file_candidates"] = [{field: duplicate_of[field] for field in ("row_number", "name", "model")}]
        for product in exact:
            diff = field_diff(row, product)
            can_update = any(value["changed"] for value in diff.values())
            row["candidate"].append({"id": product.id, "version": product.version, "name": product.name, "model": product.model,
                                     "reason": "同型号、同名称", "can_update": can_update, "diff": diff})
        update_candidates = [c for c in row["candidate"] if c["can_update"]]
        if exact or duplicate_of:
            reason = "同型号、同名称已存在" if exact else f"与文件第 {duplicate_of['row_number']} 行重复"
            row.update(category="duplicate", identity_reason=reason)
            if exact:
                row.update(existing_id=exact[0].id, version=exact[0].version, diff=field_diff(row, exact[0]))
        else:
            row.update(category="new", identity_reason="同型号、同名称不存在", allowed_actions=["insert", "skip"], default_action="insert")
        if update_candidates:
            row["allowed_actions"].append("update")
            if len(update_candidates) == 1:
                candidate = update_candidates[0]
                row.update(existing_id=candidate["id"], version=candidate["version"], diff=candidate["diff"])
    return rows


def receipt(db, import_id):
    try:
        UUID(import_id)
    except ValueError:
        error(404, "import_not_found", "导入预览不存在")
    item = db.get(LibraryImport, import_id)
    if item is None:
        error(404, "import_not_found", "导入预览不存在")
    if item.status == "ready" and (item.preview or {}).get("identity_rule_version") != IDENTITY_RULE_VERSION:
        error(409, "preview_rule_changed", "预览使用旧身份规则，不能继续确认；请重新上传并核对")
    if (item.status == "ready" and (item.preview or {}).get("template_version") == 2
            and (item.preview or {}).get("import_plan_version") != V2_IMPORT_PLAN_VERSION):
        error(409, "preview_rule_changed", "预览使用旧来源保存规则，请重新上传并核对原文归属")
    if item.status == "ready" and utc(item.expires_at) <= now():
        error(410, "preview_expired", "预览已过期，请重新上传")
    return item


def decisions(payload):
    if set(payload) - {"decisions", "confirmed_attachments_not_imported"} or not isinstance(payload.get("decisions"), list):
        error(422, "invalid_decisions", "必须提供逐行决策")
    if 'confirmed_attachments_not_imported' in payload and type(payload['confirmed_attachments_not_imported']) is not bool:
        error(422, 'invalid_decisions', '附件确认必须为布尔值')
    result = {}
    for decision in payload["decisions"]:
        if (not isinstance(decision, dict) or not {"row_number", "action"}.issubset(decision)
            or set(decision) - {"row_number", "action", "target_id", "target_version"}
            or type(decision["row_number"]) is not int or decision["row_number"] < 1
            or not isinstance(decision["action"], str) or decision["action"] not in ("insert", "update", "append_source", "skip")
            or decision["row_number"] in result):
            error(422, "invalid_decisions", "决策行号或动作无效/重复")
        target = {key: decision[key] for key in ("target_id", "target_version") if key in decision}
        if target and (decision["action"] not in ("update", "append_source") or set(target) != {"target_id", "target_version"} or any(type(value) is not int or value < 1 for value in target.values())):
            error(422, "invalid_decisions", "更新目标ID和版本必须为正整数，且仅限更新操作")
        result[decision["row_number"]] = {"action": decision["action"], **target}
    return result


def list_products_service(q, manufacturer, sort, page, page_size, db,
                          price_status="all", price_mode=None, price_min=None, price_max=None):
    if sort not in ("updated_desc", "price_asc", "price_desc"):
        error(422, "invalid_product", "排序无效")
    query = select(LibraryProduct).where(LibraryProduct.archived_at.is_(None))
    if price_status not in ("all", "quoted", "unquoted"):
        error(422, "invalid_product", "报价状态无效")
    from sqlalchemy import func
    from app.models import LibrarySimpleQuoteKey, LibrarySimpleQuoteVersion, LibrarySimpleSource
    eligible = select(LibrarySimpleQuoteVersion.amount_cents).join(LibrarySimpleQuoteKey).where(
        LibrarySimpleQuoteKey.product_id == LibraryProduct.id, LibrarySimpleQuoteVersion.is_current.is_(True),
        LibrarySimpleQuoteVersion.eligibility == 'eligible', LibrarySimpleQuoteVersion.currency == 'CNY', LibrarySimpleQuoteVersion.unit == '件')
    if price_status != "all":
        quoted = eligible.exists()
        query = query.where(~quoted if price_status == "unquoted" else quoted)
    if price_mode not in (None, "below", "above", "between"):
        error(422, "invalid_product", "价格条件无效")
    try:
        low, high = price(price_min), price(price_max)
        if (price_mode is None and (low is not None or high is not None)
            or price_mode == "below" and (high is None or low is not None)
            or price_mode == "above" and (low is None or high is not None)
            or price_mode == "between" and (low is None or high is None or Decimal(low) > Decimal(high))):
            raise ValueError("价格边界缺失或与所选条件不一致")
    except ValueError as exc:
        error(422, "invalid_product", str(exc))
    amount = eligible.with_only_columns(func.min(LibrarySimpleQuoteVersion.amount_cents)).scalar_subquery()
    if low is not None:
        cents = int(Decimal(low) * 100)
        query = query.where(amount > cents if price_mode == "above" else amount >= cents)
    if high is not None:
        cents = int(Decimal(high) * 100)
        query = query.where(amount < cents if price_mode == "below" else amount <= cents)
    if q.strip():
        literal = (
            q.strip().replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
        )
        query = query.where(
            or_(
                *(
                    getattr(LibraryProduct, field).like(f"%{literal}%", escape="\\")
                    for field in (
                        "model",
                        "name",
                        "parameters",
                        "selling_points",
                    )
                ),
                select(LibrarySimpleSource.id).where(LibrarySimpleSource.product_id == LibraryProduct.id,
                    LibrarySimpleSource.archived_at.is_(None), LibrarySimpleSource.supplier_name.like(f"%{literal}%", escape="\\")).exists(),
            )
        )
    if manufacturer is not None:
        try:
            manufacturer = normalized_text(manufacturer, "manufacturer_name")
        except ValueError as exc:
            error(422, "invalid_product", str(exc))
        sources = select(LibrarySimpleSource.id).where(LibrarySimpleSource.product_id == LibraryProduct.id, LibrarySimpleSource.supplier_name == manufacturer, LibrarySimpleSource.archived_at.is_(None))
        query = query.where(sources.exists())

    total = db.scalar(select(func.count()).select_from(query.subquery()))
    ordering = (
        [LibraryProduct.updated_at.desc(), LibraryProduct.id.desc()]
        if sort == "updated_desc"
        else [
            amount.is_(None),
            (
                amount.asc()
                if sort == "price_asc"
                else amount.desc()
            ),
            LibraryProduct.id.desc(),
        ]
    )
    items = db.scalars(
        query.order_by(*ordering).offset((page - 1) * (page_size or 0)).limit(page_size)
    ).all()
    return {
        "items": [product_view(p) for p in items],
        "total": total,
        "page": page,
        "page_size": page_size,
    }


def product_versions(items):
    if not isinstance(items, list) or not items:
        error(422, "invalid_product", "请选择产品")
    seen = set()
    for item in items:
        if (not isinstance(item, dict) or set(item) != {"id", "version"}
                or type(item["id"]) is not int or item["id"] < 1
                or type(item["version"]) is not int or item["version"] < 1
                or item["id"] in seen):
            error(422, "invalid_product", "产品 ID/version 非法或重复")
        seen.add(item["id"])
    return items


def batch_delete_service(payload, db):
    if set(payload) != {"items"}:
        error(422, "invalid_product", "必须仅提供 items")
    items = product_versions(payload["items"])
    try:
        begin_write(db)
        rows = []
        for item in items:
            product = db.get(LibraryProduct, item["id"])
            outcome = "missing" if product is None or product.archived_at is not None else "stale" if product.version != item["version"] else "deleted"
            if outcome == "deleted":
                from app.product_library_simple_quotes import archive_product
                archive_product(product, db)
            rows.append({**item, "outcome": outcome})
        db.commit()
        return {"items": rows, **{key: sum(row["outcome"] == key for row in rows) for key in ("deleted", "missing", "stale")}}
    except SQLAlchemyError as exc:
        database_error(db, exc, "product_save_failed")


def manufacturers_service(db):
    from app.models import LibrarySimpleSource
    companies = set(db.scalars(select(LibrarySimpleSource.supplier_name).join(LibraryProduct).where(LibraryProduct.archived_at.is_(None), LibrarySimpleSource.archived_at.is_(None))))
    return {"items": sorted(companies)}


def get_product_service(product_id, db):
    from app.product_library_simple_quotes import simple_details
    product = product_or_404(db, product_id)
    return {**product_view(product), **simple_details(product, db), "attachment_warning": "原始文件和图片不会保存，请自行保留原件。"}


def simple_product_payload(payload, editing=False):
    limits = {"category": 200, "model": 200, "name": 200, "parameters": 10000, "control_price": 2000, "selling_points": 10000}
    allowed = set(limits) | {"allow_suspected_duplicate"} | ({"version", "confirm_model_completion"} if editing else set())
    if set(payload) - allowed or "name" not in payload:
        error(422, "invalid_product", "产品字段缺失或含未知字段")
    if "allow_suspected_duplicate" in payload and type(payload["allow_suspected_duplicate"]) is not bool:
        error(422, "invalid_product", "重复确认必须是布尔值")
    if "confirm_model_completion" in payload and type(payload["confirm_model_completion"]) is not bool:
        error(422, "invalid_product", "补充型号确认必须是布尔值")
    values = {}
    for field, limit in limits.items():
        value = payload.get(field)
        if value is not None and not isinstance(value, str):
            error(422, "invalid_product", f"{field}必须是文本", field=field)
        value = value.replace('\r\n', '\n').replace('\r', '\n').strip() if value is not None else None
        if len(value or '') > limit or field == 'name' and not value:
            error(422, "invalid_product", f"{field}为空或超过{limit}字符", field=field)
        values[field] = value or None
    return values


def create_product_service(payload, db):
    values = simple_product_payload(payload)
    try:
        begin_write(db)
        require_identity_schema(db)
        check_duplicates(db, values)
        potential = [p for p in all_products(db) if p.name == values['name'] and (p.model is None) != (values['model'] is None)]
        if potential and payload.get('allow_suspected_duplicate') is not True:
            error(409, 'potential_duplicate', '存在同名且型号为空的潜在重复产品，请核对后明确确认独立新增', candidates=[product_view(p) for p in potential])
        product = LibraryProduct(created_at=now(), updated_at=now(), version=1)
        assign(product, values)
        db.add(product)
        db.flush()
        result = get_product_service(product.id, db)
        db.commit()
        return result
    except HTTPException:
        db.rollback()
        raise
    except SQLAlchemyError as exc:
        database_error(db, exc, "product_save_failed")


def update_product_service(product_id, payload, db):
    values = simple_product_payload(payload, editing=True)
    version = version_value(payload)
    try:
        begin_write(db)
        require_identity_schema(db)
        product = product_or_404(db, product_id)
        if product.version != version:
            error(409, "product_stale", "产品已变化，请重新读取")
        check_duplicates(db, values, product_id)
        if (product.model, product.name) != (values['model'], values['name']):
            potential = [p for p in all_products(db) if p.id != product_id and p.name == values['name'] and (p.model is None) != (values['model'] is None)]
            if potential and payload.get('allow_suspected_duplicate') is not True:
                error(409, 'potential_duplicate', '存在同名且型号为空的潜在重复产品，请核对后明确确认', candidates=[product_view(p) for p in potential])
        if product.model is None and values['model'] is not None and payload.get('confirm_model_completion') is not True:
            error(409, 'model_completion_confirmation_required', '补充型号前请确认更新原产品')
        if any(getattr(product, field) != value for field, value in values.items()):
            assign(product, values)
            product.version += 1
            product.updated_at = now()
        db.flush()
        result = get_product_service(product.id, db)
        db.commit()
        return result
    except HTTPException:
        db.rollback()
        raise
    except SQLAlchemyError as exc:
        database_error(db, exc, "product_save_failed")


def delete_product_service(product_id, payload, db):
    if set(payload) != {"version"}:
        error(422, "invalid_product", "必须仅提供 version")
    version = version_value(payload)
    try:
        begin_write(db)
        product = product_or_404(db, product_id)
        if product.version != version:
            error(409, "product_stale", "产品已变化，请重新读取")
        from app.product_library_simple_quotes import archive_product
        archive_product(product, db)
        db.commit()
        return Response(status_code=204)
    except HTTPException:
        db.rollback()
        raise
    except SQLAlchemyError as exc:
        database_error(db, exc, "product_save_failed")


def preview_import_service(content, manufacturer_name, db, workbook_name=None):
    try:
        manufacturer = normalized_text(manufacturer_name, "manufacturer_name")
    except ValueError as exc:
        error(422, "invalid_product", str(exc), field="manufacturer_name")
    rows, blank = read_excel(bytes(content))
    try:
        db.execute(text("BEGIN"))
        require_identity_schema(db)
        for row in rows:
            row['workbook_name'] = workbook_name
        products = identity_products(db, rows)
        stamp = baseline(products)
        rows = preview_rows(rows, products)
        from app.product_library_details import records
        for row in rows:
            if row['errors'] or not row['existing_id'] or row['file_candidates']:
                continue
            source = v1_source_values(row, manufacturer)
            existing = next((r for r in records(db,row['existing_id'],'sources') if r.source_company == manufacturer and r.source_excerpt == source['source_excerpt']), None)
            if existing is None:
                row['allowed_actions'].append('append_source')
                row['default_action'] = 'append_source'
        db.rollback()
        created = now()
        import_id = str(uuid4())
        counts = {
            key: sum((row["category"] == key for row in rows))
            for key in ("new", "duplicate", "invalid")
        }
        counts.update(
            total_rows=len(rows),
            ignored_blank_rows=blank,
            planned_insert=counts["new"],
            planned_update=0,
        )
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            attachments = any('/media/' in name.lower() or '/drawings/' in name.lower() or name.lower().endswith('.rels') and b'/hyperlink' in archive.read(name) for name in archive.namelist())
        from app.product_library_details import ATTACHMENT_WARNING
        result = {
            "attachment_warning": ATTACHMENT_WARNING if attachments else None,
            "import_id": import_id,
            "status": "ready",
            "identity_rule_version": IDENTITY_RULE_VERSION,
            "expires_at": (created + timedelta(minutes=30)).isoformat(),
            "manufacturer_name": manufacturer,
            "counts": counts,
            "rows": rows,
        }
        begin_write(db)
        db.add(
            LibraryImport(
                import_id=import_id,
                manufacturer_name=manufacturer,
                created_at=created,
                expires_at=created + timedelta(minutes=30),
                status="ready",
                preview=result,
                baseline_digest=stamp,
            )
        )
        db.commit()
        return result
    except HTTPException:
        db.rollback()
        raise
    except SQLAlchemyError as exc:
        database_error(db, exc, "import_save_failed")


def get_import_service(import_id, db):
    item = receipt(db, import_id)
    return item.result if item.status == "committed" else item.preview


def confirm_import_service(import_id, payload, db):
    selected = decisions(payload)
    decision_hash = digest(sorted(selected.items()))
    if 'confirmed_attachments_not_imported' in payload:
        decision_hash = digest([sorted(selected.items()),payload['confirmed_attachments_not_imported']])
    try:
        begin_write(db)
        item = receipt(db, import_id)
        if item.status == "committed":
            if (item.result or {}).get("identity_rule_version") != IDENTITY_RULE_VERSION:
                if (item.result or {}).get("identity_rule_version") in (2, 3):
                    pass
                else:
                    if any(set(value) != {"action"} for value in selected.values()):
                        error(409, "idempotency_conflict", "历史回执不能改变更新目标")
                    decision_hash = digest(sorted((number, value["action"]) for number, value in selected.items()))
            if item.decision_digest != decision_hash:
                error(409, "idempotency_conflict", "同一导入已按不同决策提交")
            result = item.result
            db.rollback()
            return result
        require_identity_schema(db)
        rows = item.preview["rows"]
        if item.preview.get('attachment_warning') and payload.get('confirmed_attachments_not_imported') is not True:
            error(422,'attachments_not_confirmed',item.preview['attachment_warning'])
        if set(selected) != {row["row_number"] for row in rows} or any(
            (selected[row["row_number"]]["action"] not in row["allowed_actions"] for row in rows)
        ):
            error(422, "invalid_decisions", "必须对全部行提供合法决策")
        if not any((value["action"] != "skip" for value in selected.values())):
            error(422, "no_changes_selected", "没有选择新增或更新")
        if (
            baseline(identity_products(db, item.preview["rows"]))
            != item.baseline_digest
        ):
            error(409, "preview_stale", "该厂家产品已变化，请重新上传并确认")
        # Validate the whole plan against the immutable preview, before any writes.
        used_targets = set()
        reserved = {(p.model, p.name) for p in all_products(db)}
        for row in rows:
            choice = selected[row["row_number"]]
            if choice["action"] == "skip":
                continue
            if choice["action"] in ("update", "append_source"):
                candidate = next((c for c in row["candidate"] if c["id"] == choice.get("target_id") and c["version"] == choice.get("target_version") and (c["can_update"] or choice['action']=='append_source')), None)
                if candidate is None:
                    error(422, "invalid_update_target", "请明确选择预览中的可更新产品及版本")
                if candidate["id"] in used_targets:
                    error(409, "duplicate_update_target", "多行更新同一产品，请保留一行后重新预览")
                used_targets.add(candidate["id"])
                # An unchanged key belongs to this target; other keys remain reserved.
                same_key = row["model"] == candidate["model"] and row["name"] == candidate["name"]
            else:
                same_key = False
            if not same_key:
                key = (row["model"], row["name"])
                if key in reserved:
                    error(409, "identity_conflict", "本批计划与已有或其他计划产品精确重复，全部未写入；请重新核对")
                reserved.add(key)
        result = {
            "import_id": import_id,
            "status": "committed",
            "identity_rule_version": IDENTITY_RULE_VERSION,
            "attachments_not_imported": bool(item.preview.get("attachment_warning")),
            "inserted": 0,
            "updated": 0,
            "skipped": 0,
            "invalid": 0,
            "rows": [],
        }
        for row in rows:
            choice = selected[row["row_number"]]
            action = choice["action"]
            product_id = None
            if row["category"] == "invalid":
                outcome, reason = ("invalid", "invalid_row")
            elif action == "skip":
                outcome, reason = ("skipped", "user_skipped")
            else:
                values = {field: row[field] for field in FIELDS}
                if action == "insert":
                    product = LibraryProduct(
                        manufacturer_name=item.manufacturer_name,
                        created_at=now(),
                        updated_at=now(),
                        version=1,
                    )
                    amount = values.pop('wholesale_price', None)
                    assign(product, values)
                    db.add(product)
                    db.flush()
                    save_v1_source(db, product, row, item.manufacturer_name, amount, creating=True)
                    outcome = "inserted"
                else:
                    product = db.get(LibraryProduct, choice["target_id"])
                    if product is None or product.version != choice["target_version"]:
                        error(409, "preview_stale", "更新目标已变化，请重新预览")
                    if action == 'update':
                        assign(product, {field:value for field,value in values.items() if value is not None and field!='wholesale_price'})
                    save_v1_source(db, product, row, item.manufacturer_name, values.get('wholesale_price'))
                    product.version += 1
                    product.updated_at = now()
                    outcome = "updated"
                db.flush()
                product_id, reason = (product.id, outcome)
            result[outcome] += 1
            result["rows"].append(
                {
                    "row_number": row["row_number"],
                    "outcome": outcome,
                    "product_id": product_id,
                    "reason_code": reason,
                    "model": row["model"],
                    "name": row["name"],
                    "errors": row["errors"],
                }
            )
        item.status = "committed"
        item.decision_digest = decision_hash
        item.result = result
        item.preview = None
        item.baseline_digest = None
        db.flush()
        db.commit()
        return result
    except HTTPException:
        db.rollback()
        raise
    except SQLAlchemyError as exc:
        database_error(db, exc, "import_save_failed")


def v1_source_values(row, company):
    return dict(source_company=company,source_type='quote',notes=f'整理工作簿：{row.get("workbook_name") or "未提供"}；原件文件名未知',source_excerpt='\n'.join(f'{header}: {row.get('raw_values',row).get(field) if row.get('raw_values',row).get(field) is not None else ''}' for header,field in zip(HEADERS,FIELDS)).strip())


def save_v1_source(db, product, row, company, amount, *, creating=False):
    from app.models import LibraryProductSource, LibraryProductPrice
    from app.product_library_details import records, decimal_text, project_price
    data = v1_source_values(row,company)
    source = next((r for r in records(db,product.id,'sources') if r.source_company==company and r.source_excerpt==data['source_excerpt']),None)
    if source is None:
        source = LibraryProductSource(product_id=product.id,**data,created_at=now(),updated_at=now(),version=1)
        db.add(source)
        db.flush()
    if amount is not None:
        candidate = next((r for r in records(db,product.id,'prices') if r.source_id==source.id and r.price_type=='批发价' and r.variant_id is None),None)
        if candidate is None:
            db.add(LibraryProductPrice(product_id=product.id,source_id=source.id,price_type='批发价',amount_text=decimal_text(amount,currency='CNY'),currency_code='CNY',price_unit='件',tax_mode='unknown',freight_mode='unknown',is_default_wholesale=creating,terms_text=row.get('price_note'),created_at=now(),updated_at=now(),version=1))
    db.flush()
    project_price(db,product)
