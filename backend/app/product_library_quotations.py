"""Ephemeral customer quotations; no quotation state or product writes."""

import hashlib
import io
import json
import re
from decimal import Decimal, DecimalException, ROUND_CEILING, ROUND_HALF_UP, localcontext

from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from sqlalchemy import select, text

from app.models import LibraryProduct
from app.product_library_service import error, list_products_service, price, product_versions, product_view

RISK_WORDS = ("成本", "采购价", "进货价", "拿货价", "供货价", "底价", "批发价", "内部价", "利润", "加价率", "开票调整", "运费成本")
CUSTOMER_FIELDS = ("model", "name", "parameters", "selling_points")
INVALID_EXCEL_TEXT = re.compile(r"[\x00-\x08\x0b-\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
MAX_PRICE = Decimal("999999999.99")
RULE_DEFAULTS = {"markup_mode": "none", "markup_value": "0.00", "tax_mode": "excluded", "tax_value": "0.00", "freight_mode": "separate", "freight_value": "0.00", "rounding": "cents"}
TAX_NOTES = {
    "excluded": "以上报价为未税单价，税费另计。",
    "included": "以上报价为含税单价。",
    "adjust": "以上报价按指定开票报价口径调整，具体开票内容及适用条件以双方确认为准。",
}
FREIGHT_NOTES = {
    "separate": "运费另计，金额以订单确认为准。",
    "included": "以上报价已含运费。",
    "fixed": "以上报价已计入单件运费，整单实际费用以订单确认为准。",
}


def quotation_rules(raw):
    if not isinstance(raw, dict) or set(raw) - set(RULE_DEFAULTS):
        error(422, "invalid_quotation", "报价规则含未知字段")
    rules = RULE_DEFAULTS | raw
    for key, options in (("markup_mode", ("none", "percent", "fixed")), ("tax_mode", tuple(TAX_NOTES)), ("freight_mode", tuple(FREIGHT_NOTES)), ("rounding", ("cents", "integer"))):
        if rules[key] not in options:
            error(422, "invalid_quotation", "报价模式无效", field=key)
    for key in ("markup_value", "tax_value", "freight_value"):
        try:
            raw = rules[key]
            if not isinstance(raw, str) or not re.fullmatch(r"[0-9]+(?:\.[0-9]{1,2})?", raw):
                raise ValueError("请输入非负十进制金额或比例文本，最多两位小数")
            amount = Decimal(raw)
            if (key == "tax_value" or key == "markup_value" and rules["markup_mode"] == "percent") and amount > 1000:
                raise ValueError("比例必须为0～1000%，最多两位小数")
            rules[key] = f"{amount:.2f}"
        except ValueError as exc:
            error(422, "invalid_quotation", str(exc), field=key)
    return rules


def quotation_options(payload):
    moq = payload.get("moq", 1)
    if type(moq) is not int or not 1 <= moq <= 999999999:
        error(422, "invalid_quotation", "MOQ 必须为1～999999999的整数", field="moq")
    note = payload.get("note", "")
    if not isinstance(note, str) or len(note) > 2000 or INVALID_EXCEL_TEXT.search(note):
        error(422, "invalid_quotation", "报价说明最多2000字符，且不能含 Excel 不支持的字符", field="note")
    return moq, note.replace("\r\n", "\n").replace("\r", "\n")


def quotation_digest(items, rules, overrides=None, moq=1, note="", snapshot=None):
    return hashlib.sha256(json.dumps({"contract_version": 2, "snapshot": snapshot, "items": items, "rules": rules, "overrides": overrides or {}, "moq": moq, "note": note}, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")).hexdigest()


def text_risks(product, fields=CUSTOMER_FIELDS):
    risks = []
    for field in fields:
        value = product[field] or ""
        for word in RISK_WORDS:
            for match in re.finditer(re.escape(word), value):
                risks.append({"field": field, "word": word, "snippet": value[max(0, match.start() - 16):match.end() + 24]})
    return risks


def quote_row(product, rules, override=None):
    cost = None if product["wholesale_price"] is None else Decimal(product["wholesale_price"])
    risks, calculation_error = [], None
    final = None
    try:
        with localcontext() as context:
            context.prec = 60
            value = Decimal(override) if override is not None else cost
            if value is not None:
                if override is None:
                    markup = Decimal(rules["markup_value"])
                    if rules["markup_mode"] == "percent":
                        value *= 1 + markup / 100
                    elif rules["markup_mode"] == "fixed":
                        value += markup
                    if rules["tax_mode"] == "adjust":
                        value *= 1 + Decimal(rules["tax_value"]) / 100
                    if rules["freight_mode"] == "fixed":
                        value += Decimal(rules["freight_value"])
                value = value.quantize(Decimal("1") if rules["rounding"] == "integer" else Decimal("0.01"), rounding=ROUND_CEILING if rules["rounding"] == "integer" else ROUND_HALF_UP)
                if not value.is_finite() or value > MAX_PRICE:
                    raise ValueError("最终报价超过999999999.99")
                final = f"{value:.2f}"
                if cost is not None and value < cost:
                    risks.append("低于采购成本")
    except (DecimalException, ValueError):
        calculation_error = "最终报价超限或无法计算，请调整规则、改价或移出范围"
    if cost is None:
        risks.append("无法核对成本")
    if final is None:
        risks.append("待报价")
    invalid_fields = [field for field in CUSTOMER_FIELDS if INVALID_EXCEL_TEXT.search(product[field] or "")]
    if invalid_fields:
        calculation_error = "客户文本含 Excel 不支持的字符，请在产品资料中核查字段：" + "、".join(invalid_fields)
    return {"id": product["id"], "version": product["version"], **{field: product[field] for field in CUSTOMER_FIELDS}, "original_price": product["wholesale_price"], "final_price": final, "risks": risks, "text_risks": text_risks(product), "invalid_text_fields": invalid_fields, "error": calculation_error, "manual": override is not None}


def quotation_overrides(raw, items):
    if not isinstance(raw, dict) or set(raw) - {str(item["id"]) for item in items}:
        error(422, "invalid_quotation", "人工报价必须属于预览产品")
    result = {}
    for key, value in raw.items():
        try:
            result[key] = price(value)
            if result[key] is None:
                raise ValueError("人工报价不能为空")
        except ValueError as exc:
            error(422, "invalid_quotation", str(exc), product_id=int(key))
    return result


def preview_quotation_service(payload, db):
    if set(payload) - {"scope", "items", "filters", "rules", "include_unquoted", "overrides", "moq", "note", "selections", "conditions_confirmed_ids"}:
        error(422, "invalid_quotation", "预览含未知字段")
    if not db.in_transaction():
        db.execute(text('BEGIN'))
    moq, note = quotation_options(payload)
    scope = payload.get("scope")
    if scope not in ("selected", "filtered", "all") or type(payload.get("include_unquoted", False)) is not bool:
        error(422, "invalid_quotation", "报价范围无效")
    rules = quotation_rules(payload.get("rules", {}))
    if scope == "selected":
        items = product_versions(payload.get("items"))
        products = read_snapshot(items, db)
    else:
        filters = payload.get("filters", {}) if scope == "filtered" else {}
        allowed = {"q", "manufacturer", "sort", "price_status", "price_mode", "price_min", "price_max"}
        if not isinstance(filters, dict) or set(filters) - allowed:
            error(422, "invalid_quotation", "筛选条件无效")
        if any(not isinstance(value, str) for key, value in filters.items() if value is not None or key in ("q", "sort", "price_status")) or len(filters.get("q", "") or "") > 200:
            error(422, "invalid_quotation", "筛选字段必须为文本")
        products = list_products_service(**({"q": "", "manufacturer": None, "sort": "updated_desc"} | filters), page=1, page_size=None, db=db)
        products = products["items"]
    if not payload.get("include_unquoted", False):
        products = [product for product in products if product.get("price_state") != "unquoted"]
    items = [{"id": product["id"], "version": product["version"]} for product in products]
    # Ignore overrides outside the newly resolved scope; never silently attach them to other rows.
    raw = payload.get("overrides", {})
    if not isinstance(raw, dict):
        error(422, "invalid_quotation", "人工报价格式无效")
    overrides = quotation_overrides({key: value for key, value in raw.items() if key in {str(item["id"]) for item in items}}, items)
    selections = payload.get('selections', {})
    conditions = payload.get('conditions_confirmed_ids', [])
    if not isinstance(selections,dict) or not isinstance(conditions,list) or any(type(value) is not int for value in conditions):
        error(422,'invalid_quotation','价格选择或条件确认格式无效')
    if set(selections) - {str(product['id']) for product in products}:
        error(422,'invalid_quotation','价格选择不属于本次产品')
    rows, snapshots = [], []
    for original in products:
        product, selection_info, blocked, internal = quotation_selection(original, selections.get(str(original['id']),{}), conditions, moq, rules, db)
        snapshots.append(selection_info)
        automatic = quote_row(product, rules)
        if blocked:
            automatic['final_price'] = None
            automatic['error'] = blocked
        row = quote_row(product, rules, overrides[str(product["id"])]) if str(product["id"]) in overrides else automatic.copy()
        if blocked:
            row['final_price'] = None
            row['error'] = blocked
        row["automatic_price"] = automatic["final_price"]
        if internal.get('variant_required'):
            row['error'] = '请选择本次客户报价规格'
        row.update(internal)
        supplier_risks = source_name_risks(product, internal['sources'])
        row['text_risks'] += supplier_risks
        row['supplier_leak'] = bool(supplier_risks or any(source['source_company'] in note for source in internal['sources'] if source.get('source_company')))
        rows.append(row)
    return {"items": items, "rules": rules, "digest": quotation_digest(items, rules, overrides, moq, note, snapshots), "snapshots": snapshots, "selections": selections, "conditions_confirmed_ids": conditions, "rows": rows, "moq": moq, "note": note, "note_risks": text_risks({"note": note}, ("note",)), "terms": TAX_NOTES[rules["tax_mode"]] + FREIGHT_NOTES[rules["freight_mode"]]}


def read_snapshot(items, db):
    products = {product.id: product for product in db.scalars(select(LibraryProduct).where(LibraryProduct.id.in_([item["id"] for item in items]), LibraryProduct.archived_at.is_(None))).all()}
    affected = [{**item, "reason": "missing" if item["id"] not in products else "stale"} for item in items if item["id"] not in products or products[item["id"]].version != item["version"]]
    if affected:
        error(409, "quotation_stale", "产品已变化，请刷新受影响项并重新预览", items=affected)
    return [product_view(products[item["id"]]) for item in items]


def export_quotation_service(payload, db):
    if set(payload) - {"moq", "note", "confirmed_note", "selections", "conditions_confirmed_ids", "confirmed_below_cost_ids"} != {"items", "rules", "digest", "overrides", "confirmed_risk_ids", "confirmed_all_text"}:
        error(422, "invalid_quotation", "导出字段缺失或含未知字段")
    items = product_versions(payload["items"])
    rules = quotation_rules(payload["rules"])
    moq, note = quotation_options(payload)
    overrides = quotation_overrides(payload["overrides"], items)
    if not isinstance(payload.get('selections', {}), dict):
        error(422, 'invalid_quotation', '价格选择格式无效')
    if set(payload.get('selections', {})) - {str(item['id']) for item in items}:
        error(409, 'quotation_preview_mismatch', '报价范围已变化，请重新预览')
    preview = preview_quotation_service(dict(scope='selected',items=items,rules=rules,include_unquoted=True,overrides=overrides,moq=moq,note=note,selections=payload.get('selections',{}),conditions_confirmed_ids=payload.get('conditions_confirmed_ids',[])),db)
    if payload['digest'] != preview['digest']:
        error(409, 'quotation_preview_mismatch', '产品、选择或条件已变化，请重新预览')
    rows = preview['rows']
    invalid = [{"id": row["id"], "fields": row["invalid_text_fields"]} for row in rows if row["invalid_text_fields"]]
    if invalid:
        error(422, "customer_text_invalid", "客户文本含 Excel 不支持的字符，请核查产品字段后重新预览", errors=invalid)
    confirmations = payload["confirmed_risk_ids"]
    required = {row["id"] for row in rows if row["text_risks"]}
    if (payload["confirmed_all_text"] is not True or not isinstance(confirmations, list)
            or any(type(value) is not int for value in confirmations) or not required.issubset(confirmations)):
        error(422, "customer_text_unconfirmed", "请逐项核查敏感产品，并确认已核对全部客户可见文本")
    if any(word in note for word in RISK_WORDS) and payload.get("confirmed_note") is not True:
        error(422, "customer_text_unconfirmed", "请核查报价说明中的敏感文本")
    if any(row["error"] for row in rows):
        error(422, "quotation_overflow", "存在超限报价，请调整后重新预览")
    if any(row['supplier_leak'] for row in rows):
        error(422, 'supplier_leak', '客户内容含供应商名称，请移除后重新预览')
    below_cost = {row['id'] for row in rows if '低于采购成本' in row['risks']}
    confirmed_below = payload.get('confirmed_below_cost_ids', [])
    if not isinstance(confirmed_below, list) or any(type(value) is not int for value in confirmed_below) or not below_cost.issubset(confirmed_below):
        error(422, 'below_cost_unconfirmed', '低于采购成本的报价须明确确认')
    wb = Workbook()
    ws = wb.active
    ws.title = "客户报价"
    ws.append(["产品型号", "产品名称", "完整参数", "卖点", "最终报价", "MOQ", "报价说明"])
    for row in rows:
        values = [row[field] or "" for field in CUSTOMER_FIELDS] + [Decimal(row["final_price"]) if row["final_price"] is not None else "待报价", moq, note]
        ws.append(values)
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(ws.max_row, 5).number_format = "0.00"
    for column, width in zip("ABCDEFG", (20, 28, 55, 45, 18, 16, 65)):
        ws.column_dimensions[column].width = width
        ws[f"{column}1"].font = Font(bold=True)
    ws.freeze_panes = "A2"
    wb.properties.creator = ""
    output = io.BytesIO()
    wb.save(output)
    wb.close()
    return Response(output.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": 'attachment; filename="customer-quotation.xlsx"'})


def source_name_risks(product, sources):
    risks=[]
    for source in sources:
        name=source.get('source_company')
        if name:
            for field in CUSTOMER_FIELDS:
                if name in (product[field] or ''):
                    risks.append(dict(field=field,word=name,snippet=product[field]))
    return risks


def quotation_selection(original, selection, confirmed, moq, rules, db):
    from app.models import LibrarySimpleQuoteEvidence, LibrarySimpleQuoteKey, LibrarySimpleQuoteVersion, LibrarySimpleSource
    if not isinstance(selection, dict):
        error(422, 'invalid_quotation', '价格选择格式无效')
    product=dict(original)
    product['wholesale_price'] = None
    internal = dict(variants=[], prices=[], sources=[], selection=selection, conditions=[], variant_required=False)
    snapshot = dict(product_version=original['version'], selection=selection, customer={field: product[field] for field in CUSTOMER_FIELDS})
    if not selection:
        return product, snapshot, '请选择供应商、规格、合格报价及来源', internal
    if set(selection) != {'quote_version_id', 'source_id', 'supplier', 'variant', 'final_conditions', 'rules_confirmed'} or selection.get('rules_confirmed') is not True:
        error(422, 'invalid_quotation', '必须明确选择报价、供应商、规格、来源及最终条件')
    if any(type(selection[field]) is not int or selection[field] < 1 for field in ('quote_version_id', 'source_id')):
        error(422, 'invalid_quotation', '报价和来源 ID 无效')
    if not isinstance(selection['supplier'], str) or (selection['variant'] is not None and not isinstance(selection['variant'], str)):
        error(422, 'invalid_quotation', '供应商或规格无效')
    final_conditions = selection['final_conditions']
    if not isinstance(final_conditions, str) or len(final_conditions) > 10000 or not final_conditions.strip() or any(word in final_conditions for word in ('未知', '不详', '待确认', '待定', '另议')):
        error(422, 'invalid_quotation', '请填写明确的最终采购条件')
    quote = db.get(LibrarySimpleQuoteVersion, selection['quote_version_id'])
    key = db.get(LibrarySimpleQuoteKey, quote.quote_key_id) if quote else None
    source = db.get(LibrarySimpleSource, selection['source_id'])
    evidence = db.get(LibrarySimpleQuoteEvidence, (selection['quote_version_id'], selection['source_id']))
    if (quote is None or key is None or source is None or evidence is None or
        key.product_id != original['id'] or source.product_id != original['id'] or source.archived_at is not None or
        key.supplier_name != selection['supplier'] or source.supplier_name != key.supplier_name or
        key.variant != selection['variant'] or not quote.is_current or quote.eligibility != 'eligible' or
        quote.currency != 'CNY' or quote.unit != '件'):
        error(409, 'quotation_stale', '所选报价、来源或资格无效，请重新选择')
    source_terms = source.notes or ''
    if '含税' in source_terms and '不含税' not in source_terms and rules['tax_mode'] == 'excluded' and '客户报价未税' not in final_conditions:
        error(422, 'conditions_conflict', '来源含税而客户报价未税，请在最终条件明确写明客户报价未税')
    if ('不含税' in source_terms or '未税' in source_terms) and rules['tax_mode'] == 'included' and '客户报价含税' not in final_conditions:
        error(422, 'conditions_conflict', '来源未税而客户报价含税，请在最终条件明确写明客户报价含税')
    if '含税' in source_terms and '不含税' not in source_terms and rules['tax_mode'] == 'adjust' and '客户报价另加税费' not in final_conditions:
        error(422, 'conditions_conflict', '来源含税而客户报价另加税费，请在最终条件明确说明')
    if '含运费' in source_terms and '不含运费' not in source_terms and rules['freight_mode'] == 'fixed' and '客户报价另计单件运费' not in final_conditions:
        error(422, 'conditions_conflict', '来源含运费而客户报价增加运费，请在最终条件明确说明')
    if ('运费另计' in source_terms or '不含运费' in source_terms) and rules['freight_mode'] == 'included' and '客户报价含运费' not in final_conditions:
        error(422, 'conditions_conflict', '来源运费另计而客户报价含运费，请在最终条件明确说明')
    product['wholesale_price'] = f'{Decimal(quote.amount_cents) / 100:.2f}'
    if key.variant:
        product['name'] += '（' + key.variant + '）'
    blocked = None if original['id'] in confirmed else '请确认实际采用的来源和最终条件'
    internal['sources'] = [{'source_company': name} for name in db.scalars(select(LibrarySimpleSource.supplier_name).where(LibrarySimpleSource.product_id == original['id']).distinct())]
    internal['conditions'] = [source.notes or '来源条件未知', final_conditions.strip()]
    snapshot.update(quote_version_id=quote.id, quote_key_id=key.id, amount_cents=quote.amount_cents,
                    eligibility=quote.eligibility, supplier=key.supplier_name, variant=key.variant,
                    source_id=source.id, source_version=source.version, source_conditions=source.notes,
                    source_price_text=source.price_text, source_file_sha256=source.file_sha256,
                    source_file_name=source.file_name, source_sheet=source.sheet_name, source_row=source.row_number,
                    final_conditions=final_conditions.strip(), conditions_confirmed=original['id'] in confirmed)
    return product,snapshot,blocked,internal
